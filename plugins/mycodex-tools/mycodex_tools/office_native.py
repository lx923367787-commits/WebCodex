"""Office document backend for the local Windows desktop (no GUI automation).

Public API: inspect(source, options), edit(source, destination, operations, options),
render(source, output_dir, options), convert(source, destination, options).
The caller owns authorization, stable snapshots, backups and final publication.
This module never edits source and never attaches to the user's Office instance.

Read options: block_start/block_count (Word); sheet/range and sheet_start/count
(Excel); slide_start/slide_count/shape_start/shape_count (PowerPoint); text_offset,
max_chars, max_items; revision_view=final|original|markup. Ranges/indices returned
by inspect belong to that source snapshot. Ordinary OOXML inspection is offline.

Operations (all use an `op` discriminator):
Word OOXML target={id:'word/document.xml#p0'}; set_text(text), replace_text(find,
text, occurrence=0), set_format, append_paragraph(text,style), set_cell with
target={part,table,row,column,paragraph}; accept_revision/reject_revision with
target={id:'word/document.xml#r0'}. Legacy Word target={story:1,story_index:0,
start:0,end:10}; offsets are Word's native UTF-16 positions from inspect.
Excel set_cells(sheet,cells:[{coordinate,type,value}]), set_format(sheet,range,
...), recalculate, insert_rows/delete_rows/insert_columns/delete_columns
(sheet,start,count). Cell types: text, number, boolean, formula, empty, date.
PPT target={slide:0,shape_id:2,paragraph:0}; set_text/replace_text/set_format;
set_cell adds row/column; set_notes uses slide; set_geometry uses point units
and rotation degrees; replace_image(image_path:caller-approved absolute snapshot);
set_chart_data(categories:[str],series:[{name,values}]).
Formatting keys: font_name,font_size,bold,italic,font_color (RRGGBB), alignment
(left|center|right|justify), plus Excel number_format,fill_color,wrap_text,
column_width,row_height. Native Office saves and reopens its own working copy.
"""
from __future__ import annotations

from contextlib import contextmanager
import ctypes
from ctypes import wintypes
from datetime import date, datetime, time as datetime_time, timedelta
import gc
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import threading
import time
import traceback
import zipfile

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from mycodex_tools.common import MAX_FILE_BYTES, MAX_RESULT_BYTES, ToolError, bounded_json
from mycodex_tools import office_ooxml as ox

FORMATS = {".docx": "word", ".doc": "word", ".xlsx": "excel", ".xls": "excel",
           ".pptx": "powerpoint", ".ppt": "powerpoint", ".pdf": "word"}
EXECUTABLES = {"word": "WINWORD.EXE", "excel": "EXCEL.EXE", "powerpoint": "POWERPNT.EXE"}
PROGIDS = {"word": "Word.Application", "excel": "Excel.Application", "powerpoint": "PowerPoint.Application"}


def _input(source):
    source = Path(source).resolve(strict=True)
    if not source.is_file() or source.stat().st_size > MAX_FILE_BYTES:
        ox.fail("file_too_large", "Input is not an ordinary file within the configured size limit.")
    if source.suffix.lower() not in FORMATS:
        ox.fail("unsupported_format", "Supported files: DOC/DOCX, XLS/XLSX, PPT/PPTX and PDF-to-DOCX.")
    return source


def _destination(source, destination):
    destination = Path(destination).absolute()
    if source == destination.resolve(strict=False):
        ox.fail("invalid_path", "The backend writes a distinct working output, never source.")
    if destination.exists():
        ox.fail("target_exists", "Output already exists.")
    if not destination.parent.is_dir():
        ox.fail("parent_missing", "The output directory must exist.")
    return destination


def _options(options):
    if not isinstance(options, dict):
        ox.fail("invalid_argument", "options must be an object.")
    result = dict(options)
    result["max_chars"] = ox.integer(result.get("max_chars", 4000), 1, 32768, "max_chars")
    result["max_items"] = ox.integer(result.get("max_items", 80), 1, 200, "max_items")
    return result


def _finish(result, options, count_key=None, start_key=None):
    result.setdefault("warnings", [])
    result.setdefault("provenance", {"engine": result.get("backend", "ooxml"),
                                     "formula_evaluation": "none"})
    selection = result.get("selection", {})
    if selection.get("next_start") is not None and start_key:
        result["next_options"] = {start_key: selection["next_start"], "text_offset": 0}
        if count_key:
            result["next_options"][count_key] = options.get(count_key, options.get("max_items", 80))
    if result.get("blocks"):
        for offset, block in enumerate(result["blocks"]):
            if block.get("has_more"):
                result["next_options"] = {"block_start": selection.get("start", 0) + offset, "block_count": 1,
                                           "text_offset": block.get("text_offset", 0) + options["max_chars"]}
                break
    if result.get("slides"):
        for slide in result["slides"]:
            shapes = slide.get("shape_selection", {})
            if shapes.get("next_start") is not None:
                result["next_options"] = {"slide_start": slide["slide"], "slide_count": 1,
                                           "shape_start": shapes["next_start"], "text_offset": 0}
                break
    bounded_json(result, MAX_RESULT_BYTES)
    return result


def inspect(source: Path, options: dict) -> dict:
    source, options = _input(source), _options(options)
    ext = source.suffix.lower()
    if ext == ".docx":
        result = ox.inspect_word(ox.Package(source), {**options, "start": options.get("block_start", 0),
                                                     "count": options.get("block_count", options["max_items"])})
        return _finish(result, options, "block_count", "block_start")
    if ext == ".pptx":
        result = ox.inspect_presentation(ox.Package(source), options)
        return _finish(result, options, "slide_count", "slide_start")
    if ext == ".xlsx":
        return _finish(_inspect_excel(source, options), options, "sheet_count", "sheet_start")
    if ext == ".pdf":
        ox.fail("unsupported_operation", "Use the PDF reader; Office is only used for explicit PDF-to-DOCX conversion.")
    return _finish(_run_native("inspect", source, None, [], options), options,
                   "block_count" if ext == ".doc" else "slide_count" if ext == ".ppt" else "sheet_count",
                   "block_start" if ext == ".doc" else "slide_start" if ext == ".ppt" else "sheet_start")


def edit(source: Path, destination: Path, operations: list[dict], options: dict) -> dict:
    source, options = _input(source), _options(options)
    destination = _destination(source, destination)
    if source.suffix.lower() != destination.suffix.lower():
        ox.fail("unsupported_format", "Editing must retain the source format; use convert explicitly.")
    if not isinstance(operations, list) or not 1 <= len(operations) <= 200 or any(not isinstance(o, dict) for o in operations):
        ox.fail("invalid_argument", "operations requires 1..200 objects.")
    ext = source.suffix.lower()
    if ext in (".docx", ".pptx"):
        package = ox.Package(source)
        before = dict(package.parts)
        (ox.edit_word if ext == ".docx" else ox.edit_presentation)(package, operations)
        package.save(destination)
        unchanged = [n for n in before if n not in package.changed]
        saved = ox.Package(destination)
        if any(saved.parts.get(n) != before[n] for n in unchanged):
            ox.fail("verification_failed", "An unrelated document part changed unexpectedly.")
        return {"backend": "ooxml", "operation_count": len(operations), "verified": True,
                "changed_parts": sorted(package.changed), "unchanged_parts_verified": len(unchanged),
                "warnings": ["Replacement text inherits the first replaced run's formatting; surrounding runs remain intact."],
                "provenance": {"engine": "ooxml", "visual_verification": "use document_render"}}
    if ext == ".pdf":
        ox.fail("unsupported_operation", "PDF edits use the PDF backend.")
    result = _run_native("edit", source, destination, operations, options)
    if ext == ".xlsx":
        result["unedited_numeric_cells_verified"] = ox.verify_excel_numeric_preservation(source, destination, operations)
    return result


def render(source: Path, output_dir: Path, options: dict) -> list[dict]:
    source, options = _input(source), _options(options)
    output_dir = Path(output_dir).resolve(strict=True)
    if not output_dir.is_dir():
        ox.fail("invalid_path", "Output must be an existing directory.")
    if source.suffix.lower() == ".pdf":
        ox.fail("unsupported_operation", "Use the PDF backend to render a PDF.")
    return _run_native("render", source, output_dir, [], options)["artifacts"]


def convert(source: Path, destination: Path, options: dict) -> dict:
    source, options = _input(source), _options(options)
    destination = _destination(source, destination)
    app = FORMATS[source.suffix.lower()]
    supported = {"word": {".doc", ".docx", ".pdf"}, "excel": {".xls", ".xlsx", ".pdf"},
                 "powerpoint": {".ppt", ".pptx", ".pdf"}}
    if destination.suffix.lower() not in supported[app] or (source.suffix.lower() == ".pdf" and destination.suffix.lower() != ".docx"):
        ox.fail("unsupported_format", "Unsupported Office conversion pair.")
    if source.suffix.lower() == ".pdf":
        # Native Word PDF Reflow was observed to block on an interactive prompt.
        # Use the explicitly selected editable-text/OCR workflow, never UI input.
        from . import pdf_native
        return pdf_native.convert(source, destination, {**options, "allow_text_reflow": True})
    return _run_native("convert", source, destination, [], options)


def _inspect_excel(source, options):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries
    package = ox.Package(source)
    if package.get("xl/workbook.xml").tag != ox.q("s", "workbook"):
        ox.fail("unsupported_format", "Strict or nonstandard Excel XML requires native Office conversion first.")
    feature = ox.features(package)
    formulas = load_workbook(source, read_only=True, data_only=False, keep_links=True)
    values = load_workbook(source, read_only=True, data_only=True, keep_links=True)
    try:
        sheets = [{"name": s.title, "max_row": s.max_row, "max_column": s.max_column,
                   "state": s.sheet_state} for s in formulas.worksheets]
        selected, selection = ox._window(sheets, {**options, "start": options.get("sheet_start", 0),
                                                  "count": options.get("sheet_count", 20)})
        result = {"format": "xlsx", "backend": "ooxml", "sheets": selected,
                  "review_sheet_scopes": [f"sheet:{sheet['name']}:all_used_cells_charts_and_print_areas" for sheet in sheets],
                  "selection": selection, "sheet_count": len(sheets), "features": feature,
                  "warnings": ["Formula results are saved caches, not a fresh calculation; edit with recalculate to refresh local formulas."],
                  "provenance": {"engine": "openpyxl-read-only", "formula_evaluation": "cached"}}
        if "sheet" not in options:
            return result
        sheet = options["sheet"]
        if sheet not in formulas.sheetnames:
            ox.fail("invalid_argument", "The selected worksheet does not exist.")
        area = options.get("range", "A1:J20")
        if not isinstance(area, str) or not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]{0,6}(?::[A-Za-z]{1,3}[1-9][0-9]{0,6})?", area):
            ox.fail("invalid_argument", "range must be an A1 cell or rectangular range.")
        c1, r1, c2, r2 = range_boundaries(area)
        if c2 > 16384 or r2 > 1048576 or c2 < c1 or r2 < r1 or (c2-c1+1)*(r2-r1+1) > 2000:
            ox.fail("invalid_argument", "The Excel read range exceeds 2000 cells or worksheet bounds.")
        rows = []
        for fr, vr in zip(formulas[sheet].iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2),
                          values[sheet].iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2)):
            row = []
            for f, v in zip(fr, vr):
                item = {"coordinate": getattr(f, "coordinate", None), "value": _value(v.value),
                        "number_format": f.number_format, "type": f.data_type}
                if f.data_type == "f":
                    item["formula"] = f.value if isinstance(f.value, str) else getattr(f.value, "text", None)
                    item["value_status"] = "missing_cache" if v.value is None else "cached"
                if v.data_type == "e":
                    item["type"] = "error"
                    item["display_value"] = str(v.value)
                elif isinstance(item["value"], str):
                    item["value"] = ox._clip(item["value"], options)
                row.append(item)
            rows.append(row)
        result.update({"sheet": sheet, "range": area.upper(), "rows": rows})
        return result
    finally:
        formulas.close()
        values.close()


def _value(value):
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (date, datetime, datetime_time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return {"seconds": value.total_seconds()}
    return str(value)[:1000]


def _preflight(source, action):
    ext = source.suffix.lower()
    if ext in (".docx", ".xlsx", ".pptx"):
        feature = ox.features(ox.Package(source))
        if feature["external_link_count"] > 100:
            ox.fail("document_too_complex", "Too many external relationships for safe native opening.")
        if feature["native_open_risks"]:
            ox.fail("active_content_rejected", "Native opening is disabled for DDE, active formulas or automatic data refresh; offline reading remains available.")
        if feature["signed"] and action in ("edit", "convert"):
            ox.fail("signed_document", "Editing/conversion would invalidate the document signature.")
        unsafe_links = [x for x in feature["external_links"] if x["type"] != "hyperlink" and
                        not (ext == ".xlsx" and x["type"] == "externalLinkPath")]
        if unsafe_links:
            ox.fail("external_content_rejected", "This document may fetch external templates, images or objects; use offline operations.")
        return feature
    if ext == ".pdf":
        return {}
    try:
        import olefile
    except ImportError:
        ox.fail("dependency_missing", "Legacy Office preflight requires olefile.")
    try:
        with olefile.OleFileIO(source) as ole:
            paths = ole.listdir()
            if any(p[-1] in ("EncryptedPackage", "EncryptionInfo") for p in paths):
                ox.fail("encrypted_document", "Encrypted documents require a separate password workflow.")
            if ext == ".xls":
                workbook = next((p for p in paths if p[-1] in ("Workbook", "Book")), None)
                if workbook is None:
                    ox.fail("invalid_document", "Legacy workbook stream is missing.")
                data = ole.openstream(workbook).read(MAX_FILE_BYTES + 1)
                offset = 0
                while offset + 4 <= len(data):
                    record, size = struct.unpack_from("<HH", data, offset)
                    payload = data[offset + 4:offset + 4 + size]
                    if len(payload) != size:
                        ox.fail("invalid_document", "Legacy workbook contains a truncated record.")
                    if record == 0x2F:
                        ox.fail("encrypted_document", "Encrypted legacy workbooks are unsupported.")
                    if record == 0x85 and len(payload) >= 6 and payload[5] in (1, 6):
                        ox.fail("xlm_macro_rejected", "Excel 4.0 macro sheets are not allowed.")
                    # EXTERNNAME with DDE/OLE flags or DDE external references.
                    if record == 0x23 and len(payload) >= 2 and (struct.unpack_from("<H", payload)[0] & 0x001E):
                        ox.fail("active_content_rejected", "Legacy DDE/OLE external names cannot be opened safely.")
                    offset += 4 + size
            data = source.read_bytes()
            for token in (b"DDEAUTO", b"DDEAUTO".decode().encode("utf-16le")):
                if token in data.upper():
                    ox.fail("active_content_rejected", "DDE fields cannot be opened through native Office.")
    except ToolError:
        raise
    except (OSError, ValueError, struct.error):
        ox.fail("invalid_document", "Cannot validate the legacy Office compound document.")
    return {"legacy_format": True}


def _kernel():
    k = ctypes.WinDLL("kernel32", use_last_error=True)
    k.OpenProcess.argtypes, k.OpenProcess.restype = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD], wintypes.HANDLE
    k.CloseHandle.argtypes = [wintypes.HANDLE]
    k.GetProcessTimes.argtypes = [wintypes.HANDLE] + [ctypes.POINTER(wintypes.FILETIME)] * 4
    k.QueryFullProcessImageNameW.argtypes = [wintypes.HANDLE, wintypes.DWORD, wintypes.LPWSTR, ctypes.POINTER(wintypes.DWORD)]
    return k


def _process_identity(pid):
    k = _kernel()
    handle = k.OpenProcess(0x1000, False, pid)
    if not handle:
        return None
    try:
        times = [wintypes.FILETIME() for _ in range(4)]
        size, path = wintypes.DWORD(32768), ctypes.create_unicode_buffer(32768)
        if not k.GetProcessTimes(handle, *(ctypes.byref(t) for t in times)) or not k.QueryFullProcessImageNameW(handle, 0, path, ctypes.byref(size)):
            return None
        return {"pid": int(pid), "created": (times[0].dwHighDateTime << 32) | times[0].dwLowDateTime,
                "executable": path.value}
    finally:
        k.CloseHandle(handle)


def _processes(executable):
    if os.name != "nt":
        return {}
    k = _kernel()
    class Entry(ctypes.Structure):
        _fields_ = [("dwSize", wintypes.DWORD), ("cntUsage", wintypes.DWORD), ("th32ProcessID", wintypes.DWORD),
                    ("th32DefaultHeapID", ctypes.c_size_t), ("th32ModuleID", wintypes.DWORD),
                    ("cntThreads", wintypes.DWORD), ("th32ParentProcessID", wintypes.DWORD),
                    ("pcPriClassBase", wintypes.LONG), ("dwFlags", wintypes.DWORD), ("szExeFile", wintypes.WCHAR * 260)]
    k.CreateToolhelp32Snapshot.argtypes, k.CreateToolhelp32Snapshot.restype = [wintypes.DWORD, wintypes.DWORD], wintypes.HANDLE
    k.Process32FirstW.argtypes = k.Process32NextW.argtypes = [wintypes.HANDLE, ctypes.POINTER(Entry)]
    handle = k.CreateToolhelp32Snapshot(2, 0)
    if handle == ctypes.c_void_p(-1).value:
        ox.fail("process_identity_failed", "Cannot inventory running Office processes.")
    found, entry = {}, Entry()
    entry.dwSize = ctypes.sizeof(entry)
    try:
        more = k.Process32FirstW(handle, ctypes.byref(entry))
        while more:
            if entry.szExeFile.casefold() == executable.casefold():
                identity = _process_identity(entry.th32ProcessID)
                if identity:
                    found[entry.th32ProcessID] = identity
                else:
                    # An inaccessible matching process is still busy, not absent.
                    found[entry.th32ProcessID] = {"pid": entry.th32ProcessID, "created": None, "executable": executable}
            more = k.Process32NextW(handle, ctypes.byref(entry))
    finally:
        k.CloseHandle(handle)
    return found


@contextmanager
def _application_lock(app):
    if os.name != "nt":
        ox.fail("office_unavailable", "Native Microsoft Office operations require Windows.")
    k = _kernel()
    k.CreateMutexW.argtypes, k.CreateMutexW.restype = [ctypes.c_void_p, wintypes.BOOL, wintypes.LPCWSTR], wintypes.HANDLE
    k.WaitForSingleObject.argtypes = [wintypes.HANDLE, wintypes.DWORD]
    k.ReleaseMutex.argtypes = [wintypes.HANDLE]
    handle = k.CreateMutexW(None, False, "Local\\WebCodex.OfficeNative." + app)
    if not handle:
        ox.fail("office_busy", "Cannot acquire the Office worker lock.")
    acquired = k.WaitForSingleObject(handle, 0) in (0, 0x80)
    try:
        if not acquired:
            ox.fail("office_busy", "Another local task is using this Office application; retry after it finishes.")
        yield
    finally:
        if acquired:
            k.ReleaseMutex(handle)
        k.CloseHandle(handle)


def _terminate_owned(identity):
    if not identity or _process_identity(identity["pid"]) != identity:
        return False
    k = _kernel()
    k.TerminateProcess.argtypes = [wintypes.HANDLE, wintypes.UINT]
    handle = k.OpenProcess(0x1001, False, identity["pid"])
    if not handle:
        return False
    try:
        # Recheck immediately before termination to guard PID reuse.
        if _process_identity(identity["pid"]) != identity:
            return False
        return bool(k.TerminateProcess(handle, 1))
    finally:
        k.CloseHandle(handle)


def _guard_owned_process(identity):
    """The kernel owns cleanup even if Runner kills the whole Python process tree."""
    import win32api
    import win32job
    job = win32job.CreateJobObject(None, "")
    try:
        info = win32job.QueryInformationJobObject(job, win32job.JobObjectExtendedLimitInformation)
        info["BasicLimitInformation"]["LimitFlags"] |= win32job.JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE
        win32job.SetInformationJobObject(job, win32job.JobObjectExtendedLimitInformation, info)
        process = win32api.OpenProcess(0x0100 | 0x0001 | 0x1000, False, identity["pid"])
        try:
            if _process_identity(identity["pid"]) != identity:
                ox.fail("process_identity_failed", "Office process identity changed before lifecycle registration.")
            win32job.AssignProcessToJobObject(job, process)
        finally:
            process.Close()
        return job
    except Exception:
        job.Close()
        _terminate_owned(identity)
        ox.fail("office_lifecycle_unavailable", "Cannot bind the owned Office process to a kill-on-close job; no unguarded automation will run.")


def _run_native(action, source, destination, operations, options):
    app = FORMATS[source.suffix.lower()]
    _preflight(source, action)
    timeout = ox.integer(options.get("timeout_seconds", 55), 5, 300, "timeout_seconds")
    with _application_lock(app), tempfile.TemporaryDirectory(prefix="webcodex-office-worker-") as temp:
        journal = Path(temp) / "owned.json"
        request = {"action": action, "source": str(source), "destination": str(destination) if destination else None,
                   "operations": operations, "options": options, "journal": str(journal),
                   "owner": _process_identity(os.getpid())}
        child = subprocess.Popen([sys.executable, "-B", str(Path(__file__).resolve()), "--worker"],
                                 stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                 creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        try:
            output, diagnostic = child.communicate(json.dumps(request, ensure_ascii=False).encode("utf-8"), timeout=timeout)
        except subprocess.TimeoutExpired:
            child.kill()
            child.communicate()
            identity = json.loads(journal.read_text("utf-8")) if journal.exists() else None
            _terminate_owned(identity)
            ox.fail("office_timeout", "Office exceeded the task deadline; the source is unchanged. Inspect the working output before retrying.")
        finally:
            if child.poll() is None:
                child.kill()
                child.wait()
        try:
            result = json.loads(output)
        except (ValueError, UnicodeError):
            if journal.exists():
                _terminate_owned(json.loads(journal.read_text("utf-8")))
            ox.fail("office_failed", "Office worker did not return a valid result.")
        if "error" in result:
            raise ToolError(result["error"]["code"], result["error"]["message"])
        return result["result"]


def _owner_from_application(app, kind, before, probe=None):
    import pythoncom
    import win32process
    hwnd = None
    if kind == "excel":
        hwnd = app.Hwnd
    elif kind == "word" and probe is not None:
        hwnd = probe.Windows(1).Hwnd
    else:
        try:
            hwnd = app._oleobj_.QueryInterface(pythoncom.IID_IOleWindow).GetWindow()
        except (pythoncom.com_error, TypeError):
            try:
                hwnd = app.HWND
            except (pythoncom.com_error, AttributeError):
                pass
    if callable(hwnd):
        try:
            hwnd = hwnd()
        except (pythoncom.com_error, TypeError):
            hwnd = None
    if hwnd:
        pid = win32process.GetWindowThreadProcessId(int(hwnd))[1]
        identity = _process_identity(pid)
    else:
        # PowerPoint is a singleton. Require a new, automation-launched server,
        # not merely the first process whose executable has a familiar name.
        candidates = {p: i for p, i in _processes(EXECUTABLES[kind]).items() if p not in before}
        if kind != "powerpoint" or len(candidates) != 1 or app.Presentations.Count != 0:
            ox.fail("process_identity_failed", "Cannot prove ownership of the Office instance.")
        pid, identity = next(iter(candidates.items()))
        from win32com.client import dynamic
        locator = dynamic.Dispatch("WbemScripting.SWbemLocator")
        service = locator.ConnectServer(".", "root\\cimv2")
        rows = list(service.ExecQuery(f"SELECT CommandLine FROM Win32_Process WHERE ProcessId={pid}"))
        command = str(rows[0].CommandLine if rows else "")
        if not re.search(r"(?:/automation|/embedding|-embedding)(?:\s|$)", command, re.I):
            ox.fail("process_identity_failed", "The Office process was not launched for automation.")
    if not identity or identity["pid"] in before or not identity["created"] or Path(identity["executable"]).name.upper() != EXECUTABLES[kind]:
        ox.fail("office_busy", "Office returned an existing or unidentifiable application; it will not be controlled.")
    return identity


@contextmanager
def _office_session(kind, journal):
    import pythoncom
    import pywintypes
    from win32com.client import dynamic
    pythoncom.CoInitializeEx(pythoncom.COINIT_APARTMENTTHREADED)
    app = probe = identity = process_job = None
    original_options = []
    try:
        before = _processes(EXECUTABLES[kind])
        if kind == "powerpoint" and before:
            ox.fail("office_busy", "PowerPoint is already running. Offline PPTX editing works; native rendering/conversion requires it to be closed.")
        dispatch = pythoncom.CoCreateInstance(pywintypes.IID(PROGIDS[kind]), None,
                                             pythoncom.CLSCTX_LOCAL_SERVER, pythoncom.IID_IDispatch)
        app = dynamic.Dispatch(dispatch)
        if kind == "word":
            probe = app.Documents.Add(Visible=False)
        identity = _owner_from_application(app, kind, before, probe)
        Path(journal).write_text(json.dumps(identity), encoding="utf-8")
        process_job = _guard_owned_process(identity)
        if probe is not None:
            probe.Close(SaveChanges=0)
            probe = None
        app.AutomationSecurity = 3  # msoAutomationSecurityForceDisable
        if kind == "word":
            app.Visible = False
            app.DisplayAlerts = 0
            for name in ("UpdateLinksAtOpen", "UpdateLinksAtPrint", "UpdateFieldsAtPrint", "SaveNormalPrompt"):
                original_options.append((name, getattr(app.Options, name)))
                setattr(app.Options, name, False)
        elif kind == "excel":
            app.Visible = False
            app.DisplayAlerts = False
            app.EnableEvents = False
            app.IgnoreRemoteRequests = True
            probe = app.Workbooks.Add()
            app.Calculation = -4135  # manual, before opening any user content
            app.CalculateBeforeSave = False
        else:
            app.DisplayAlerts = 1  # ppAlertsNone (not Excel's Boolean)
        yield app, probe
    finally:
        if identity and app is not None and _process_identity(identity["pid"]) == identity:
            try:
                for name, value in original_options:
                    setattr(app.Options, name, value)
                app.Quit()
            except Exception:
                _terminate_owned(identity)
        probe = app = None
        gc.collect()
        if process_job is not None:
            process_job.Close()
        pythoncom.CoUninitialize()


def _open(app, kind, path, readonly):
    app.AutomationSecurity = 3
    path = _com_path(path, kind)
    if kind == "word":
        return app.Documents.Open(FileName=str(path), ConfirmConversions=False, ReadOnly=readonly,
                                  AddToRecentFiles=False, Visible=False, OpenAndRepair=False,
                                  NoEncodingDialog=True, PasswordDocument="", WritePasswordDocument="")
    if kind == "excel":
        return app.Workbooks.Open(Filename=str(path), UpdateLinks=0, ReadOnly=readonly,
                                  Password="", WriteResPassword="", IgnoreReadOnlyRecommended=True,
                                  Notify=False, AddToMru=False, CorruptLoad=0)
    return app.Presentations.Open(str(path), -1 if readonly else 0, 0, 0)


def _com_path(path, kind):
    """Normalize only the Office API spelling, never the authorized filesystem path."""
    value = str(path)
    if value.startswith("\\\\?\\UNC\\"):
        value = "\\\\" + value[8:]
    elif value.startswith("\\\\?\\"):
        value = value[4:]
    limit = 218 if kind == "excel" else 259
    if len(value) > limit:
        ox.fail("office_path_too_long", f"The native {kind} path exceeds {limit} characters; use a shorter project/cache path.")
    return value


def _close(doc, kind):
    if kind == "powerpoint":
        doc.Saved = True
        doc.Close()
    else:
        doc.Close(SaveChanges=False if kind == "excel" else 0)


def _word_stories(doc):
    for first in doc.StoryRanges:
        story, index = first, 0
        while story is not None:
            yield int(story.StoryType), index, story
            story, index = story.NextStoryRange, index + 1
            if index > 1000:
                ox.fail("document_too_complex", "Word story traversal exceeded its limit.")


def _inspect_word_com(doc, options):
    blocks = []
    for story_type, story_index, story in _word_stories(doc):
        for p in story.Paragraphs:
            rng = p.Range
            blocks.append({"type": "paragraph", "target": {"story": story_type, "story_index": story_index,
                           "start": int(rng.Start), "end": int(rng.End) - 1},
                           **ox._clip(str(rng.Text).rstrip("\r\x07"), options),
                           "in_table": bool(rng.Information(12))})
            if len(blocks) > 20000:
                ox.fail("document_too_complex", "Word contains too many paragraphs for the bounded reader.")
    selected, selection = ox._window(blocks, {**options, "start": options.get("block_start", 0),
                                              "count": options.get("block_count", options["max_items"])})
    revisions = []
    for revision in doc.Revisions:
        if len(revisions) >= 50:
            break
        revisions.append({"index": int(revision.Index) - 1, "type": int(revision.Type),
                          "author": str(revision.Author)[:200], **ox._clip(str(revision.Range.Text), options)})
    return {"format": "doc", "backend": "microsoft_office", "blocks": selected, "selection": selection,
            "revisions": revisions, "revision_count": int(doc.Revisions.Count), "warnings": []}


def _word_target(doc, target):
    st = ox.integer(target.get("story", 1), 1, 17, "target.story")
    si = ox.integer(target.get("story_index", 0), 0, 1000, "target.story_index")
    story = next((s for t, i, s in _word_stories(doc) if t == st and i == si), None)
    if story is None:
        ox.fail("target_not_found", "The requested Word story does not exist.")
    start = ox.integer(target.get("start"), int(story.Start), int(story.End), "target.start")
    end = ox.integer(target.get("end"), start, int(story.End), "target.end")
    rng = story.Duplicate
    rng.SetRange(start, end)
    return rng


def _rgb(value):
    if not isinstance(value, str) or not re.fullmatch(r"#?[0-9A-Fa-f]{6}", value):
        ox.fail("invalid_argument", "Colors must be six hexadecimal RGB digits.")
    v = value.lstrip("#")
    return int(v[0:2], 16) | (int(v[2:4], 16) << 8) | (int(v[4:6], 16) << 16)


def _font(font, op, powerpoint=False):
    for key, attr in (("font_name", "Name"), ("font_size", "Size"), ("bold", "Bold"), ("italic", "Italic")):
        if key in op:
            value = op[key]
            if key == "font_name":
                value = ox.text(value, key, 200)
            elif key == "font_size":
                value = ox.finite(value, 1, 400, key)
            elif type(value) is not bool:
                ox.fail("invalid_argument", f"{key} must be Boolean.")
            setattr(font, attr, (-1 if value else 0) if powerpoint and key in ("bold", "italic") else value)
    if "font_color" in op:
        if powerpoint:
            font.Color.RGB = _rgb(op["font_color"])
        else:
            font.Color = _rgb(op["font_color"])


def _edit_word_com(doc, operations):
    for op in operations:
        kind = op.get("op")
        if kind == "append_paragraph":
            rng = doc.Range(doc.Content.End - 1, doc.Content.End - 1)
            rng.InsertAfter("\r" + ox.text(op.get("text")))
            continue
        if kind in ("accept_revision", "reject_revision"):
            index = ox.integer(op.get("target", {}).get("revision"), 0, int(doc.Revisions.Count) - 1, "target.revision")
            getattr(doc.Revisions(index + 1), "Accept" if kind == "accept_revision" else "Reject")()
            continue
        rng = _word_target(doc, op.get("target", {}))
        if "expected_text" in op and str(rng.Text) != op["expected_text"]:
            ox.fail("stale_target", "Word target does not match expected_text.")
        if kind == "set_format":
            _font(rng.Font, op)
            if "alignment" in op:
                values = {"left": 0, "center": 1, "right": 2, "justify": 3}
                if op["alignment"] not in values:
                    ox.fail("invalid_argument", "Unsupported text alignment.")
                rng.ParagraphFormat.Alignment = values[op["alignment"]]
        elif kind == "set_text":
            rng.Text = ox.text(op.get("text"))
        elif kind == "replace_text":
            found = rng.Find
            found.ClearFormatting()
            found.Replacement.ClearFormatting()
            success = found.Execute(FindText=ox.text(op.get("find"), "find"), MatchCase=True,
                                    MatchWholeWord=False, MatchWildcards=False, Forward=True, Wrap=0,
                                    ReplaceWith=ox.text(op.get("text")), Replace=1)
            if not success:
                ox.fail("target_not_found", "Word replacement text was not found in the selected range.")
        else:
            ox.fail("unsupported_operation", "Unsupported legacy Word operation.")


def _excel_sheet(doc, name):
    if not isinstance(name, str) or name not in [str(s.Name) for s in doc.Worksheets]:
        ox.fail("target_not_found", "The worksheet does not exist.")
    return doc.Worksheets(name)


def _excel_range(sheet, area, maximum=2000):
    from openpyxl.utils.cell import range_boundaries
    if not isinstance(area, str) or not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]{0,6}(?::[A-Za-z]{1,3}[1-9][0-9]{0,6})?", area):
        ox.fail("invalid_argument", "Expected a bounded A1 cell or rectangle.")
    c1, r1, c2, r2 = range_boundaries(area)
    if c2 > 16384 or r2 > 1048576 or c2 < c1 or r2 < r1 or (c2-c1+1)*(r2-r1+1) > maximum:
        ox.fail("invalid_argument", "The Excel range exceeds its allowed bounds.")
    return sheet.Range(area)


def _inspect_excel_com(doc, options):
    sheets = [{"name": str(s.Name), "max_row": int(s.UsedRange.Row + s.UsedRange.Rows.Count - 1),
               "max_column": int(s.UsedRange.Column + s.UsedRange.Columns.Count - 1)} for s in doc.Worksheets]
    chosen, selection = ox._window(sheets, {**options, "start": options.get("sheet_start", 0),
                                           "count": options.get("sheet_count", 20)})
    result = {"format": "xls", "backend": "microsoft_office", "sheets": chosen, "sheet_count": len(sheets),
              "selection": selection, "warnings": ["Opening is in manual calculation mode; values may be saved caches."]}
    if "sheet" in options:
        sheet = _excel_sheet(doc, options["sheet"])
        rng = _excel_range(sheet, options.get("range", "A1:J20"))
        cells = []
        for cell in rng.Cells:
            value = _value(cell.Value2)
            item = {"coordinate": str(cell.Address).replace("$", ""), "value": value,
                    "formula": str(cell.Formula) if cell.HasFormula else None,
                    "number_format": str(cell.NumberFormat)}
            # pywin32 projects VT_ERROR/CVErr as a signed integer. The very same
            # number can be a legitimate cell value, so never infer error status
            # from its integer range. Ask Excel about the actual cell reference.
            if bool(doc.Application.WorksheetFunction.IsError(cell)):
                item.update(type="error", value=str(cell.Text), display_value=str(cell.Text))
            cells.append(item)
        result["cells"] = cells
    return result


def _edit_excel_com(app, doc, operations):
    recalculate = False
    for op in operations:
        kind = op.get("op")
        if kind == "recalculate":
            recalculate = True
            continue
        sheet = _excel_sheet(doc, op.get("sheet"))
        if kind == "set_cells":
            cells = op.get("cells")
            if not isinstance(cells, list) or not 1 <= len(cells) <= 2000:
                ox.fail("invalid_argument", "set_cells requires 1..2000 cells.")
            for item in cells:
                cell = _excel_range(sheet, item.get("coordinate"), 1)
                typ, value = item.get("type"), item.get("value")
                if typ == "empty":
                    cell.ClearContents()
                elif typ == "text":
                    value = ox.text(value, "cell.value", 32767)
                    cell.Value2 = "'" + value  # Excel quote prefix prevents formula execution.
                elif typ == "number":
                    cell.Value2 = ox.finite(value, -1e300, 1e300, "cell.value")
                elif typ == "boolean":
                    if type(value) is not bool:
                        ox.fail("invalid_argument", "Boolean cells need Boolean value.")
                    cell.Value2 = value
                elif typ == "formula":
                    value = ox.text(value, "formula", 8192)
                    if not value.startswith("=") or ox.formula_external_risk(value, reject_new_workbook_links=True):
                        ox.fail("active_content_rejected", "New formulas cannot reference external workbooks or execute data-source functions.")
                    cell.Formula2 = value
                elif typ == "date":
                    try:
                        dt = datetime.fromisoformat(value)
                        if dt.tzinfo:
                            ox.fail("invalid_argument", "Excel dates must not contain a timezone.")
                    except (TypeError, ValueError):
                        ox.fail("invalid_argument", "Date cells require ISO-8601 text.")
                    base = datetime(1904, 1, 1) if doc.Date1904 else datetime(1899, 12, 30)
                    if not doc.Date1904 and dt < datetime(1900, 3, 1):
                        base = datetime(1899, 12, 31)  # Excel's fictitious 1900-02-29
                    cell.Value2 = (dt - base).total_seconds() / 86400
                    cell.NumberFormat = "yyyy-mm-dd hh:mm:ss" if dt.time() != datetime_time() else "yyyy-mm-dd"
                else:
                    ox.fail("invalid_argument", "Unsupported Excel cell type.")
        elif kind == "set_format":
            rng = _excel_range(sheet, op.get("range"), 100000)
            _font(rng.Font, op)
            for key, attribute in (("number_format", "NumberFormat"), ("wrap_text", "WrapText"),
                                   ("column_width", "ColumnWidth"), ("row_height", "RowHeight")):
                if key in op:
                    value = op[key]
                    if key == "number_format":
                        value = ox.text(value, key, 255)
                    elif key == "wrap_text":
                        if type(value) is not bool:
                            ox.fail("invalid_argument", "wrap_text must be Boolean.")
                    else:
                        value = ox.finite(value, 0, 255 if key == "column_width" else 409, key)
                    setattr(rng, attribute, value)
            if "fill_color" in op:
                rng.Interior.Color = _rgb(op["fill_color"])
            if "alignment" in op:
                aligns = {"left": -4131, "center": -4108, "right": -4152, "justify": -4130}
                if op["alignment"] not in aligns:
                    ox.fail("invalid_argument", "Unsupported alignment.")
                rng.HorizontalAlignment = aligns[op["alignment"]]
        elif kind in ("insert_rows", "delete_rows", "insert_columns", "delete_columns"):
            columns = kind.endswith("columns")
            bound = int(sheet.Columns.Count if columns else sheet.Rows.Count)
            start = ox.integer(op.get("start"), 1, bound, "start")
            count = ox.integer(op.get("count", 1), 1, min(1000, bound - start + 1), "count")
            collection = sheet.Columns if columns else sheet.Rows
            rng = sheet.Range(collection(start), collection(start + count - 1))
            getattr(rng, "Insert" if kind.startswith("insert") else "Delete")()
        else:
            ox.fail("unsupported_operation", "Unsupported Excel operation.")
    if recalculate:
        if int(app.Workbooks.Count) != 1:
            ox.fail("office_instance_not_isolated", "Full calculation requires exactly the task workbook in the owned Excel instance.")
        _assert_no_external_recalculation(doc)
        app.CalculateFullRebuild()
        deadline = time.monotonic() + 120
        while int(app.CalculationState) != 0:
            if time.monotonic() > deadline:
                ox.fail("office_timeout", "Excel calculation did not finish in time.")
            import pythoncom
            pythoncom.PumpWaitingMessages()
            time.sleep(0.05)
    return recalculate


def _assert_no_external_recalculation(doc):
    # Also covers legacy XLS, whose BIFF expressions are not available to the
    # OOXML preflight. Reading Formula is local; no evaluation or refresh occurs.
    checked = 0
    for sheet in doc.Worksheets:
        used = sheet.UsedRange
        if used.HasFormula is False:
            continue
        for cell in used.SpecialCells(-4123).Cells:  # xlCellTypeFormulas
            checked += 1
            if checked > 20000:
                ox.fail("formula_preflight_too_large", "Recalculation formula preflight exceeds 20000 cells; use a smaller workbook.")
            if ox.formula_external_risk(str(cell.Formula), reject_new_workbook_links=True):
                ox.fail("external_refresh_rejected", "Recalculation would evaluate an external workbook reference or data-source function; existing links and caches are preserved.")


def _ppt_shapes(shapes):
    for shape in shapes:
        yield shape
        if int(shape.Type) == 6:
            yield from _ppt_shapes(shape.GroupItems)


def _ppt_shape(doc, target):
    index = ox.integer(target.get("slide"), 0, int(doc.Slides.Count) - 1, "target.slide")
    slide = doc.Slides(index + 1)
    identifier = target.get("shape_id")
    shape = next((s for s in _ppt_shapes(slide.Shapes) if int(s.Id) == identifier), None)
    if shape is None:
        ox.fail("target_not_found", "The selected PowerPoint shape does not exist.")
    return slide, shape


def _inspect_ppt_com(doc, options):
    total = int(doc.Slides.Count)
    start = ox.integer(options.get("slide_start", 0), 0, total, "slide_start")
    count = ox.integer(options.get("slide_count", 5), 1, 20, "slide_count")
    slides = []
    for i in range(start, min(total, start + count)):
        slide = doc.Slides(i + 1)
        shapes = list(_ppt_shapes(slide.Shapes))
        selected, selection = ox._window(shapes, {**options, "start": options.get("shape_start", 0),
                                                  "count": options.get("shape_count", 25)})
        items = []
        for shape in selected:
            item = {"shape_id": int(shape.Id), "name": str(shape.Name), "kind": int(shape.Type),
                    "geometry": {"left": float(shape.Left), "top": float(shape.Top),
                                 "width": float(shape.Width), "height": float(shape.Height), "units": "pt"}}
            if shape.HasTextFrame:
                item.update(ox._clip(str(shape.TextFrame.TextRange.Text), options))
            if shape.HasTable:
                table = shape.Table
                item["table"] = {"row_count": int(table.Rows.Count), "column_count": int(table.Columns.Count),
                                 "rows": [[str(table.Cell(r, c).Shape.TextFrame.TextRange.Text)[:1000]
                                           for c in range(1, min(20, table.Columns.Count) + 1)]
                                          for r in range(1, min(20, table.Rows.Count) + 1)]}
            if shape.HasChart:
                chart = shape.Chart
                item["chart_type"] = int(chart.ChartType)
                item["series"] = [{"name": str(s.Name), "values": [_value(v) for v in list(s.Values)[:100]]}
                                  for s in list(chart.SeriesCollection())[:20]]
            items.append(item)
        notes = [str(s.TextFrame.TextRange.Text) for s in slide.NotesPage.Shapes
                 if s.Type == 14 and s.PlaceholderFormat.Type == 2 and s.HasTextFrame]
        slides.append({"slide": i, "slide_id": int(slide.SlideID), "shapes": items,
                       "shape_selection": selection, "notes": ox._clip("\n".join(notes), options)})
    end = min(total, start + count)
    return {"format": "ppt", "backend": "microsoft_office", "slides": slides, "slide_count": total,
            "selection": {"start": start, "count": end-start, "total": total, "next_start": end if end < total else None},
            "warnings": []}


def _edit_ppt_com(doc, operations):
    for op in operations:
        kind, target = op.get("op"), op.get("target", {})
        if kind == "set_notes":
            index = ox.integer(target.get("slide"), 0, int(doc.Slides.Count)-1, "target.slide")
            shapes = doc.Slides(index+1).NotesPage.Shapes
            body = next((s for s in shapes if s.Type == 14 and s.PlaceholderFormat.Type == 2), None)
            if body is None:
                ox.fail("unsupported_edit", "The slide has no notes body placeholder.")
            body.TextFrame.TextRange.Text = ox.text(op.get("text"))
            continue
        slide, shape = _ppt_shape(doc, target)
        if kind == "set_geometry":
            for key, attr in (("left", "Left"), ("top", "Top"), ("width", "Width"), ("height", "Height"), ("rotation", "Rotation")):
                if key in op:
                    setattr(shape, attr, ox.finite(op[key], 0 if key in ("width", "height") else -100000, 100000, key))
        elif kind in ("set_text", "replace_text", "set_format", "set_cell"):
            if kind == "set_cell":
                if not shape.HasTable:
                    ox.fail("unsupported_edit", "The selected shape is not a table.")
                row = ox.integer(target.get("row"), 0, shape.Table.Rows.Count-1, "target.row")
                column = ox.integer(target.get("column"), 0, shape.Table.Columns.Count-1, "target.column")
                rng = shape.Table.Cell(row+1, column+1).Shape.TextFrame.TextRange
            else:
                if not shape.HasTextFrame:
                    ox.fail("unsupported_edit", "The selected shape has no text frame.")
                rng = shape.TextFrame.TextRange
            p = ox.integer(target.get("paragraph", 0), 0, rng.Paragraphs().Count-1, "target.paragraph")
            rng = rng.Paragraphs(p+1, 1)
            if kind == "set_format":
                _font(rng.Font, op, True)
                if "alignment" in op:
                    align = {"left": 1, "center": 2, "right": 3, "justify": 4}.get(op["alignment"])
                    if align is None:
                        ox.fail("invalid_argument", "Unsupported alignment.")
                    rng.ParagraphFormat.Alignment = align
            elif kind == "replace_text":
                if rng.Replace(ox.text(op.get("find"), "find"), ox.text(op.get("text")), 0, -1, 0) is None:
                    ox.fail("target_not_found", "PowerPoint replacement text was not found.")
            else:
                rng.Text = ox.text(op.get("text"))
        else:
            ox.fail("unsupported_operation", "Legacy PPT supports text, formatting, cells, geometry and notes. Convert to PPTX for surgical image/chart updates.")


def _save_as(doc, kind, destination):
    ext = destination.suffix.lower()
    output_path = _com_path(destination, kind)
    if ext == ".pdf":
        if kind == "word":
            doc.ExportAsFixedFormat(output_path, 17, OpenAfterExport=False)
        elif kind == "excel":
            doc.ExportAsFixedFormat(0, output_path, OpenAfterPublish=False)
        else:
            doc.ExportAsFixedFormat(output_path, 2)
    elif kind == "word":
        doc.SaveAs2(output_path, FileFormat=16 if ext == ".docx" else 0, AddToRecentFiles=False)
    elif kind == "excel":
        doc.SaveAs(output_path, FileFormat=51 if ext == ".xlsx" else 56, AddToMru=False)
    else:
        doc.SaveAs(output_path, 24 if ext == ".pptx" else 1)


def _saved_signature(doc, kind, operations):
    """Verify requested changes after reopening, without claiming visual identity."""
    if kind == "word":
        data = [(t, i, str(s.Text)) for t, i, s in _word_stories(doc)]
    elif kind == "powerpoint":
        data = []
        for slide in doc.Slides:
            data.append((int(slide.SlideID), [
                (int(s.Id), int(s.Type), str(s.TextFrame.TextRange.Text) if s.HasTextFrame else None,
                 round(float(s.Left), 3), round(float(s.Top), 3), round(float(s.Width), 3), round(float(s.Height), 3))
                for s in _ppt_shapes(slide.Shapes)]))
    else:
        data = {"sheets": [(str(s.Name), int(s.Shapes.Count), int(s.ChartObjects().Count), int(s.PivotTables().Count))
                            for s in doc.Worksheets], "cells": [], "formats": []}
        for op in operations:
            if op.get("op") == "set_cells":
                sheet = _excel_sheet(doc, op["sheet"])
                for item in op["cells"]:
                    c = _excel_range(sheet, item["coordinate"], 1)
                    data["cells"].append((op["sheet"], item["coordinate"], _value(c.Value2),
                                          str(c.Formula), str(c.NumberFormat)))
            elif op.get("op") == "set_format":
                rng = _excel_range(_excel_sheet(doc, op["sheet"]), op["range"], 100000)
                data["formats"].append((op["sheet"], op["range"], str(rng.NumberFormat),
                                        _value(rng.Font.Name), _value(rng.Font.Size), _value(rng.Font.Bold),
                                        _value(rng.Font.Italic), _value(rng.Font.Color)))
    return hashlib.sha256(json.dumps(data, ensure_ascii=False, default=str).encode("utf-8")).hexdigest()


def _native_work(request):
    source = Path(request["source"])
    destination = Path(request["destination"]) if request["destination"] else None
    action, options, operations = request["action"], request["options"], request["operations"]
    kind = FORMATS[source.suffix.lower()]
    with _office_session(kind, request["journal"]) as session:
        app, scratch = session
        working = source
        if action == "edit":
            shutil.copy2(source, destination)
            working = destination
        doc = _open(app, kind, working, action in ("inspect", "render", "convert"))
        if scratch is not None:
            scratch.Close(SaveChanges=False)
            scratch = None
        try:
            if action == "inspect":
                return {"word": _inspect_word_com, "excel": _inspect_excel_com, "powerpoint": _inspect_ppt_com}[kind](doc, options)
            if action == "edit":
                calculated = False
                if kind == "word":
                    _edit_word_com(doc, operations)
                elif kind == "excel":
                    calculated = _edit_excel_com(app, doc, operations)
                else:
                    _edit_ppt_com(doc, operations)
                expected = _saved_signature(doc, kind, operations)
                doc.Save()
                _close(doc, kind)
                doc = None
                doc = _open(app, kind, destination, True)
                if _saved_signature(doc, kind, operations) != expected:
                    ox.fail("verification_failed", "Requested values, formatting or object inventory changed during the Office save/reopen round trip.")
                return {"backend": "microsoft_office", "verified": True,
                        "operation_count": len(operations), "calculated": calculated,
                        "verification": "saved-and-reopened; requested values, formatting and object inventory",
                        "provenance": {"engine": PROGIDS[kind], "version": str(app.Version)},
                        "warnings": ["External data was not refreshed."] if kind == "excel" else []}
            if action == "convert":
                _save_as(doc, kind, destination)
                if not destination.is_file() or destination.stat().st_size == 0:
                    ox.fail("verification_failed", "Office did not create the converted document.")
                _close(doc, kind)
                doc = None
                if destination.suffix.lower() != ".pdf":
                    doc = _open(app, kind, destination, True)
                else:
                    from pypdf import PdfReader
                    if len(PdfReader(destination).pages) == 0:
                        ox.fail("verification_failed", "The converted PDF contains no pages.")
                return {"backend": "microsoft_office", "verified": True,
                        "output_format": destination.suffix.lstrip("."),
                        "warnings": ["PDF-to-Word conversion may change page layout; render the result before relying on its appearance."] if source.suffix.lower() == ".pdf" else [],
                        "provenance": {"engine": PROGIDS[kind], "version": str(app.Version)}}
            if action == "render":
                artifacts = []
                if kind == "powerpoint":
                    start = ox.integer(options.get("slide_start", 0), 0, int(doc.Slides.Count)-1, "slide_start")
                    count = ox.integer(options.get("slide_count", options.get("page_count", 2)), 1, 20, "slide_count")
                    for i in range(start, min(doc.Slides.Count, start+count)):
                        path = destination / f"slide-{i+1}.png"
                        if path.exists():
                            ox.fail("target_exists", "A rendered slide output already exists.")
                        width = ox.integer(options.get("width", 1280), 320, 2560, "width")
                        height = round(width * float(doc.PageSetup.SlideHeight) / float(doc.PageSetup.SlideWidth))
                        doc.Slides(i+1).Export(_com_path(path, kind), "PNG", width, height)
                        artifacts.append({"path": str(path.resolve()), "slide": i,
                                          "page_count_total": int(doc.Slides.Count),
                                          "render_source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                                          "renderer": "Microsoft PowerPoint " + str(app.Version)})
                else:
                    pdf = destination / "office-render.pdf"
                    if pdf.exists():
                        ox.fail("target_exists", "The render output already exists.")
                    _save_as(doc, kind, pdf)
                    import pypdfium2 as pdfium
                    rendered = pdfium.PdfDocument(pdf.read_bytes())
                    try:
                        start = ox.integer(options.get("page_start", 1), 1, len(rendered), "page_start")
                        count = ox.integer(options.get("page_count", 2), 1, 20, "page_count")
                        dpi = ox.integer(options.get("dpi", 120), 72, 200, "dpi")
                        for i in range(start-1, min(len(rendered), start-1+count)):
                            path = destination / f"page-{i+1}.png"
                            if path.exists():
                                ox.fail("target_exists", "A rendered page output already exists.")
                            page = rendered[i]
                            bitmap = page.render(scale=dpi/72)
                            try:
                                bitmap.to_pil().save(path)
                            finally:
                                bitmap.close()
                                page.close()
                            artifacts.append({"path": str(path.resolve()), "page": i+1,
                                              "page_count_total": len(rendered),
                                              "render_source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                                              "renderer": PROGIDS[kind] + " " + str(app.Version)})
                    finally:
                        rendered.close()
                return {"artifacts": artifacts}
            ox.fail("unsupported_operation", "Unknown native Office operation.")
        finally:
            if doc is not None:
                _close(doc, kind)
            doc = None


def _worker_main():
    try:
        request = json.loads(sys.stdin.buffer.read(1024 * 1024 + 1))
        owner = request.get("owner")
        if not owner or _process_identity(owner["pid"]) != owner:
            ox.fail("worker_owner_missing", "The native worker's owning process no longer exists.")
        kernel = _kernel()
        parent_handle = kernel.OpenProcess(0x100000 | 0x1000, False, owner["pid"])
        if not parent_handle:
            ox.fail("worker_owner_missing", "Cannot observe the native worker's owning process.")
        kernel.WaitForSingleObject.argtypes = [wintypes.HANDLE, wintypes.DWORD]
        def watch_owner():
            kernel.WaitForSingleObject(parent_handle, 0xFFFFFFFF)
            # This immediately closes the non-inheritable Office Job handle.
            os._exit(1)
        threading.Thread(target=watch_owner, name="office-owner-watch", daemon=True).start()
        result = _native_work(request)
        reply = {"result": result}
    except ToolError as error:
        reply = {"error": {"code": error.code, "message": error.message}}
    except Exception as error:
        hresult = getattr(error, "hresult", None)
        exception_info = getattr(error, "excepinfo", None)
        native_code = exception_info[5] if exception_info and len(exception_info) > 5 else None
        frames = traceback.extract_tb(error.__traceback__)
        frame = next((f for f in reversed(frames) if Path(f.filename).name == "office_native.py"), frames[-1])
        detail = ""
        if os.environ.get("MYCODEX_OFFICE_DIAGNOSTICS") == "1" and exception_info:
            description = str(exception_info[2] or "")[:800]
            description = re.sub(r"https?://\S+|[A-Za-z]:[\\/][^\r\n\"']*|\\\\[^\r\n\"']*", "<fixture-path>", description)
            detail = " Diagnostic: " + description
        reply = {"error": {"code": "office_failed", "message": "Native Office failed (" + type(error).__name__ +
                           (f", HRESULT {hresult}" if hresult is not None else "") +
                           (f", Office code {native_code}" if native_code is not None else "") +
                           f", {Path(frame.filename).name}:{frame.lineno}). The source file was not changed." + detail}}
    sys.stdout.buffer.write(json.dumps(reply, ensure_ascii=False, allow_nan=False).encode("utf-8"))


if __name__ == "__main__" and sys.argv[1:] == ["--worker"]:
    _worker_main()
