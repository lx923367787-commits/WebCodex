"""Bounded OOXML inspection and surgical edits; untouched ZIP parts are preserved.

These helpers deliberately do not render documents or evaluate spreadsheet formulas.
Object identifiers are valid only for the source snapshot from which they were read.
"""
from __future__ import annotations

from copy import deepcopy
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import posixpath
import re
import zipfile

from lxml import etree as ET

from .common import MAX_FILE_BYTES, MAX_ZIP_BYTES, ToolError, bounded_json, MAX_RESULT_BYTES

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}


def fail(code, message):
    raise ToolError(code, message)


def integer(value, low, high, name):
    if type(value) is not int or not low <= value <= high:
        fail("invalid_argument", f"{name} must be an integer from {low} to {high}.")
    return value


def finite(value, low, high, name):
    if type(value) not in (int, float) or not math.isfinite(value) or not low <= value <= high:
        fail("invalid_argument", f"{name} is outside the supported range.")
    return value


def text(value, name="text", maximum=32768):
    if not isinstance(value, str) or len(value) > maximum:
        fail("invalid_argument", f"{name} must be text of at most {maximum} characters.")
    if any(ord(c) < 32 and c not in "\t\n\r" for c in value):
        fail("invalid_argument", f"{name} contains an XML control character.")
    return value


def xml(data):
    if b"<!DOCTYPE" in data or b"<!ENTITY" in data:
        fail("invalid_document", "Document XML entities are not supported.")
    try:
        return ET.fromstring(data, ET.XMLParser(resolve_entities=False, no_network=True,
                                               load_dtd=False, huge_tree=False))
    except ET.XMLSyntaxError:
        fail("invalid_document", "Document contains malformed XML.")


def encoded(element):
    return ET.tostring(element, xml_declaration=True, encoding="UTF-8", standalone=True)


def q(prefix, local):
    return "{" + NS[prefix] + "}" + local


def formula_external_risk(formula, reject_new_workbook_links=False):
    """Classify executable references, not similarly named text/table columns.

This is a narrow no-refresh policy check, not a spreadsheet expression engine.
Existing ordinary workbook links remain readable/preservable with UpdateLinks=0;
new links and functions that themselves connect to a data source are rejected.
    """
    expression = re.sub(r'"(?:[^"]|"")*"', '""', formula)
    outside = []
    reference_bangs = set()
    quoted, depth, index = False, 0, 0
    while index < len(expression):
        character = expression[index]
        if character == "'":
            if quoted and index + 1 < len(expression) and expression[index + 1] == "'":
                outside.extend("  ")
                index += 2
                continue
            quoted = not quoted
        elif not quoted:
            if character == "[":
                depth += 1
            elif character == "]":
                depth = max(0, depth - 1)
            elif character == "!" and depth == 0:
                reference_bangs.add(index)
        outside.append(" " if quoted or depth or character in "'[]" else character)
        index += 1
    executable = "".join(outside)
    network_functions = (
        r"WEBSERVICE|IMAGE|RTD|DDE|REGISTER(?:\.ID)?|CALL|EXEC|"
        r"CUBEKPIMEMBER|CUBEMEMBER|CUBEMEMBERPROPERTY|CUBERANKEDMEMBER|"
        r"CUBESET|CUBESETCOUNT|CUBEVALUE"
    )
    if re.search(r"(?<![A-Za-z0-9_])(?:_xlfn\.)?(?:_xlws\.)?(?:" + network_functions + r")\s*\(",
                 executable, re.I):
        return "external_data_function"
    if re.search(r"\|[^!]*!", expression):
        return "dde_reference"
    if not reject_new_workbook_links:
        return None
    # A local sheet name cannot contain [, ], :, / or \\. Structured table
    # brackets occur after a sheet qualifier (if any), not in the qualifier.
    qualifiers = [m.group(1) for m in re.finditer(r"'((?:[^']|'')*)'\s*!", expression)
                  if m.end()-1 in reference_bangs]
    qualifiers.extend(m.group(1) for m in re.finditer(r"([^\s'\"!()+\-*/^&=<>%,;:{}]+)\s*!", expression)
                      if m.end()-1 in reference_bangs)
    if any(any(character in qualifier for character in "[]:/\\") for qualifier in qualifiers):
        return "external_workbook_reference"
    # External defined names may omit the sheet/! component, e.g. [Book]Name.
    # A legitimate structured reference Table1[Column] or [@Column] has no
    # identifier following its closing bracket in the same reference token.
    if re.search(r"(?<!\w)\[[^\]]+\](?:[^\W\d]|\\)[\w.\\]*", expression):
        return "external_workbook_reference"
    return None


class Package:
    def __init__(self, source):
        try:
            with zipfile.ZipFile(source) as archive:
                infos = archive.infolist()
                if len(infos) > 10000 or sum(i.file_size for i in infos) > MAX_ZIP_BYTES:
                    fail("archive_too_large", "Office package exceeds the expansion budget.")
                if len({i.filename for i in infos}) != len(infos):
                    fail("invalid_document", "Office package contains duplicate members.")
                if any(i.flag_bits & 1 for i in infos):
                    fail("encrypted_document", "Encrypted Office packages cannot be processed without a password workflow.")
                self.infos = {i.filename: i for i in infos}
                self.parts = {i.filename: archive.read(i) for i in infos}
        except (zipfile.BadZipFile, KeyError, OSError):
            fail("invalid_document", "Cannot read the Office package.")
        if "[Content_Types].xml" not in self.parts:
            fail("invalid_document", "The file is not an Office Open XML package.")
        self.changed = set()

    def get(self, name):
        if name not in self.parts:
            fail("invalid_document", "A required Office package part is missing.")
        return xml(self.parts[name])

    def put(self, name, element):
        self.parts[name] = encoded(element)
        self.changed.add(name)

    def relationships(self, part):
        folder, base = posixpath.split(part)
        relname = posixpath.join(folder, "_rels", base + ".rels")
        if relname not in self.parts:
            return relname, ET.Element(q("rel", "Relationships")), {}
        root = self.get(relname)
        links = {}
        for rel in root:
            if rel.get("TargetMode") == "External":
                links[rel.get("Id")] = None
                continue
            target = posixpath.normpath(posixpath.join(folder, rel.get("Target", ""))).lstrip("/")
            if target.startswith("../") or target not in self.parts:
                fail("invalid_document", "Office package contains an invalid relationship.")
            links[rel.get("Id")] = target
        return relname, root, links

    def bytes(self):
        stream = BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in self.parts.items():
                archive.writestr(self.infos.get(name, name), data)
        return stream.getvalue()

    def save(self, destination):
        if any(n.lower().startswith("_xmlsignatures/") for n in self.parts):
            fail("signed_document", "Editing would invalidate the document signature; create an explicitly unsigned copy first.")
        data = self.bytes()
        if len(data) > MAX_FILE_BYTES:
            fail("file_too_large", "The edited Office file exceeds 32 MiB.")
        try:
            with Path(destination).open("xb") as stream:
                stream.write(data)
        except FileExistsError:
            fail("target_exists", "Output already exists.")
        # Validate the published package, including every unchanged part.
        saved = Package(destination)
        if any(saved.parts.get(n) != data for n, data in self.parts.items()):
            fail("verification_failed", "The saved package differs from the intended changes.")


def features(package):
    names = list(package.parts)
    links = []
    risky = []
    for name in names:
        lower = name.lower()
        if "macrosheet" in lower:
            fail("xlm_macro_rejected", "Excel 4.0 macro sheets are not allowed.")
        if name.endswith(".rels"):
            for rel in package.get(name):
                if rel.get("TargetMode") == "External":
                    item = {"part": name, "type": rel.get("Type", "").rsplit("/", 1)[-1]}
                    links.append(item)
        if name.endswith(".xml") and ("word/" in name or "xl/" in name):
            data = package.parts[name]
            # Inspect executable field/formula nodes, not ordinary prose that
            # happens to mention DDE or a function name.
            root = package.get(name)
            fields = root.xpath("//w:instrText/text()|//w:fldSimple/@w:instr|//s:f/text()", namespaces=NS)
            if root.find(".//s:ddeLink", NS) is not None or any(re.search(r"\b(?:DDEAUTO|DDE)\b|\|[^!]*!", f, re.I) for f in fields):
                risky.append("dde")
            if name.startswith("word/") and any(re.search(r"\b(?:INCLUDETEXT|INCLUDEPICTURE|LINK)\b", f, re.I) for f in fields):
                risky.append("external_field")
            if any(formula_external_risk(f) for f in root.xpath("//s:f/text()", namespaces=NS)):
                risky.append("active_formula")
            if ("connections" in lower or "querytable" in lower or "pivotcachedefinition" in lower) and re.search(
                    rb'(?:refreshOnLoad|refreshOnOpen)\s*=\s*[\"\'](?:1|true)[\"\']', data, re.I):
                risky.append("automatic_external_refresh")
    return {
        "external_links": links[:100], "external_link_count": len(links),
        "macros_present": any(n.lower().endswith("vbaproject.bin") for n in names),
        "signed": any(n.lower().startswith("_xmlsignatures/") for n in names),
        "charts": sum(bool(re.fullmatch(r"(?:xl|ppt|word)/charts/chart\d+\.xml", n)) for n in names),
        "images": sum("/media/" in n for n in names),
        "embedded_objects": sum("/embeddings/" in n for n in names),
        "pivot_tables": sum(bool(re.fullmatch(r"xl/pivotTables/pivotTable\d+\.xml", n)) for n in names),
        "native_open_risks": sorted(set(risky)),
    }


def _window(items, options, prefix=""):
    start = integer(options.get(prefix + "start", 0), 0, len(items), prefix + "start")
    count = integer(options.get(prefix + "count", options.get("max_items", 60)), 1, 200, prefix + "count")
    end = min(len(items), start + count)
    return items[start:end], {"start": start, "count": end - start, "total": len(items),
                              "next_start": end if end < len(items) else None}


def _clip(value, options):
    offset = integer(options.get("text_offset", 0), 0, 2**31 - 1, "text_offset")
    limit = integer(options.get("max_chars", 2000), 1, 32768, "max_chars")
    return {"text": value[offset:offset + limit], "text_length": len(value),
            "text_offset": offset, "has_more": offset + limit < len(value)}


def _paragraph_text(paragraph, view="final", prefix="w"):
    values = []
    for node in paragraph.iter():
        if node is paragraph:
            continue
        ancestors = list(node.iterancestors())
        if any(a is not paragraph and a.tag in (q(prefix, "p"), q("w", "txbxContent"))
               for a in ancestors[:ancestors.index(paragraph)]):
            continue
        local = ET.QName(node).localname
        if prefix == "w":
            kinds = {ET.QName(a).localname for a in ancestors}
            if view == "final" and kinds & {"del", "moveFrom"}:
                continue
            if view == "original" and kinds & {"ins", "moveTo"}:
                continue
        if local in ("t", "delText"):
            values.append(node.text or "")
        elif local == "tab":
            values.append("\t")
        elif local in ("br", "cr"):
            values.append("\n")
    return "".join(values)


def inspect_word(package, options):
    view = options.get("revision_view", "final")
    if view not in ("final", "original", "markup"):
        fail("invalid_argument", "revision_view must be final, original or markup.")
    parts = ["word/document.xml"] + sorted(n for n in package.parts if re.fullmatch(
        r"word/(?:header\d+|footer\d+|footnotes|endnotes|comments)\.xml", n))
    requested = options.get("part")
    if requested is not None:
        if requested not in parts:
            fail("invalid_argument", "The selected Word story does not exist.")
        parts = [requested]
    blocks, stories, revision_items = [], [], []
    for part in parts:
        root = package.get(part)
        if ET.QName(root).namespace != NS["w"]:
            fail("unsupported_format", "Strict or nonstandard Word XML requires native Office conversion to ordinary DOCX first.")
        # Keep references: lxml may otherwise recycle Python wrapper identities.
        paragraph_nodes = list(root.iter(q("w", "p")))
        table_nodes = list(root.iter(q("w", "tbl")))
        paragraphs = {p: i for i, p in enumerate(paragraph_nodes)}
        tables = {t: i for i, t in enumerate(table_nodes)}
        before = len(blocks)
        for node in root.iter():
            if node.tag == q("w", "tbl"):
                rows = node.findall(q("w", "tr"))
                blocks.append({"id": f"{part}#t{tables[node]}", "part": part, "type": "table",
                               "rows": len(rows), "columns": max((len(r.findall(q("w", "tc"))) for r in rows), default=0)})
            elif node.tag == q("w", "p"):
                item = {"id": f"{part}#p{paragraphs[node]}", "part": part, "type": "paragraph"}
                ancestors = list(node.iterancestors())
                item["story"] = "textbox" if any(a.tag == q("w", "txbxContent") for a in ancestors) else part
                parent_table = next((a for a in ancestors if a.tag == q("w", "tbl")), None)
                if parent_table is not None:
                    item["table_id"] = f"{part}#t{tables[parent_table]}"
                    tc = next(a for a in ancestors if a.tag == q("w", "tc"))
                    tr = tc.getparent()
                    item["row"] = list(parent_table.findall(q("w", "tr"))).index(tr)
                    item["column"] = list(tr.findall(q("w", "tc"))).index(tc)
                style = node.find("w:pPr/w:pStyle", NS)
                if style is not None:
                    item["style"] = style.get(q("w", "val"))
                item.update(_clip(_paragraph_text(node, view), options))
                if options.get("include_revisions", True):
                    revisions = []
                    for r in node.iter():
                        if r.tag in {q("w", k) for k in ("ins", "del", "moveFrom", "moveTo", "rPrChange", "pPrChange")}:
                            revisions.append({"kind": ET.QName(r).localname,
                                              "id": r.get(q("w", "id")),
                                              "author": r.get(q("w", "author"), "")[:200],
                                              "date": r.get(q("w", "date"), "")})
                    if revisions:
                        item["revisions"] = revisions[:50]
                        item["revision_count"] = len(revisions)
                blocks.append(item)
        stories.append({"part": part, "block_count": len(blocks) - before})
        for index, revision in enumerate(_revision_nodes(root)):
            revision_items.append({"id": f"{part}#r{index}", "kind": ET.QName(revision).localname,
                                   "author": revision.get(q("w", "author"), "")[:200],
                                   "date": revision.get(q("w", "date"), ""),
                                   **_clip("".join(revision.xpath(".//w:t/text()|.//w:delText/text()", namespaces=NS)), options)})
    selected, page = _window(blocks, options)
    return {"format": "docx", "backend": "ooxml", "revision_view": view,
            "stories": stories[:100], "story_count": len(stories), "blocks": selected,
            "selection": page, "features": features(package),
            "revisions": revision_items[options.get("revision_start", 0):options.get("revision_start", 0)+30],
            "revision_count": len(revision_items)}


def _find_word_paragraph(package, target, roots=None, paragraph_lists=None):
    identifier = target.get("id")
    if not isinstance(identifier, str) or not re.fullmatch(
            r"word/(?:document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml#p\d+", identifier):
        fail("invalid_argument", "Word target.id must be a paragraph id from inspect.")
    part, index = identifier.rsplit("#p", 1)
    root = roots[part] if roots is not None and part in roots else package.get(part)
    paragraphs = paragraph_lists[part] if paragraph_lists is not None and part in paragraph_lists else list(root.iter(q("w", "p")))
    index = integer(int(index), 0, len(paragraphs) - 1, "paragraph index")
    if root not in paragraphs[index].iterancestors():
        fail("stale_target", "An earlier operation removed this paragraph.")
    return part, root, paragraphs[index]


def replace_paragraph(paragraph, operation, prefix="w"):
    protected = {q("w", k) for k in ("ins", "del", "moveFrom", "moveTo", "fldChar", "instrText", "txbxContent")}
    if prefix == "w" and any(n.tag in protected for n in paragraph.iter()):
        fail("unsupported_edit", "This paragraph contains revisions, fields or a text box; target a simpler text range.")
    if any(n.tag in {q(prefix, "tab"), q(prefix, "br"), q(prefix, "cr")} for n in paragraph.iter()):
        fail("unsupported_edit", "Text replacement across tabs or explicit line breaks requires separate ranges.")
    nodes = list(paragraph.iter(q(prefix, "t")))
    old = "".join(n.text or "" for n in nodes)
    if "expected_text" in operation and old != operation["expected_text"]:
        fail("stale_target", "The target text does not match expected_text.")
    replacement = text(operation.get("text"))
    if operation["op"] == "replace_text":
        find = text(operation.get("find"), "find")
        if not find:
            fail("invalid_argument", "find cannot be empty.")
        matches = [m.start() for m in re.finditer(re.escape(find), old)]
        occurrence = integer(operation.get("occurrence", 0), 0, max(0, len(matches) - 1), "occurrence")
        if not matches:
            fail("target_not_found", "The replacement text was not found in the selected paragraph.")
        start, end = matches[occurrence], matches[occurrence] + len(find)
    else:
        start, end = 0, len(old)
    if not nodes:
        run = _new_run(paragraph, prefix)
        node = ET.SubElement(run, q(prefix, "t"))
        node.text = replacement
        node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        return
    cursor = 0
    inserted = False
    for node in nodes:
        value = node.text or ""
        right = cursor + len(value)
        if right > start and cursor < end or (start == end and not inserted and right >= start):
            left_kept = value[:max(0, start - cursor)]
            right_kept = value[max(0, end - cursor):] if end < right else ""
            node.text = left_kept + (replacement if not inserted else "") + right_kept
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            inserted = True
        cursor = right
    if not inserted and old == "":
        nodes[0].text = replacement


def _new_run(paragraph, prefix):
    run = ET.Element(q(prefix, "r"))
    end = paragraph.find(q(prefix, "endParaRPr")) if prefix == "a" else None
    paragraph.insert(list(paragraph).index(end) if end is not None else len(paragraph), run)
    return run


def _revision_nodes(root):
    tags = {q("w", k) for k in ("ins", "del", "moveFrom", "moveTo", "rPrChange", "pPrChange", "tblPrChange", "trPrChange", "tcPrChange", "sectPrChange")}
    return [n for n in root.iter() if n.tag in tags]


def _revision_edit(package, operation, roots=None, revision_lists=None):
    identifier = operation.get("target", {}).get("id")
    if not isinstance(identifier, str) or not re.fullmatch(r"word/[A-Za-z0-9]+\.xml#r\d+", identifier):
        fail("invalid_argument", "Revision target.id must come from inspect.revisions.")
    part, index = identifier.rsplit("#r", 1)
    root = roots[part] if roots is not None and part in roots else package.get(part)
    nodes = revision_lists[part] if revision_lists is not None and part in revision_lists else _revision_nodes(root)
    revision = nodes[integer(int(index), 0, len(nodes)-1, "revision index")]
    if root not in revision.iterancestors():
        fail("stale_target", "An earlier operation already resolved this revision.")
    kind, parent = ET.QName(revision).localname, revision.getparent()
    accept = operation["op"] == "accept_revision"
    if kind.endswith("PrChange"):
        if not accept:
            previous = next(iter(revision), None)
            if previous is None or previous.tag != parent.tag:
                fail("unsupported_edit", "This property revision cannot be restored safely.")
            restored = deepcopy(previous)
            parent.getparent().replace(parent, restored)
        else:
            parent.remove(revision)
    elif kind in ("ins", "del", "moveFrom", "moveTo"):
        remove_content = (kind in ("del", "moveFrom")) == accept
        if parent.tag == q("w", "trPr"):
            row = parent.getparent()
            if remove_content:
                row.getparent().remove(row)
            else:
                parent.remove(revision)
        elif parent.tag == q("w", "rPr"):
            # A paragraph-mark revision can merge paragraphs and affect tables.
            fail("unsupported_edit", "Paragraph-mark revision decisions require the native Word revision backend.")
        elif remove_content:
            parent.remove(revision)
        else:
            offset = list(parent).index(revision)
            for child in list(revision):
                revision.remove(child)
                for text_node in child.iter(q("w", "delText")):
                    text_node.tag = q("w", "t")
                parent.insert(offset, child)
                offset += 1
            parent.remove(revision)
    package.put(part, root)


def _color(value):
    if not isinstance(value, str) or not re.fullmatch(r"#?[0-9A-Fa-f]{6}", value):
        fail("invalid_argument", "Colors must be six hexadecimal RGB digits.")
    return value.lstrip("#").upper()


def format_paragraph(paragraph, operation, prefix="w"):
    runs = [r for r in paragraph.iter(q(prefix, "r"))
            if next((p for p in r.iterancestors() if p.tag == q(prefix, "p")), None) is paragraph]
    if not runs:
        runs = [_new_run(paragraph, prefix)]
        if prefix == "a":
            ET.SubElement(runs[0], q("a", "t")).text = ""
    for run in runs:
        prop = run.find(q(prefix, "rPr"))
        if prop is None:
            prop = ET.Element(q(prefix, "rPr"))
            run.insert(0, prop)
        if prefix == "w":
            for key, tag in (("font_size", "sz"), ("bold", "b"), ("italic", "i"), ("font_color", "color")):
                if key not in operation:
                    continue
                node = prop.find(q("w", tag))
                if node is None:
                    node = ET.SubElement(prop, q("w", tag))
                value = operation[key]
                if key == "font_size":
                    value = str(round(finite(value, 1, 400, key) * 2))
                elif key == "font_color":
                    value = _color(value)
                    for attr in ("themeColor", "themeTint", "themeShade"):
                        node.attrib.pop(q("w", attr), None)
                else:
                    if type(value) is not bool:
                        fail("invalid_argument", f"{key} must be Boolean.")
                    value = "1" if value else "0"
                node.set(q("w", "val"), value)
            if "font_name" in operation:
                fonts = prop.find(q("w", "rFonts"))
                if fonts is None:
                    fonts = ET.SubElement(prop, q("w", "rFonts"))
                for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                    fonts.set(q("w", key), text(operation["font_name"], "font_name", 200))
                for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                    fonts.attrib.pop(q("w", key), None)
        else:
            for key, attr in (("font_size", "sz"), ("bold", "b"), ("italic", "i")):
                if key in operation:
                    value = operation[key]
                    if key == "font_size":
                        value = str(round(finite(value, 1, 400, key) * 100))
                    else:
                        if type(value) is not bool:
                            fail("invalid_argument", f"{key} must be Boolean.")
                        value = "1" if value else "0"
                    prop.set(attr, value)
            if "font_color" in operation:
                for fill in list(prop):
                    if ET.QName(fill).localname in ("solidFill", "gradFill", "pattFill", "noFill", "grpFill", "blipFill"):
                        prop.remove(fill)
                fill = ET.SubElement(prop, q("a", "solidFill"))
                ET.SubElement(fill, q("a", "srgbClr")).set("val", _color(operation["font_color"]))
            if "font_name" in operation:
                for tag in ("latin", "ea", "cs"):
                    node = prop.find(q("a", tag))
                    if node is None:
                        node = ET.SubElement(prop, q("a", tag))
                    node.set("typeface", text(operation["font_name"], "font_name", 200))
            # DrawingML character properties have a schema-defined child order.
            # In particular, solidFill must precede latin/ea/cs font elements.
            order = ["ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill",
                     "effectLst", "effectDag", "highlight", "uLnTx", "uLn", "uFillTx", "uFill",
                     "latin", "ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst"]
            children = list(prop)
            if any(ET.QName(child).localname not in order for child in children):
                fail("unsupported_edit", "Unknown character-format extensions cannot be reordered safely.")
            for child in children:
                prop.remove(child)
            prop.extend(sorted(children, key=lambda child: order.index(ET.QName(child).localname)))
    if "alignment" in operation:
        alignments = {"left": "left" if prefix == "w" else "l", "right": "right" if prefix == "w" else "r",
                      "center": "center" if prefix == "w" else "ctr", "justify": "both" if prefix == "w" else "just"}
        if operation["alignment"] not in alignments:
            fail("invalid_argument", "Unsupported alignment.")
        prop = paragraph.find(q(prefix, "pPr"))
        if prop is None:
            prop = ET.Element(q(prefix, "pPr"))
            paragraph.insert(0, prop)
        if prefix == "w":
            node = prop.find(q("w", "jc"))
            if node is None:
                node = ET.SubElement(prop, q("w", "jc"))
            node.set(q("w", "val"), alignments[operation["alignment"]])
        else:
            prop.set("algn", alignments[operation["alignment"]])


def edit_word(package, operations):
    roots = {n: package.get(n) for n in package.parts if re.fullmatch(
        r"word/(?:document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml", n)}
    paragraphs = {n: list(r.iter(q("w", "p"))) for n, r in roots.items()}
    revisions = {n: _revision_nodes(r) for n, r in roots.items()}
    for op in operations:
        kind = op.get("op")
        if kind in ("accept_revision", "reject_revision"):
            _revision_edit(package, op, roots, revisions)
        elif kind in ("replace_text", "set_text", "set_format", "replace_cell"):
            part, root, paragraph = _find_word_paragraph(package, op.get("target", {}), roots, paragraphs)
            if kind == "replace_cell":
                cell = paragraph.getparent()
                if cell.tag != q("w", "tc") or cell.find("w:tbl", NS) is not None:
                    fail("unsupported_edit", "Whole-cell replacement is supported only for ordinary text cells.")
                replace_paragraph(paragraph, {**op, "op": "set_text"})
                for other in list(cell.findall("w:p", NS)):
                    if other is not paragraph:
                        cell.remove(other)
            else:
                (format_paragraph if kind == "set_format" else replace_paragraph)(paragraph, op)
            package.put(part, root)
        elif kind == "set_cell":
            target = op.get("target", {})
            part = target.get("part", "word/document.xml")
            if not re.fullmatch(r"word/(?:document|header\d+|footer\d+)\.xml", part):
                fail("invalid_argument", "Unsupported Word table part.")
            root = roots[part]
            tables = list(root.iter(q("w", "tbl")))
            table = tables[integer(target.get("table"), 0, len(tables)-1, "target.table")]
            rows = table.findall("w:tr", NS)
            row = rows[integer(target.get("row"), 0, len(rows)-1, "target.row")]
            cells = row.findall("w:tc", NS)
            cell = cells[integer(target.get("column"), 0, len(cells)-1, "target.column")]
            paragraphs = cell.findall("w:p", NS)
            p = paragraphs[integer(target.get("paragraph", 0), 0, len(paragraphs)-1, "target.paragraph")]
            replace_paragraph(p, {**op, "op": "set_text"})
            package.put(part, root)
        elif kind == "append_paragraph":
            root = roots["word/document.xml"]
            body = root.find("w:body", NS)
            p = ET.Element(q("w", "p"))
            if "style" in op:
                prop = ET.SubElement(p, q("w", "pPr"))
                ET.SubElement(prop, q("w", "pStyle")).set(q("w", "val"), text(op["style"], "style", 200))
            ET.SubElement(ET.SubElement(p, q("w", "r")), q("w", "t")).text = text(op.get("text"))
            sect = body.find("w:sectPr", NS)
            body.insert(list(body).index(sect) if sect is not None else len(body), p)
            package.put("word/document.xml", root)
        else:
            fail("unsupported_operation", "Unsupported Word operation.")


def slide_parts(package):
    root = package.get("ppt/presentation.xml")
    if root.tag != q("p", "presentation"):
        fail("unsupported_format", "Strict or nonstandard PowerPoint XML requires native Office conversion first.")
    _, _, links = package.relationships("ppt/presentation.xml")
    return [(s.get("id"), links.get(s.get(q("r", "id")))) for s in root.findall("p:sldIdLst/p:sldId", NS)]


def _shape_nodes(root):
    return [n for n in root.iter() if n.tag in {q("p", k) for k in ("sp", "pic", "graphicFrame", "grpSp", "cxnSp")}]


def _shape_id(shape):
    nv = shape.find(".//p:cNvPr", NS)
    return int(nv.get("id")) if nv is not None else -1


def _shape_text(shape):
    body = shape.find("p:txBody", NS)
    return "\n".join(_paragraph_text(p, prefix="a") for p in body.findall("a:p", NS)) if body is not None else ""


def inspect_presentation(package, options):
    entries = slide_parts(package)
    chosen, selection = _window(entries, {**options, "start": options.get("slide_start", 0),
                                         "count": options.get("slide_count", 5)})
    result = []
    for offset, (slide_id, part) in enumerate(chosen):
        root = package.get(part)
        _, rels, links = package.relationships(part)
        shapes = _shape_nodes(root)
        selected, shape_selection = _window(shapes, {**options, "start": options.get("shape_start", 0),
                                                     "count": options.get("shape_count", 25)})
        items = []
        for shape in selected:
            item = {"shape_id": _shape_id(shape), "kind": ET.QName(shape).localname,
                    **_clip(_shape_text(shape), options)}
            xfrm = shape.find("p:spPr/a:xfrm", NS)
            if xfrm is None:
                xfrm = shape.find("p:xfrm", NS)
            if xfrm is not None:
                off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
                if off is not None and ext is not None:
                    item["geometry"] = {"left": int(off.get("x", 0)) / 12700,
                                        "top": int(off.get("y", 0)) / 12700,
                                        "width": int(ext.get("cx", 0)) / 12700,
                                        "height": int(ext.get("cy", 0)) / 12700,
                                        "rotation": int(xfrm.get("rot", 0)) / 60000,
                                        "units": "pt", "coordinate_space": "parent"}
            blip = shape.find(".//a:blip", NS)
            if blip is not None:
                item["image_part"] = links.get(blip.get(q("r", "embed")))
            chart = shape.find(".//c:chart", NS)
            if chart is not None:
                chart_part = links.get(chart.get(q("r", "id")))
                item["chart_part"] = chart_part
                if chart_part and options.get("include_chart_data", True):
                    chart_root = package.get(chart_part)
                    series = []
                    for s in chart_root.findall(".//c:ser", NS)[:30]:
                        series.append({"name": "".join(s.xpath("./c:tx//c:v/text()", namespaces=NS))[:200],
                                       "categories": s.xpath("./c:cat//c:pt/c:v/text()", namespaces=NS)[:100],
                                       "values": s.xpath("./c:val//c:pt/c:v/text()", namespaces=NS)[:100]})
                    item["chart_series"] = series
            table = shape.find(".//a:tbl", NS)
            if table is not None:
                rows = table.findall("a:tr", NS)
                item["table"] = {"row_count": len(rows), "rows": [
                    ["\n".join(_paragraph_text(p, prefix="a") for p in tc.findall("a:txBody/a:p", NS))[:1000]
                     for tc in row.findall("a:tc", NS)[:20]] for row in rows[:20]]}
            items.append(item)
        slide = {"slide": selection["start"] + offset, "slide_id": slide_id,
                 "shapes": items, "shape_selection": shape_selection}
        if options.get("include_notes", True):
            notes_rel = next((r for r in rels if r.get("Type", "").endswith("/notesSlide")), None)
            if notes_rel is not None and links.get(notes_rel.get("Id")):
                notes_root = package.get(links[notes_rel.get("Id")])
                notes = [s for s in _shape_nodes(notes_root) if s.find('.//p:ph[@type="body"]', NS) is not None]
                slide["notes"] = _clip("\n".join(_shape_text(s) for s in notes), options)
        result.append(slide)
    return {"format": "pptx", "backend": "ooxml", "slide_count": len(entries),
            "selection": selection, "slides": result, "features": features(package)}


def _locate_shape(package, target):
    entries = slide_parts(package)
    index = integer(target.get("slide"), 0, len(entries) - 1, "target.slide")
    part = entries[index][1]
    root = package.get(part)
    shape_id = integer(target.get("shape_id"), 0, 2**31 - 1, "target.shape_id")
    matches = [s for s in _shape_nodes(root) if _shape_id(s) == shape_id]
    if len(matches) != 1:
        fail("target_not_found", "The selected PowerPoint shape does not exist uniquely.")
    return part, root, matches[0]


def _replace_image(package, part, shape, image_path):
    from PIL import Image
    try:
        data = Path(image_path).read_bytes()
        if len(data) > MAX_FILE_BYTES:
            fail("file_too_large", "Image exceeds the supported file size.")
        with Image.open(BytesIO(data)) as img:
            fmt = img.format
            img.verify()
        if fmt not in ("PNG", "JPEG"):
            fail("unsupported_format", "Replacement images must be PNG or JPEG.")
    except (OSError, ValueError, Image.DecompressionBombError):
        fail("invalid_image", "The replacement image cannot be read.")
    if shape.tag != q("p", "pic"):
        fail("unsupported_edit", "replace_image requires a picture shape.")
    blip = shape.find(".//a:blip", NS)
    if blip is None:
        fail("unsupported_edit", "The selected picture has no embedded image.")
    ext, mime = ("png", "image/png") if fmt == "PNG" else ("jpg", "image/jpeg")
    name = "ppt/media/mycodex_" + hashlib.sha256(data).hexdigest()[:24] + "." + ext
    package.parts[name] = data
    package.changed.add(name)
    relname, rels, _ = package.relationships(part)
    existing = {r.get("Id") for r in rels}
    index = 1
    while f"rId{index}" in existing:
        index += 1
    identifier = f"rId{index}"
    ET.SubElement(rels, q("rel", "Relationship"), Id=identifier,
                  Type=NS["r"] + "/image", Target=posixpath.relpath(name, posixpath.dirname(part)))
    blip.set(q("r", "embed"), identifier)
    blip.attrib.pop(q("r", "link"), None)
    package.put(relname, rels)
    content = package.get("[Content_Types].xml")
    if not any(e.get("Extension") == ext for e in content):
        ET.SubElement(content, q("ct", "Default"), Extension=ext, ContentType=mime)
        package.put("[Content_Types].xml", content)


def _replace_chart_data(package, target, operation):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    presentation = Presentation(BytesIO(package.bytes()))
    slide = presentation.slides[target["slide"]]
    def walk(shapes):
        for shape in shapes:
            yield shape
            if shape.shape_type == 6:
                yield from walk(shape.shapes)
    shape = next((s for s in walk(slide.shapes) if s.shape_id == target["shape_id"]), None)
    if shape is None or not shape.has_chart:
        fail("unsupported_edit", "set_chart_data requires a chart shape.")
    chart = shape.chart
    chart_part = str(chart.part.partname).lstrip("/")
    references = 0
    for _, slide_part in slide_parts(package):
        _, _, relationships = package.relationships(slide_part)
        for chart_node in package.get(slide_part).findall(".//c:chart", NS):
            references += relationships.get(chart_node.get(q("r", "id"))) == chart_part
    if references != 1:
        fail("unsupported_edit", "This chart is shared by multiple shapes; replacing its dataset would change other shapes.")
    root = package.get(chart_part)
    plots = root.findall("c:chart/c:plotArea/*", NS)
    plot_tags = [ET.QName(p).localname for p in plots if ET.QName(p).localname.endswith("Chart")]
    if len(plot_tags) != 1 or plot_tags[0] not in ("barChart", "lineChart", "pieChart", "doughnutChart", "areaChart", "radarChart"):
        fail("unsupported_edit", "Only a single ordinary category chart can replace its dataset.")
    workbook_part = chart.part.chart_workbook.xlsx_part
    if workbook_part is None:
        fail("unsupported_edit", "The chart has no local embedded workbook; linked data is not refreshed.")
    workbook_name = str(workbook_part.partname).lstrip("/")
    workbook_references = 0
    for name in package.parts:
        if re.fullmatch(r"ppt/charts/chart\d+\.xml", name):
            _, _, chart_links = package.relationships(name)
            workbook_references += sum(target_part == workbook_name for target_part in chart_links.values())
    if workbook_references != 1:
        fail("unsupported_edit", "This chart workbook is shared; replacing it could change another chart.")
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    embedded_package = Package(BytesIO(workbook_part.blob))
    embedded_features = features(embedded_package)
    if any(embedded_features[k] for k in ("charts", "images", "external_link_count", "macros_present", "pivot_tables", "embedded_objects")):
        fail("unsupported_edit", "The embedded chart workbook contains other objects; replacing it would lose unrelated data.")
    allowed_cells = set()
    for formula in root.xpath(".//c:f/text()", namespaces=NS):
        match = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^!]+))!([A-Za-z$0-9:]+)", formula)
        if not match:
            fail("unsupported_edit", "Chart data references use an unsupported form.")
        sheet_name = (match.group(1) or match.group(2)).replace("''", "'")
        c1, r1, c2, r2 = range_boundaries(match.group(3))
        if (r2-r1+1)*(c2-c1+1) > 5000:
            fail("unsupported_edit", "Chart data exceeds the bounded dataset editor.")
        allowed_cells.update((sheet_name, f"{get_column_letter(c)}{r}") for r in range(r1, r2+1) for c in range(c1, c2+1))
    embedded = load_workbook(BytesIO(workbook_part.blob), data_only=False)
    try:
        if len(embedded.worksheets) != 1 or any((s.title, c.coordinate) not in allowed_cells
                for s in embedded.worksheets for c in s._cells.values() if c.value is not None):
            fail("unsupported_edit", "The embedded workbook contains data beyond this chart; it cannot be replaced safely.")
    finally:
        embedded.close()
    categories, series = operation.get("categories"), operation.get("series")
    if not isinstance(categories, list) or not 1 <= len(categories) <= 200 or not isinstance(series, list) or not 1 <= len(series) <= 20:
        fail("invalid_argument", "Charts require 1..200 categories and 1..20 series.")
    data = CategoryChartData()
    data.categories = [text(v, "category", 200) for v in categories]
    for item in series:
        values = item.get("values")
        if not isinstance(values, list) or len(values) != len(categories):
            fail("invalid_argument", "Each chart series needs one value per category.")
        data.add_series(text(item.get("name"), "series.name", 200),
                        [None if v is None else finite(v, -1e100, 1e100, "chart value") for v in values])
    # Only the chart and its explicitly replaced dataset are serialized. All
    # unrelated parts, including animation/SmartArt/master extensions, stay exact.
    chart.replace_data(data)
    package.parts[chart_part] = chart.part.blob
    package.changed.add(chart_part)
    workbook_part = chart.part.chart_workbook.xlsx_part
    workbook_name = str(workbook_part.partname).lstrip("/")
    package.parts[workbook_name] = workbook_part.blob
    package.changed.add(workbook_name)


def _create_notes(package, slide_index, value):
    """Use the library's notes template, importing only new/notes-related parts."""
    from pptx import Presentation
    presentation = Presentation(BytesIO(package.bytes()))
    notes = presentation.slides[slide_index].notes_slide
    if notes.notes_text_frame is None:
        fail("unsupported_edit", "Cannot create a standard notes body placeholder.")
    notes.notes_text_frame.text = text(value)
    stream = BytesIO()
    presentation.save(stream)
    created = Package(BytesIO(stream.getvalue()))
    selected_part = slide_parts(package)[slide_index][1]
    relname = posixpath.join(posixpath.dirname(selected_part), "_rels", posixpath.basename(selected_part) + ".rels")
    allowed = {"[Content_Types].xml", "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels", relname}
    for name, data in created.parts.items():
        if name not in package.parts or name in allowed or name.startswith("ppt/notesSlides/") and name not in package.parts:
            if package.parts.get(name) != data:
                package.parts[name] = data
                package.changed.add(name)


def edit_presentation(package, operations):
    for op in operations:
        kind, target = op.get("op"), op.get("target", {})
        if kind == "set_notes":
            entries = slide_parts(package)
            slide = integer(target.get("slide"), 0, len(entries) - 1, "target.slide")
            part = entries[slide][1]
            _, rels, links = package.relationships(part)
            rel = next((r for r in rels if r.get("Type", "").endswith("/notesSlide")), None)
            if rel is None or not links.get(rel.get("Id")):
                _create_notes(package, slide, op.get("text"))
                continue
            notes_part = links[rel.get("Id")]
            root = package.get(notes_part)
            shape = next((s for s in _shape_nodes(root) if s.find('.//p:ph[@type="body"]', NS) is not None), None)
            if shape is None:
                fail("unsupported_edit", "This notes page has no body placeholder.")
            body = shape.find("p:txBody", NS)
            paragraphs = body.findall("a:p", NS)
            if not paragraphs:
                paragraphs = [ET.SubElement(body, q("a", "p"))]
            replace_paragraph(paragraphs[0], {"op": "set_text", "text": op.get("text")}, "a")
            for p in paragraphs[1:]:
                body.remove(p)
            package.put(notes_part, root)
            continue
        if kind == "set_chart_data":
            _locate_shape(package, target)
            _replace_chart_data(package, target, op)
            continue
        part, root, shape = _locate_shape(package, target)
        if kind in ("set_text", "replace_text", "set_format"):
            body = shape.find("p:txBody", NS)
            if body is None:
                fail("unsupported_edit", "The selected shape has no text body.")
            paragraphs = body.findall("a:p", NS)
            index = integer(target.get("paragraph", 0), 0, len(paragraphs) - 1, "target.paragraph")
            (format_paragraph if kind == "set_format" else replace_paragraph)(paragraphs[index], op, "a")
        elif kind == "set_cell":
            table = shape.find(".//a:tbl", NS)
            if table is None:
                fail("unsupported_edit", "The selected shape is not a table.")
            rows = table.findall("a:tr", NS)
            row = rows[integer(target.get("row"), 0, len(rows) - 1, "target.row")]
            cells = row.findall("a:tc", NS)
            cell = cells[integer(target.get("column"), 0, len(cells) - 1, "target.column")]
            paragraphs = cell.findall("a:txBody/a:p", NS)
            index = integer(target.get("paragraph", 0), 0, len(paragraphs) - 1, "target.paragraph")
            replace_paragraph(paragraphs[index], {**op, "op": "set_text"}, "a")
        elif kind == "set_geometry":
            xfrm = shape.find("p:spPr/a:xfrm", NS)
            if xfrm is None:
                xfrm = shape.find("p:xfrm", NS)
            if xfrm is None:
                fail("unsupported_edit", "This shape inherits its geometry; explicit geometry is required.")
            for key, tag, attr in (("left", "off", "x"), ("top", "off", "y"), ("width", "ext", "cx"), ("height", "ext", "cy")):
                if key in op:
                    node = xfrm.find("a:" + tag, NS)
                    if node is None:
                        node = ET.SubElement(xfrm, q("a", tag))
                    value = finite(op[key], 0 if key in ("width", "height") else -100000, 100000, key)
                    node.set(attr, str(round(value * 12700)))
            if "rotation" in op:
                xfrm.set("rot", str(round(finite(op["rotation"], -3600, 3600, "rotation") * 60000)))
        elif kind == "replace_image":
            _replace_image(package, part, shape, op.get("image_path"))
        else:
            fail("unsupported_operation", "Unsupported PowerPoint operation.")
        package.put(part, root)


def edit_workbook_cells(package, updates):
    """Basic cell-only path for the existing XLSX API; never resave other objects."""
    from openpyxl.utils.cell import coordinate_to_tuple
    workbook = package.get("xl/workbook.xml")
    _, _, links = package.relationships("xl/workbook.xml")
    sheets = {s.get("name"): links.get(s.get(q("r", "id"))) for s in workbook.findall("s:sheets/s:sheet", NS)}
    for update in updates:
        part = sheets.get(update.get("sheet"))
        if part is None:
            fail("invalid_argument", "The worksheet does not exist.")
        root = package.get(part)
        sheet_data = root.find("s:sheetData", NS)
        if sheet_data is None:
            fail("invalid_document", "The worksheet has no sheetData element.")
        for item in update["cells"]:
            coordinate = item["coordinate"].upper()
            row_index, column_index = coordinate_to_tuple(coordinate)
            if row_index > 1048576 or column_index > 16384:
                fail("invalid_argument", "Cell coordinate exceeds worksheet bounds.")
            row = next((r for r in sheet_data if int(r.get("r", 0)) == row_index), None)
            if row is None:
                row = ET.Element(q("s", "row"), r=str(row_index))
                preceding = next((i for i, r in enumerate(sheet_data) if int(r.get("r", 0)) > row_index), len(sheet_data))
                sheet_data.insert(preceding, row)
            cell = next((c for c in row if c.get("r") == coordinate), None)
            if cell is None:
                cell = ET.Element(q("s", "c"), r=coordinate)
                preceding = next((i for i, c in enumerate(row) if coordinate_to_tuple(c.get("r"))[1] > column_index), len(row))
                row.insert(preceding, cell)
            for child in list(cell):
                if child.tag in (q("s", "v"), q("s", "f"), q("s", "is")):
                    cell.remove(child)
            cell.attrib.pop("t", None)
            typ, value = item["type"], item.get("value")
            if typ == "text":
                cell.set("t", "inlineStr")
                node = ET.SubElement(ET.SubElement(cell, q("s", "is")), q("s", "t"))
                node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                node.text = value
            elif typ == "formula":
                if formula_external_risk(value, reject_new_workbook_links=True):
                    fail("active_content_rejected", "New formulas cannot reference external workbooks or execute data-source functions.")
                ET.SubElement(cell, q("s", "f")).text = value[1:]
            elif typ == "number":
                serial = str(value) if type(value) is int else format(value, ".16g")
                if float(serial) != value:
                    fail("verification_failed", "Excel would lose numeric precision; no output was published.")
                ET.SubElement(cell, q("s", "v")).text = serial
            elif typ == "boolean":
                cell.set("t", "b")
                ET.SubElement(cell, q("s", "v")).text = "1" if value else "0"
            elif typ != "empty":
                fail("invalid_argument", "Unsupported basic cell type.")
        # Dimension is advisory, but a stale value hides newly inserted cells
        # from read-only readers. Recompute from actual stored cells, not a rectangle.
        positions = [coordinate_to_tuple(c.get("r")) for c in sheet_data.iter(q("s", "c")) if c.get("r")]
        from openpyxl.utils.cell import get_column_letter
        dimension = root.find("s:dimension", NS)
        if dimension is not None and positions:
            r1, r2 = min(r for r, c in positions), max(r for r, c in positions)
            c1, c2 = min(c for r, c in positions), max(c for r, c in positions)
            dimension.set("ref", f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}")
        package.put(part, root)
    # Dependent formulas can span every sheet. Remove caches that are no longer
    # authoritative, preserving expressions, external-link parts and all objects.
    for part in sheets.values():
        root = package.get(part)
        changed = False
        for cell in root.findall(".//s:c", NS):
            formula = cell.find("s:f", NS)
            if formula is not None and not formula_external_risk(formula.text or "", reject_new_workbook_links=True):
                for value in cell.findall("s:v", NS):
                    cell.remove(value)
                    changed = True
        if changed:
            package.put(part, root)
    calc = workbook.find("s:calcPr", NS)
    if calc is None:
        calc = ET.SubElement(workbook, q("s", "calcPr"))
    calc.set("fullCalcOnLoad", "1")
    calc.set("forceFullCalc", "1")
    package.put("xl/workbook.xml", workbook)


def verify_excel_numeric_preservation(source, destination, operations):
    """Detect native Excel rounding of untouched numbers before final publication.

Compare raw worksheet numbers with Decimal, independent of number/date display
formats. Track explicitly requested row/column shifts and replacements. Formula
result cells are excluded because native recalculation intentionally changes them.
"""
    from decimal import Decimal, InvalidOperation
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
    before, after = Package(source), Package(destination)

    def sheet_parts(package):
        _, _, links = package.relationships("xl/workbook.xml")
        return {s.get("name"): links.get(s.get(q("r", "id"))) for s in package.get("xl/workbook.xml").findall("s:sheets/s:sheet", NS)}

    original_sheets, saved_sheets = sheet_parts(before), sheet_parts(after)
    if list(original_sheets) != list(saved_sheets):
        fail("verification_failed", "Native Excel changed the worksheet inventory unexpectedly.")
    checked = 0
    for name, part in original_sheets.items():
        root, saved = before.get(part), after.get(saved_sheets[name])
        target_cells = {c.get("r"): c for c in saved.findall(".//s:sheetData/s:row/s:c", NS)}
        derived_ranges = [range_boundaries(f.get("ref")) for f in root.findall(".//s:f", NS)
                          if f.get("ref") and f.get("t") in ("array", "dataTable")]
        for cell in root.findall(".//s:sheetData/s:row/s:c", NS):
            value = cell.find("s:v", NS)
            if cell.find("s:f", NS) is not None or cell.get("t", "n") != "n" or value is None or not value.text:
                continue
            row, column = coordinate_to_tuple(cell.get("r"))
            if any(c1 <= column <= c2 and r1 <= row <= r2 for c1, r1, c2, r2 in derived_ranges):
                continue
            replaced = False
            for op in operations:
                if op.get("sheet") != name:
                    continue
                kind = op.get("op")
                if kind == "set_cells":
                    current = f"{get_column_letter(column)}{row}"
                    if any(c.get("coordinate", "").upper() == current for c in op["cells"]):
                        replaced = True
                        break
                elif kind in ("insert_rows", "delete_rows", "insert_columns", "delete_columns"):
                    axis = column if kind.endswith("columns") else row
                    start, count = op["start"], op.get("count", 1)
                    if kind.startswith("delete") and start <= axis < start + count:
                        replaced = True
                        break
                    if axis >= start:
                        axis += count if kind.startswith("insert") else -count
                    if kind.endswith("columns"):
                        column = axis
                    else:
                        row = axis
            if replaced:
                continue
            actual = target_cells.get(f"{get_column_letter(column)}{row}")
            actual_value = actual.find("s:v", NS) if actual is not None else None
            try:
                equal = actual is not None and actual.get("t", "n") == "n" and actual.find("s:f", NS) is None and actual_value is not None and Decimal(actual_value.text) == Decimal(value.text)
            except (InvalidOperation, TypeError):
                equal = False
            if not equal:
                fail("verification_failed", "Native Excel changed an untouched numeric value or its type; the original will not be replaced.")
            checked += 1
    return checked
