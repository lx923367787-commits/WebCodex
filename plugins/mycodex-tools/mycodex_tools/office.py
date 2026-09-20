"""Bounded Office operations for the native WebCodex plugin."""

from __future__ import annotations

import math
import re
import unicodedata
from datetime import date, datetime, time, timedelta
from io import BytesIO

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from .common import ToolError, arr, boolean, integer, number, obj, string


def _check(condition, message):
    if not condition:
        raise ToolError("invalid_argument", message)


def _bounded(value, low, high, name):
    _check(type(value) is int and low <= value <= high, f"{name} must be {low}..{high}")
    return value


def _selection(args, prefix, total, default, maximum=200):
    start = _bounded(args.get(f"{prefix}_start", 0), 0, total, f"{prefix}_start")
    count = _bounded(args.get(f"{prefix}_count", default), 0, maximum, f"{prefix}_count")
    end = min(total, start + count)
    return start, end, {"start": start, "count": end - start, "total": total, "has_more": end < total}


def _text(text, args):
    offset = _bounded(args.get("text_offset", 0), 0, 2**31 - 1, "text_offset")
    count = _bounded(args.get("max_chars", 4000), 1, 32768, "max_chars")
    return {"text": text[offset : offset + count], "text_offset": offset,
            "text_length": len(text), "has_more": offset + count < len(text)}


def _load(ctx, args, extension, loader):
    data, digest = ctx.read_bytes(args["path"], extension,
                                  expected_sha256=args.get("expected_sha256"))
    try:
        return loader(BytesIO(data)), digest
    except Exception as exc:
        raise ToolError("invalid_document", f"Cannot open {extension} document: {type(exc).__name__}") from exc


def _table_read(table, args, index):
    start, end, selected_rows = _selection(args, "row", len(table.rows), 20)
    column_start, column_end, selected_columns = _selection(args, "column", len(table.columns), 10)
    return {"index": index, "row_count": len(table.rows), "column_count": len(table.columns),
            "selected_rows": selected_rows, "selected_columns": selected_columns,
            "rows": [{"index": row, "cells": [
                {"index": col, **_text(table.cell(row, col).text, args)}
                for col in range(column_start, column_end)]} for row in range(start, end)]}


def docx_read(ctx, args):
    doc, digest = _load(ctx, args, ".docx", Document)
    start, end, paragraphs = _selection(args, "paragraph", len(doc.paragraphs), 20)
    tstart, tend, tables = _selection(args, "table", len(doc.tables), 2, 20)
    return {"path": args["path"], "sha256": digest, "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables), "selected_paragraphs": paragraphs, "selected_tables": tables,
            "paragraphs": [{"index": i, "style": doc.paragraphs[i].style.name,
                            **_text(doc.paragraphs[i].text, args)} for i in range(start, end)],
            "tables": [_table_read(doc.tables[i], args, i) for i in range(tstart, tend)],
            "warnings": ["Only top-level body paragraphs and tables are returned; headers, footers, text boxes and tracked changes are not included."]}


def _append_blocks(doc, blocks):
    for block in blocks:
        kind = block["type"]
        if kind in ("paragraph", "heading"):
            _check("text" in block and "rows" not in block, "Text blocks require text and cannot contain rows")
            if kind == "heading":
                level = _bounded(block.get("level", 1), 0, 9, "heading level")
                doc.add_heading(block["text"], level=level)
            else:
                _check("level" not in block, "Only heading blocks accept level")
                doc.add_paragraph(block["text"])
        elif kind == "table":
            _check("rows" in block and "text" not in block and "level" not in block,
                   "Table blocks require rows only")
            rows = block["rows"]
            _check(bool(rows) and bool(rows[0]), "Tables need at least one row and column")
            width = len(rows[0])
            _check(all(len(row) == width for row in rows), "Table rows must have equal length")
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    table.cell(r, c).text = text
        else:
            raise ToolError("invalid_argument", "Unknown block type")


def _docx_signature(doc):
    return ([p.text for p in doc.paragraphs],
            [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables])


def _publish_docx(ctx, output_path, doc):
    expected = _docx_signature(doc)

    def verify(path):
        if _docx_signature(Document(path)) != expected:
            raise ToolError("verification_failed", "Saved Word content differs from expected content")

    return {**ctx.publish(output_path, ".docx", doc.save, verify),
            "paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)}


def docx_create(ctx, args):
    doc = Document()
    _append_blocks(doc, args["blocks"])
    return _publish_docx(ctx, args["output_path"], doc)


def docx_update(ctx, args):
    doc, digest = _load(ctx, args, ".docx", Document)
    _check(any(args.get(k) for k in ("paragraph_updates", "cell_updates", "append_blocks")),
           "At least one edit or appended block is required")
    for update in args.get("paragraph_updates", []):
        index = _bounded(update["index"], 0, len(doc.paragraphs) - 1, "paragraph index")
        doc.paragraphs[index].text = update["text"]
    for update in args.get("cell_updates", []):
        index = _bounded(update["table"], 0, len(doc.tables) - 1, "table index")
        table = doc.tables[index]
        row = _bounded(update["row"], 0, len(table.rows) - 1, "table row")
        column = _bounded(update["column"], 0, len(table.columns) - 1, "table column")
        table.cell(row, column).text = update["text"]
    _append_blocks(doc, args.get("append_blocks", []))
    return {**_publish_docx(ctx, args["output_path"], doc), "source_sha256": digest,
            "warnings": ["Replacing paragraph/cell text replaces its runs and inline formatting. Original file is retained; complex Word layout is outside the supported scope."]}


def _xlsx_load(stream):
    return load_workbook(stream, data_only=False, keep_links=False)


def _sheet_name(name):
    _check(bool(name) and len(name) <= 31 and not re.search(r"[\\/*?:\[\]\x00-\x1f]", name)
           and not name.startswith("'") and not name.endswith("'"), "Invalid Excel sheet name")


def _cell_coordinate(value):
    _check(bool(re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]{0,6}", value)), "Expected a single A1 cell coordinate")
    row, col = coordinate_to_tuple(value)
    _check(row <= 1048576 and col <= 16384, "Cell is outside Excel worksheet limits")
    return value.upper()


def _set_cells(sheet, cells):
    seen = set()
    for item in cells:
        coord = _cell_coordinate(item["coordinate"])
        _check(coord not in seen, f"Duplicate cell coordinate: {coord}")
        seen.add(coord)
        kind = item["type"]
        expected = {"coordinate", "type"} | ({kind} if kind != "empty" else set())
        _check(set(item) == expected, f"{kind} cells require exactly these keys: {', '.join(sorted(expected))}")
        cell = sheet[coord]
        if kind == "empty":
            cell.value = None
        elif kind == "text":
            _check(len(item["text"]) <= 32767, "Excel text cells cannot exceed 32767 characters")
            cell.value = item["text"]
            cell.data_type = "s"  # Prevent '=' text from becoming an executable formula.
        elif kind == "formula":
            _check(item["formula"].startswith("=") and len(item["formula"]) > 1,
                   "Formula must start with '=' followed by an expression")
            cell.value = item["formula"]
        elif kind == "number":
            _check(type(item["number"]) in (int, float) and math.isfinite(item["number"]),
                   "Cell number must be finite")
            cell.value = item["number"]
        elif kind == "boolean":
            _check(type(item["boolean"]) is bool, "Boolean cell requires a JSON boolean")
            cell.value = item["boolean"]
        else:
            raise ToolError("invalid_argument", "Unsupported cell type")


def _read_cell(cell, args):
    result = {"coordinate": cell.coordinate}
    value = cell.value
    if value is None:
        if cell.data_type == "inlineStr":
            return {**result, "type": "text", **_text("", args)}
        return {**result, "type": "empty"}
    if cell.data_type == "f":
        _check(isinstance(value, str), "Array/data-table formulas are not supported by this basic reader")
        selection = _text(value, args)
        return {**result, "type": "formula", "formula": selection.pop("text"), **selection}
    if cell.data_type == "e":
        return {**result, "type": "error", **_text(str(value), args)}
    if isinstance(value, bool):
        return {**result, "type": "boolean", "boolean": value}
    if isinstance(value, (int, float)):
        return {**result, "type": "number", "number": value}
    if isinstance(value, (datetime, date, time)):
        return {**result, "type": "date", "iso8601": value.isoformat()}
    if isinstance(value, timedelta):
        return {**result, "type": "duration", "seconds": value.total_seconds()}
    return {**result, "type": "text", **_text(str(value), args)}


def xlsx_read(ctx, args):
    workbook, digest = _load(ctx, args, ".xlsx", _xlsx_load)
    try:
        start, end, selected_sheets = _selection(args, "sheet", len(workbook.worksheets), 20, 100)
        result = {"path": args["path"], "sha256": digest, "sheet_count": len(workbook.worksheets),
                  "selected_sheets": selected_sheets,
                  "sheets": [{"name": s.title, "max_row": s.max_row, "max_column": s.max_column}
                             for s in workbook.worksheets[start:end]], "rows": [],
                  "warnings": ["Formulas are returned as expressions; no formulas are evaluated and cached results are not used."]}
        if "sheet" not in args:
            _check("range" not in args, "A range requires sheet")
            return result
        _check(args["sheet"] in workbook.sheetnames, "Sheet does not exist")
        selected = args.get("range", "A1:J20").upper()
        _check(bool(re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}(:[A-Z]{1,3}[1-9][0-9]{0,6})?", selected)),
               "range must be an A1 cell or rectangular A1:B2 range")
        c1, r1, c2, r2 = range_boundaries(selected)
        _check(1 <= r1 <= r2 <= 1048576 and 1 <= c1 <= c2 <= 16384, "Range is outside Excel limits")
        _check((r2 - r1 + 1) * (c2 - c1 + 1) <= 2000, "Range exceeds 2000 cells; request a smaller range")
        sheet = workbook[args["sheet"]]
        result.update({"sheet": sheet.title, "selected_range": selected,
                       "rows": [{"row": r, "cells": [_read_cell(sheet.cell(r, c), args)
                                 for c in range(c1, c2 + 1)]} for r in range(r1, r2 + 1)]})
        return result
    finally:
        workbook.close()


def _xlsx_cells(workbook, restore_empty_strings=False):
    result = {}
    for sheet in workbook.worksheets:
        # The pinned openpyxl version stores only instantiated cells here.
        # iter_rows() would expand an entire rectangle for sparse worksheets.
        for cell in sheet._cells.values():
            value = cell.value
            if value is None and cell.data_type in ("s", "inlineStr"):
                value = ""
                if restore_empty_strings:
                    cell.value = ""
                    cell.data_type = "s"
            if value is None or not isinstance(value, (str, int, float, datetime, date, time, timedelta)):
                continue
            data_type = "s" if cell.data_type == "inlineStr" else cell.data_type
            result[(sheet.title, cell.coordinate)] = (value, data_type)
    return result


def _publish_xlsx(ctx, output_path, workbook):
    names = workbook.sheetnames[:]
    expected_cells = _xlsx_cells(workbook, restore_empty_strings=True)

    def verify(path):
        saved = load_workbook(path, data_only=False, keep_links=False)
        try:
            if saved.sheetnames != names:
                raise ToolError("verification_failed", "Saved Excel sheet names differ")
            if _xlsx_cells(saved) != expected_cells:
                raise ToolError("verification_failed", "Saved Excel basic cell values or types differ; no output was published")
        finally:
            saved.close()

    return {**ctx.publish(output_path, ".xlsx", workbook.save, verify), "sheet_count": len(names),
            "warnings": ["Formulas are saved but not calculated or refreshed. Updates may remove external workbook links and complex Excel objects; the source file is retained."]}


def xlsx_create(ctx, args):
    _check(bool(args["sheets"]), "At least one sheet is required")
    workbook = Workbook()
    workbook.remove(workbook.active)
    seen = set()
    try:
        for item in args["sheets"]:
            name = item["name"]
            _sheet_name(name)
            _check(name.casefold() not in seen, "Sheet names must be unique ignoring case")
            seen.add(name.casefold())
            sheet = workbook.create_sheet(name)
            _set_cells(sheet, item["cells"])
        return _publish_xlsx(ctx, args["output_path"], workbook)
    finally:
        workbook.close()


def xlsx_update(ctx, args):
    workbook, digest = _load(ctx, args, ".xlsx", _xlsx_load)
    try:
        _check(any(update["cells"] for update in args["updates"]), "At least one cell update is required")
        seen = set()
        for update in args["updates"]:
            name = update["sheet"]
            _check(name in workbook.sheetnames, "Sheet does not exist")
            _check(name not in seen, "Only one update group per sheet is allowed")
            seen.add(name)
            _set_cells(workbook[name], update["cells"])
        return {**_publish_xlsx(ctx, args["output_path"], workbook), "source_sha256": digest}
    finally:
        workbook.close()


def _shapes(shapes, prefix=""):
    for i, shape in enumerate(shapes):
        index = f"{prefix}{i}"
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes, f"{index}.")
        else:
            yield index, shape


def pptx_read(ctx, args):
    presentation, digest = _load(ctx, args, ".pptx", Presentation)
    start, end, selected = _selection(args, "slide", len(presentation.slides), 5, 80)
    slides = []
    for index in range(start, end):
        slide = presentation.slides[index]
        texts, tables = [], []
        shapes = list(_shapes(slide.shapes))
        shape_start, shape_end, selected_shapes = _selection(args, "shape", len(shapes), 20, 100)
        for shape_index, shape in shapes[shape_start:shape_end]:
            if shape.has_text_frame:
                texts.append({"shape_index": shape_index, **_text(shape.text, args)})
            if shape.has_table:
                tables.append(_table_read(shape.table, args, shape_index))
        slides.append({"index": index, "title": _text(slide.shapes.title.text if slide.shapes.title else "", args),
                       "shape_count": len(shapes), "selected_shapes": selected_shapes,
                       "texts": texts, "tables": tables})
    return {"path": args["path"], "sha256": digest, "slide_count": len(presentation.slides),
            "selected_slides": selected, "slides": slides,
            "warnings": ["Reads slide text and tables, including groups. Notes, charts, images and rendered layout are not extracted."]}


def _line_count(text, units_per_line):
    return sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(c) in "WF" else 1
                                  for c in line) / units_per_line)) for line in text.splitlines() or [""])


def pptx_create(ctx, args):
    _check(bool(args["slides"]), "At least one slide is required")
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333333), Inches(7.5)
    for item in args["slides"]:
        body = item.get("body", "")
        bullets = item.get("bullets", [])
        _check(_line_count(item["title"], 48) <= 2, "Slide title is too long; shorten the title")
        lines = ([body] if body else []) + [f"• {b}" for b in bullets]
        # Conservative width estimates and explicit line/paragraph spacing keep
        # CJK and many short bullets inside the basic layout's vertical budget.
        _check(sum(_line_count(text, 64) * 32 + 10 for text in lines) <= 360,
               "Slide content exceeds the readable layout; split it into multiple slides")
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        title = slide.shapes.title
        title.left, title.top, title.width, title.height = Inches(.65), Inches(.4), Inches(12), Inches(1.2)
        title.text_frame.word_wrap = True
        title.text = item["title"]
        for paragraph in title.text_frame.paragraphs:
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(32)
            paragraph.font.bold = True
            paragraph.line_spacing = Pt(38)
            paragraph.font.color.rgb = RGBColor(31, 55, 88)
        content = slide.shapes.add_textbox(Inches(.7), Inches(1.65), Inches(11.9), Inches(5.35))
        frame = content.text_frame
        frame.word_wrap = True
        for i, text in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = text
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(24)
            paragraph.line_spacing = Pt(32)
            paragraph.font.color.rgb = RGBColor(43, 50, 58)
            paragraph.space_after = Pt(10)

    expected = [[s.text for s in slide.shapes if s.has_text_frame] for slide in presentation.slides]

    def verify(path):
        saved = Presentation(path)
        actual = [[s.text for s in slide.shapes if s.has_text_frame] for slide in saved.slides]
        if actual != expected:
            raise ToolError("verification_failed", "Saved PowerPoint content differs")

    return {**ctx.publish(args["output_path"], ".pptx", presentation.save, verify),
            "slide_count": len(presentation.slides),
            "warnings": ["Basic editable 16:9 slides only; rendering and font availability depend on the presentation viewer."]}


_PATH = string("Path relative to the configured plugin cwd; absolute paths and traversal are forbidden.", max_length=4096)
_TEXT_OPTIONS = {"text_offset": integer("0-based character offset applied to each text field; default 0."),
                 "max_chars": integer("1..32768 characters per text field; default 4000. text_length and has_more expose remaining text.")}
_TABLE_OPTIONS = {f"{name}_{suffix}": integer(f"0-based {name} start or count for each selected table; default start 0, count {default}.")
                  for name, default in (("row", 20), ("column", 10)) for suffix in ("start", "count")}
_BLOCK = obj({"type": {"type": "string", "enum": ["paragraph", "heading", "table"]},
              "text": string(), "level": integer("Heading level 0..9; default 1."),
              "rows": arr(arr(string(max_length=4000), max_items=50), max_items=200)}, ["type"])
_CELL = obj({"coordinate": string("Single A1 coordinate.", max_length=16),
             "type": {"type": "string", "enum": ["text", "number", "boolean", "formula", "empty"]},
             "text": string(max_length=32767), "number": number(), "boolean": boolean(),
             "formula": string("Excel expression starting with '='; no recalculation.", max_length=8192)},
            ["coordinate", "type"])
_SOURCE = {"path": _PATH, "expected_sha256": string("SHA-256 from a prior read; stale sources are rejected.", max_length=64),
           "output_path": _PATH}


def _spec(name, description, properties, required, read_only):
    return {"name": name, "title": name.replace("_", " ").title(),
            "description": description + " Paths are bound to this provider's configured cwd. " +
            ("Read-only; use range controls and text_offset/max_chars to page content. Output overflow is an error: narrow the selection."
             if read_only else "Writes a new output_path only, never overwrites. Its parent directory must exist. Saves and verifies before publishing. On OutcomeUnknown inspect the target before retrying; do not replay automatically."),
            "inputSchema": obj(properties, required),
            "annotations": {"readOnlyHint": read_only, "destructiveHint": False,
                            "idempotentHint": read_only, "openWorldHint": False}}


TOOL_SPECS = [
    _spec("docx_read", "Read DOCX top-level paragraphs, headings (style) and tables, including SHA-256 and total counts. Defaults: first 20 paragraphs and 2 tables; 0-based selections.",
          {"path": _PATH, **_TEXT_OPTIONS, **_TABLE_OPTIONS,
           **{f"{n}_{s}": integer() for n in ("paragraph", "table") for s in ("start", "count")}}, ["path"], True),
    _spec("docx_create", "Create a basic DOCX from ordered paragraph, heading and table blocks.",
          {"output_path": _PATH, "blocks": arr(_BLOCK, max_items=200)}, ["output_path", "blocks"], False),
    _spec("docx_update", "Update DOCX into a new copy using 0-based paragraph/table/cell indexes or append blocks. Replacing text loses its inline run formatting; original is preserved.",
          {**_SOURCE, "paragraph_updates": arr(obj({"index": integer(), "text": string()}, ["index", "text"]), max_items=200),
           "cell_updates": arr(obj({"table": integer(), "row": integer(), "column": integer(), "text": string()},
                                   ["table", "row", "column", "text"]), max_items=500),
           "append_blocks": arr(_BLOCK, max_items=200)}, ["path", "expected_sha256", "output_path"], False),
    _spec("xlsx_read", "Read XLSX sheet metadata and optional A1 range (default A1:J20; max 2000 cells). Sheet metadata uses 0-based sheet_start/sheet_count, default first 20, max 100. Omit sheet to list sheets only. Formula expressions are read without calculation or cached values.",
          {"path": _PATH, "sheet": string(max_length=31), "range": string(max_length=40),
           "sheet_start": integer("0-based metadata page start; default 0."),
           "sheet_count": integer("Metadata entries per call, 0..100; default 20."), **_TEXT_OPTIONS}, ["path"], True),
    _spec("xlsx_create", "Create an XLSX with named sheets and explicitly typed cells. Each cell requires exactly the value key named by type, or no value key for empty. Text starting '=' remains text; formulas are not calculated.",
          {"output_path": _PATH, "sheets": arr(obj({"name": string(max_length=31), "cells": arr(_CELL, max_items=2000)}, ["name", "cells"]), max_items=50)},
          ["output_path", "sheets"], False),
    _spec("xlsx_update", "Write cell changes into a new XLSX copy. Each cell uses an explicit type and matching value key. Original remains unchanged. External workbook links and complex Excel objects may be lost; formulas are saved without recalculation.",
          {**_SOURCE, "updates": arr(obj({"sheet": string(max_length=31), "cells": arr(_CELL, max_items=2000)}, ["sheet", "cells"]), max_items=50)},
          ["path", "expected_sha256", "output_path", "updates"], False),
    _spec("pptx_read", "Read PPTX slide count, titles, text and tables including grouped shapes. 0-based slide_start/slide_count, default first 5 slides. Each slide uses a flattened shape_start/shape_count window (default first 20, max 100) with selected_shapes continuation. Does not read notes, images or charts.",
          {"path": _PATH, "slide_start": integer(), "slide_count": integer(),
           "shape_start": integer("0-based flattened shape start in each selected slide; default 0. Use one slide for continuation."),
           "shape_count": integer("Flattened shapes per selected slide, 0..100; default 20."),
           **_TEXT_OPTIONS, **_TABLE_OPTIONS}, ["path"], True),
    _spec("pptx_create", "Create basic editable 16:9 PPTX slides with title, optional body and bullets. Content exceeding the readable 24pt body layout is rejected; split large content across slides.",
          {"output_path": _PATH, "slides": arr(obj({"title": string(max_length=140), "body": string(max_length=1000),
                                                   "bullets": arr(string(max_length=300), max_items=10)}, ["title"]), max_items=80)},
          ["output_path", "slides"], False),
]

HANDLERS = {name: globals()[name] for name in ("docx_read", "docx_create", "docx_update", "xlsx_read",
                                            "xlsx_create", "xlsx_update", "pptx_read", "pptx_create")}
