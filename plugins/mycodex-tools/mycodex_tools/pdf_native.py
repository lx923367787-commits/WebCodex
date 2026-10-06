"""Bounded PDF/image backend. Paths and write transactions belong to the caller.

PDF coordinates are points with a bottom-left origin; image coordinates are pixels
with a top-left origin. Only destinations supplied by the caller are written.
"""
from __future__ import annotations

import csv
from contextlib import closing
import ctypes
import functools
import hashlib
import io
import math
from pathlib import Path
import subprocess
import uuid

from PIL import Image, ImageOps
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import Highlight, Text
from pypdf.generic import ArrayObject, DecodedStreamObject, DictionaryObject, FloatObject, NameObject, RectangleObject, TextStringObject

from .common import ToolError

MAX_BYTES = 64 * 1024 * 1024
MAX_PIXELS = 16_000_000
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp"}


def _api(function):
    @functools.wraps(function)
    def wrapped(*args, **kwargs):
        try:
            return function(*args, **kwargs)
        except ToolError:
            raise
        except (ImportError, ModuleNotFoundError) as exc:
            raise ToolError("missing_dependency", f"PDF backend dependency unavailable: {exc.name}.") from exc
        except Exception as exc:
            raise ToolError("pdf_operation_failed", f"PDF operation failed ({type(exc).__name__}); source was not modified.") from exc
    return wrapped


def _int(options, key, default, minimum, maximum):
    value = options.get(key, default)
    if type(value) is not int or not minimum <= value <= maximum:
        raise ToolError("invalid_argument", f"{key} must be an integer between {minimum} and {maximum}.")
    return value


def _source(path):
    path = Path(path)
    if path.suffix.lower() not in IMAGE_SUFFIXES | {".pdf"}:
        raise ToolError("unsupported_format", "Use PDF, PNG, JPEG or WEBP.")
    if not path.is_file():
        raise ToolError("file_not_found", "Source file does not exist.")
    if path.stat().st_size > MAX_BYTES:
        raise ToolError("file_too_large", "PDF/image source exceeds 64 MiB.")
    return path


def _reader(data):
    reader = PdfReader(io.BytesIO(data), strict=False)
    if reader.is_encrypted:
        raise ToolError("unsupported_pdf", "Encrypted PDFs must be decrypted before processing.")
    if not 1 <= len(reader.pages) <= 1000:
        raise ToolError("page_limit", "PDF must contain between 1 and 1000 pages.")
    return reader


def _unsigned(reader):
    if reader.trailer["/Root"].get("/Perms") or any(
            field.get("/FT") == "/Sig" and field.get("/V") for field in (reader.get_fields() or {}).values()):
        raise ToolError("signed_pdf", "This PDF has a digital signature; editing would invalidate it. Use an unsigned copy.")


def _selection(total, options, maximum=20, default=2):
    start = _int(options, "page_start", 1, 1, total)
    count = _int(options, "page_count", min(default, total - start + 1), 1, maximum)
    if start + count - 1 > total:
        raise ToolError("invalid_range", "Selected page range exceeds the document.")
    return range(start - 1, start + count - 1)


def _image(path):
    with Image.open(path) as opened:
        if opened.width * opened.height > MAX_PIXELS:
            raise ToolError("image_too_large", "Image exceeds 16 million pixels; crop or resize it first.")
        return ImageOps.exif_transpose(opened).convert("RGB")


def _bitmap(page, options):
    dpi = _int(options, "dpi", 150, 72, 300)
    width, height = page.get_size()
    if width <= 0 or height <= 0 or math.ceil(width * dpi / 72) * math.ceil(height * dpi / 72) > MAX_PIXELS:
        raise ToolError("image_too_large", "Rendered page exceeds 16 million pixels; lower dpi.")
    bitmap = page.render(scale=dpi / 72)
    try:
        return bitmap.to_pil().convert("RGB")
    finally:
        bitmap.close()


def _ocr(image, options):
    executable = options.get("tesseract_path")
    if not executable or not Path(executable).is_file():
        raise ToolError("ocr_unavailable", "Configure a local Tesseract executable and chi_sim/eng language data.")
    language = options.get("language", options.get("ocr_language", "chi_sim+eng"))
    if language not in ("chi_sim+eng", "eng+chi_sim", "chi_sim", "eng"):
        raise ToolError("invalid_argument", "language must contain chi_sim and/or eng.")
    timeout = _int(options, "ocr_timeout_seconds", 45, 1, 120)
    command = [str(executable), "stdin", "stdout", "-l", language,
               "--psm", str(_int(options, "ocr_psm", 3, 3, 13))]
    if options.get("tessdata_path"):
        command.extend(["--tessdata-dir", str(options["tessdata_path"])])
    command.append("tsv")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    try:
        completed = subprocess.run(command, input=buffer.getvalue(), capture_output=True,
                                   timeout=timeout, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except subprocess.TimeoutExpired as exc:
        raise ToolError("ocr_timeout", "OCR exceeded its time limit; select a smaller page or crop.") from exc
    if completed.returncode:
        # Do not echo executable diagnostics, which may contain private paths.
        raise ToolError("ocr_failed", "Tesseract failed; verify chi_sim/eng data and the selected image.")
    if len(completed.stdout) > 4 * 1024 * 1024:
        raise ToolError("ocr_result_too_large", "OCR result exceeds 4 MiB; crop the image.")
    words, lines = [], {}
    threshold = _int(options, "ocr_confidence", 60, 0, 100)
    for row in csv.DictReader(io.StringIO(completed.stdout.decode("utf-8-sig")), delimiter="\t", quoting=csv.QUOTE_NONE):
        if row.get("level") != "5" or not row.get("text", "").strip():
            continue
        confidence = float(row["conf"])
        if confidence < 0:
            continue
        if len(words) >= 5000:
            raise ToolError("ocr_result_too_large", "More than 5000 OCR words on one page; crop the image.")
        x, y, w, h = (int(row[key]) for key in ("left", "top", "width", "height"))
        if min(x, y, w, h) < 0 or x + w > image.width or y + h > image.height or not math.isfinite(confidence):
            raise ToolError("invalid_ocr_result", "OCR returned invalid coordinates or confidence.")
        word = {"text": row["text"], "rect": [x, y, x + w, y + h],
                "confidence": round(confidence, 2), "low_confidence": confidence < threshold, "source": "ocr"}
        words.append(word)
        key = tuple(row.get(field) for field in ("block_num", "par_num", "line_num"))
        lines.setdefault(key, []).append(row["text"])
    return {"text": "\n".join(" ".join(line) for line in lines.values()), "blocks": words,
            "status": "ocr_text" if words else "no_text_detected",
            "confidence_mean": round(sum(w["confidence"] for w in words) / len(words), 2) if words else None,
            "low_confidence_count": sum(w["low_confidence"] for w in words)}


def _to_page_point(page, image, x, y):
    from pypdfium2 import raw
    px, py = ctypes.c_double(), ctypes.c_double()
    if not raw.FPDF_DeviceToPage(page, 0, 0, image.width, image.height, 0, x, y,
                                 ctypes.byref(px), ctypes.byref(py)):
        raise ToolError("invalid_pdf", "Cannot map OCR coordinates to PDF page coordinates.")
    return px.value, py.value


def _to_page_rect(page, image, rect):
    points = [_to_page_point(page, image, x, y) for x, y in ((rect[0], rect[1]), (rect[2], rect[3]))]
    return [min(p[0] for p in points), min(p[1] for p in points),
            max(p[0] for p in points), max(p[1] for p in points)]


def _native_page(page):
    from pypdfium2 import raw
    blocks = []
    with closing(page.get_textpage()) as textpage:
        text = textpage.get_text_range().replace("\r\n", "\n")
        for index, item in enumerate(page.get_objects(max_depth=1)):
            if index >= 20000:
                raise ToolError("page_too_complex", "Page contains more than 20000 objects.")
            if item.type == raw.FPDF_PAGEOBJ_TEXT:
                item.textpage = textpage
                value = item.extract()
                if value:
                    blocks.append({"object_index": index, "text": value,
                                   "rect": list(item.get_bounds()), "source": "native", "confidence": None})
        return text, blocks


def _window(page, options):
    text = page["text"]
    offset = _int(options, "text_offset", 0, 0, len(text))
    count = _int(options, "max_chars", 12000, 1, 32768)
    end = min(len(text), offset + count)
    page.update(text=text[offset:end], text_offset=offset, chars_total=len(text),
                next_offset=end if end < len(text) else None)
    blocks = page["blocks"]
    start = _int(options, "block_offset", 0, 0, len(blocks))
    limit = _int(options, "max_blocks", 80, 1, 100)
    # Bound duplicate text in object descriptors as well as the main text window.
    selected, budget = [], 12000
    for block in blocks[start:start + limit]:
        if len(block["text"]) > budget:
            if not selected:
                selected.append({**block, "text": block["text"][:budget], "text_truncated": True})
            break
        selected.append(block)
        budget -= len(block["text"])
    page.update(blocks=selected, block_offset=start, blocks_total=len(blocks),
                next_block_offset=start + len(selected) if start + len(selected) < len(blocks) else None)
    return page


def _continuation(result, options):
    safe = {key: value for key, value in options.items() if key in
            {"ocr", "language", "ocr_language", "ocr_confidence", "dpi", "max_chars", "max_blocks"}}
    result["next_options"] = None
    for page in result["pages"]:
        if page["next_offset"] is not None or page["next_block_offset"] is not None:
            result["next_options"] = {**safe, "page_start": page["page"], "page_count": 1,
                                      "text_offset": page["next_offset"] if page["next_offset"] is not None else page["chars_total"],
                                      "block_offset": page["next_block_offset"] if page["next_block_offset"] is not None else page["blocks_total"]}
            return result
    if result["next_page"] is not None:
        result["next_options"] = {**safe, "page_start": result["next_page"], "page_count": 1}
    return result


@_api
def inspect(source: Path, options: dict) -> dict:
    source = _source(source)
    mode = options.get("ocr", "auto")
    if mode not in ("auto", "always", "never"):
        raise ToolError("invalid_argument", "ocr must be auto, always or never.")
    result = {"path": str(source.resolve()), "sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "pages": []}
    if source.suffix.lower() in IMAGE_SUFFIXES:
        _selection(1, options)
        bitmap = _image(source)
        page = _ocr(bitmap, options) if mode != "never" else {"text": "", "blocks": [], "status": "ocr_not_requested"}
        page.update(page=1, width=bitmap.width, height=bitmap.height, coordinate_system="pixels_top_left")
        result.update(page_count_total=1, next_page=None)
        result["pages"].append(_window(page, options))
        return _continuation(result, options)
    import pypdfium2 as pdfium
    data = source.read_bytes()
    reader = _reader(data)
    fields = reader.get_fields() or {}
    field_start = _int(options, "field_offset", 0, 0, len(fields))
    field_count = _int(options, "max_fields", 20, 1, 40)
    result["fields"] = [{"name": name, "type": str(field.get("/FT", "")),
                         "value": str(field.get("/V", ""))[:4096],
                         "states": [str(value) for value in field.get("/_States_", [])][:30]}
                        for name, field in list(fields.items())[field_start:field_start + field_count]]
    result["fields_total"] = len(fields)
    result["next_field_offset"] = field_start + field_count if field_start + field_count < len(fields) else None
    with pdfium.PdfDocument(data) as document:
        document.init_forms()
        selected = _selection(len(document), options)
        result.update(page_count_total=len(document), next_page=selected.stop + 1 if selected.stop < len(document) else None)
        for number in selected:
            with closing(document[number]) as page:
                text, blocks = _native_page(page)
                info = {"text": text, "blocks": blocks, "status": "native_text" if text.strip() else "no_extractable_text"}
                if mode == "always" or (mode == "auto" and not text.strip()):
                    bitmap = _bitmap(page, options)
                    info = _ocr(bitmap, options)
                    for word in info["blocks"]:
                        word["rect"] = _to_page_rect(page, bitmap, word["rect"])
                info.update(page=number + 1, width=page.get_width(), height=page.get_height(), coordinate_system="pdf_points_bottom_left")
                result["pages"].append(_window(info, options))
    return _continuation(result, options)


def _encoded_image(bitmap, maximum):
    output = io.BytesIO()
    bitmap.save(output, format="PNG", optimize=True)
    if len(output.getvalue()) <= maximum:
        return output.getvalue(), ".png", "image/png"
    for quality in (90, 80, 65, 50):
        output = io.BytesIO()
        bitmap.save(output, format="JPEG", quality=quality, optimize=True)
        if len(output.getvalue()) <= maximum:
            return output.getvalue(), ".jpg", "image/jpeg"
    raise ToolError("image_result_too_large", "Page image exceeds 1 MiB; lower dpi or select a crop.")


@_api
def render(source: Path, output_dir: Path, options: dict) -> list[dict]:
    source = _source(source)
    directory = Path(output_dir).resolve(strict=True)
    if not directory.is_dir():
        raise ToolError("invalid_path", "Render output directory must already exist.")
    maximum = _int(options, "max_image_bytes", 1024 * 1024, 16384, 1024 * 1024)
    result, owned = [], []

    def save(bitmap, number):
        crop = options.get("crop")
        if crop is not None:
            if not isinstance(crop, list) or len(crop) != 4 or any(type(v) is not int for v in crop):
                raise ToolError("invalid_argument", "crop must be integer pixels [left,top,right,bottom].")
            if not (0 <= crop[0] < crop[2] <= bitmap.width and 0 <= crop[1] < crop[3] <= bitmap.height):
                raise ToolError("invalid_range", "Crop falls outside the rendered image.")
            bitmap = bitmap.crop(crop)
        content, suffix, mime = _encoded_image(bitmap, maximum)
        target = directory / f"page-{number}-{uuid.uuid4().hex}{suffix}"
        with target.open("xb") as stream:
            stream.write(content)
        owned.append(target)
        result.append({"path": str(target), "page": number, "mime_type": mime,
                       "width": bitmap.width, "height": bitmap.height, "bytes": len(content), "crop": crop,
                       "render_source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                       "renderer": "Pillow" if source.suffix.lower() in IMAGE_SUFFIXES else "PDFium"})

    try:
        if source.suffix.lower() in IMAGE_SUFFIXES:
            _selection(1, options, maximum=2)
            save(_image(source), 1)
        else:
            import pypdfium2 as pdfium
            data = source.read_bytes()
            _reader(data)
            with pdfium.PdfDocument(data) as document:
                document.init_forms()
                for index in _selection(len(document), options, maximum=2):
                    with closing(document[index]) as page:
                        save(_bitmap(page, options), index + 1)
        return result
    except Exception:
        for path in owned:
            path.unlink(missing_ok=True)
        raise


def _destination(source, destination, suffix):
    destination = Path(destination)
    if destination.resolve() == Path(source).resolve():
        raise ToolError("unsafe_destination", "Backend destination must differ from source; caller owns the write transaction.")
    if destination.suffix.lower() != suffix or not destination.parent.is_dir():
        raise ToolError("invalid_path", f"Destination must have suffix {suffix} and an existing parent.")
    return destination


def _serialize(writer):
    output = io.BytesIO()
    writer.write(output)
    content = output.getvalue()
    if len(content) > MAX_BYTES:
        raise ToolError("file_too_large", "Generated PDF exceeds 64 MiB.")
    _reader(content)
    return content


def _page_numbers(value, total):
    if not isinstance(value, list) or not value or len(value) > 1000 or any(type(v) is not int or not 1 <= v <= total for v in value):
        raise ToolError("invalid_argument", "pages must contain valid 1-based page numbers.")
    if len(set(value)) != len(value):
        raise ToolError("invalid_argument", "Duplicate page numbers are not supported.")
    return value


def _rect(value):
    if not isinstance(value, list) or len(value) != 4 or any(type(v) not in (int, float) or not math.isfinite(v) for v in value):
        raise ToolError("invalid_argument", "rect must contain four finite PDF coordinates.")
    if value[0] >= value[2] or value[1] >= value[3]:
        raise ToolError("invalid_argument", "rect must have positive width and height.")
    return value


def _utf16(text):
    content = text.encode("utf-16-le") + b"\0\0"
    return (ctypes.c_ushort * (len(content) // 2)).from_buffer_copy(content)


def _object_bounds(item):
    from pypdfium2 import raw
    coordinates = [ctypes.c_float() for _ in range(4)]
    if not raw.FPDFPageObj_GetBounds(item, *(ctypes.byref(v) for v in coordinates)):
        raise ToolError("unsafe_pdf_replacement", "Cannot determine replacement text bounds.")
    return tuple(v.value for v in coordinates)


def _font_data(path, text):
    from reportlab.pdfbase.ttfonts import TTFont
    if not path or not Path(path).is_file():
        raise ToolError("missing_font", "Replacement requires a configured Chinese TrueType font covering the new text.")
    if Path(path).stat().st_size > 32 * 1024 * 1024:
        raise ToolError("font_too_large", "Configured TrueType font exceeds 32 MiB.")
    font = TTFont("PDFNativeCoverage", str(path))
    missing = next((char for char in text if not char.isspace() and not font.face.charToGlyph.get(ord(char))), None)
    if missing is not None:
        raise ToolError("unsupported_font_glyph", f"Configured font lacks U+{ord(missing):04X}; use another font or convert to DOCX.")
    return Path(path).read_bytes()


def _replace(data, operation, options):
    import pypdfium2 as pdfium
    from pypdfium2 import raw
    old, new = operation.get("old_text"), operation.get("new_text")
    if not isinstance(old, str) or not old or not isinstance(new, str) or not new or max(len(old), len(new)) > 4096 or any(c in new for c in "\0\r\n"):
        raise ToolError("invalid_argument", "replace_text needs nonempty old_text/new_text (one line, at most 4096 characters).")
    with pdfium.PdfDocument(data) as document:
        page_number = _int(operation, "page", 1, 1, len(document))
        with closing(document[page_number - 1]) as page, closing(page.get_textpage()) as textpage:
            objects = list(page.get_objects(max_depth=1))
            if len(objects) > 20000:
                raise ToolError("page_too_complex", "Page contains too many objects for safe replacement.")
            matches = []
            for index, item in enumerate(objects):
                if item.type == raw.FPDF_PAGEOBJ_TEXT:
                    item.textpage = textpage
                    if item.extract() == old and ("object_index" not in operation or operation["object_index"] == index):
                        matches.append(item)
            if len(matches) != 1:
                raise ToolError("unsafe_pdf_replacement", "Text must match exactly one complete top-level text object; inspect object_index or convert to DOCX.")
            item = matches[0]
            # PDFium requires text-page handles to be closed before removing a
            # text object. Keeping one alive here risks stale native pointers.
            textpage.close()
            matrix = item.get_matrix()
            if abs(matrix.b) > 0.0001 or abs(matrix.c) > 0.0001 or matrix.a <= 0 or matrix.d <= 0:
                raise ToolError("unsafe_pdf_replacement", "Rotated, skewed or vertical text must be converted to DOCX before editing.")
            clip = raw.FPDFPageObj_GetClipPath(item)
            if clip and raw.FPDFClipPath_CountPaths(clip) > 0:
                raise ToolError("unsafe_pdf_replacement", "Clipped text cannot be safely replaced; convert to DOCX.")
            if raw.FPDFTextObj_GetTextRenderMode(item) != raw.FPDF_TEXTRENDERMODE_FILL or raw.FPDFPageObj_HasTransparency(item):
                raise ToolError("unsafe_pdf_replacement", "Only opaque fill text objects support local replacement.")
            if raw.FPDFPageObj_CountMarks(item):
                raise ToolError("unsafe_pdf_replacement", "Tagged/marked text may contain alternate text; convert to DOCX before editing.")
            original_bounds = item.get_bounds()
            font_name = item.get_font().get_base_name()
            # Existing subset fonts reliably encode only already used characters.
            replace_font = not set(new).issubset(set(old)) and not (
                font_name in {"Helvetica", "Times-Roman", "Courier", "Helvetica-Bold", "Times-Bold", "Courier-Bold"}
                and all(32 <= ord(char) <= 126 for char in new))
            loaded_font = None
            replacement = item
            if replace_font:
                font_bytes = _font_data(options.get("font_path"), new)
                font_buffer = (ctypes.c_ubyte * len(font_bytes)).from_buffer_copy(font_bytes)
                loaded_font = raw.FPDFText_LoadFont(document, font_buffer, len(font_bytes), raw.FPDF_FONT_TRUETYPE, True)
                if not loaded_font:
                    raise ToolError("unsupported_font", "PDFium cannot load configured TrueType font; convert to DOCX.")
                pointer = raw.FPDFPageObj_CreateTextObj(document, loaded_font, item.get_font_size())
                if not pointer:
                    raw.FPDFFont_Close(loaded_font)
                    raise ToolError("unsupported_font", "Cannot create replacement text with the configured font.")
                replacement = pdfium.PdfTextObj(pointer, pdf=document)
                replacement.set_matrix(matrix)
                channels = [ctypes.c_uint() for _ in range(4)]
                if not raw.FPDFPageObj_GetFillColor(item, *(ctypes.byref(v) for v in channels)) or not raw.FPDFPageObj_SetFillColor(replacement, *(v.value for v in channels)):
                    replacement.close()
                    raw.FPDFFont_Close(loaded_font)
                    raise ToolError("unsafe_pdf_replacement", "Cannot preserve original text color.")
            try:
                if not raw.FPDFText_SetText(replacement, _utf16(new)):
                    raise ToolError("unsafe_pdf_replacement", "PDFium cannot encode the replacement text; convert to DOCX.")
                bounds = _object_bounds(replacement)
                if bounds[0] < original_bounds[0] - 2 or bounds[2] > original_bounds[2] + 2 or bounds[1] < original_bounds[1] - 2 or bounds[3] > original_bounds[3] + 2:
                    raise ToolError("pdf_text_overflow", "Replacement does not fit the original text area; shorten it or convert to DOCX.")
                if replace_font:
                    # PDFium may place a newly embedded-font object in a new
                    # content stream; refuse overlap rather than change layering.
                    for other in objects:
                        if other is item:
                            continue
                        a, b, c, d = other.get_bounds()
                        if min(c, bounds[2]) > max(a, bounds[0]) and min(d, bounds[3]) > max(b, bounds[1]):
                            raise ToolError("unsafe_pdf_replacement", "Font substitution overlaps another object; convert to DOCX.")
                    page.remove_obj(item)
                    item.close()
                    page.insert_obj(replacement)
                page.gen_content()
                output = io.BytesIO()
                document.save(output)
            finally:
                if replace_font and replacement.page is None:
                    replacement.close()
                if loaded_font:
                    raw.FPDFFont_Close(loaded_font)
    result = output.getvalue()
    with pdfium.PdfDocument(result) as verification:
        with closing(verification[page_number - 1]) as page:
            extracted, objects_after = _native_page(page)
            if not any(block["text"] == new for block in objects_after):
                raise ToolError("verification_failed", "Replacement text did not survive PDF save/readback.")
            if old != new and sum(b["text"] == old for b in objects_after) >= sum(
                    b["text"] == old for b in inspect_bytes_page(data, page_number)):
                raise ToolError("verification_failed", "Original text object was not removed.")
            _bitmap(page, {"dpi": 72})
    return result, {"page": page_number, "font_substituted": replace_font, "replacement_verified": True}


def inspect_bytes_page(data, page_number):
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(data) as document, closing(document[page_number - 1]) as page:
        return _native_page(page)[1]


def _unicode_form_appearances(writer, values, options):
    """Embed a covering font in plain Unicode text-field appearances.

    A Unicode /V alone is insufficient: a viewer can otherwise show blank boxes
    while string readback reports success. Preserve widget geometry/background.
    """
    unicode_values = {name: value for name, value in values.items() if isinstance(value, str) and any(ord(c) > 127 for c in value)}
    if not unicode_values:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    _font_data(options.get("font_path"), "".join(unicode_values.values()))
    font_name = "PDFNativeForm_" + hashlib.sha256(str(options["font_path"]).encode()).hexdigest()[:12]
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(options["font_path"])))
    handled = set()
    for page in writer.pages:
        for reference in page.get("/Annots", []):
            widget = reference.get_object()
            names, field_type, parent, levels, value_owner = [], None, widget, 0, None
            while parent is not None:
                levels += 1
                if levels > 32:
                    raise ToolError("unsupported_form_field", "Form field nesting exceeds 32.")
                if parent.get("/T"):
                    names.append(str(parent["/T"]))
                    value_owner = value_owner or parent
                field_type = field_type or parent.get("/FT")
                parent = parent.get("/Parent")
                parent = parent.get_object() if parent else None
            name = ".".join(reversed(names))
            if name not in unicode_values:
                continue
            value = unicode_values[name]
            if field_type != "/Tx" or any(c in value for c in "\r\n") or widget.get("/MK", {}).get("/R", 0):
                raise ToolError("unsupported_form_field", "Unicode form filling supports unrotated, single-line text fields; use an editable document for this field.")
            x0, y0, x1, y1 = map(float, widget["/Rect"])
            width, height = x1 - x0, y1 - y0
            if width <= 4 or height <= 4:
                raise ToolError("unsupported_form_field", "Form field is too small to render text.")
            size = min(12, height * 0.65)
            measured = pdfmetrics.stringWidth(value, font_name, size)
            if measured > width - 4:
                size *= (width - 4) / measured
            if size < 6:
                raise ToolError("pdf_text_overflow", "Unicode form text would be smaller than 6 points; shorten it.")
            output = io.BytesIO()
            drawing = canvas.Canvas(output, pagesize=(width, height), pageCompression=1)
            appearance = widget.get("/MK", {})
            background, border = appearance.get("/BG"), appearance.get("/BC")
            for color, fill in ((background, True), (border, False)):
                if color and len(color) == 3:
                    (drawing.setFillColorRGB if fill else drawing.setStrokeColorRGB)(*map(float, color))
                    drawing.rect(0.5, 0.5, width - 1, height - 1, stroke=not fill, fill=fill)
            drawing.setFillColorRGB(0, 0, 0)
            drawing.setFont(font_name, size)
            drawing.drawString(2, max(2, (height - size) / 2), value)
            drawing.showPage()
            drawing.save()
            appearance_page = PdfReader(io.BytesIO(output.getvalue())).pages[0]
            stream = DecodedStreamObject()
            stream.set_data(appearance_page.get_contents().get_data())
            stream.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Form"),
                           NameObject("/BBox"): RectangleObject([0, 0, width, height]),
                           NameObject("/Resources"): appearance_page["/Resources"].clone(writer)})
            widget[NameObject("/AP")] = DictionaryObject({NameObject("/N"): writer._add_object(stream)})
            value_owner[NameObject("/V")] = TextStringObject(value)
            handled.add(name)
    if handled != set(unicode_values):
        raise ToolError("unsupported_form_field", "Cannot locate every Unicode field widget for appearance generation.")


@_api
def edit(source: Path, destination: Path, operations: list[dict], options: dict) -> dict:
    source = _source(source)
    if source.suffix.lower() != ".pdf":
        raise ToolError("unsupported_format", "PDF editing requires PDF input; convert images first.")
    target = _destination(source, destination, ".pdf")
    if not isinstance(operations, list) or not 1 <= len(operations) <= 100 or any(not isinstance(op, dict) for op in operations):
        raise ToolError("invalid_argument", "operations must contain 1 to 100 objects.")
    data = source.read_bytes()
    initial = _reader(data)
    _unsigned(initial)
    results = []
    for operation in operations:
        kind = operation.get("op")
        if kind == "replace_text":
            data, result = _replace(data, operation, options)
            results.append({"op": kind, **result})
            continue
        reader = _reader(data)
        writer = PdfWriter(clone_from=reader)
        total = len(writer.pages)
        if kind in ("delete_pages", "select_pages", "reorder_pages"):
            pages = _page_numbers(operation.get("pages"), total)
            if kind == "reorder_pages" and set(pages) != set(range(1, total + 1)):
                raise ToolError("invalid_argument", "reorder_pages requires a complete permutation of the pages.")
            if kind == "delete_pages":
                pages = [n for n in range(1, total + 1) if n not in pages]
            if not pages:
                raise ToolError("invalid_argument", "Cannot delete every page.")
            # Filter fields/bookmarks along with pages. Deleting then re-adding
            # pages to a cloned catalog leaves removed-page AcroForm values and
            # outline destinations reachable through the document catalog.
            writer = PdfWriter()
            writer.append(reader, pages=[number - 1 for number in pages])
            if reader.metadata:
                writer.add_metadata({str(key): str(value) for key, value in reader.metadata.items()})
        elif kind == "rotate_pages":
            angle = operation.get("angle")
            if type(angle) is not int or angle not in (-270, -180, -90, 90, 180, 270):
                raise ToolError("invalid_argument", "angle must be +/-90, +/-180 or +/-270.")
            for number in _page_numbers(operation.get("pages"), total):
                writer.pages[number - 1].rotate(angle)
        elif kind == "merge":
            paths = operation.get("paths")
            if not isinstance(paths, list) or not 1 <= len(paths) <= 20:
                raise ToolError("invalid_argument", "merge.paths must contain 1 to 20 caller-validated source paths.")
            for path in paths:
                other = _source(Path(path))
                if other.suffix.lower() != ".pdf":
                    raise ToolError("unsupported_format", "Only PDF sources can be merged.")
                additional = _reader(other.read_bytes())
                _unsigned(additional)
                writer.append(additional)
                if len(writer.pages) > 1000:
                    raise ToolError("page_limit", "Merged PDF exceeds 1000 pages.")
        elif kind == "fill_form":
            fields = operation.get("fields")
            existing = reader.get_fields() or {}
            if not isinstance(fields, dict) or not fields or len(fields) > 100 or set(fields) - set(existing):
                raise ToolError("invalid_argument", "fields must name existing PDF form fields (at most 100).")
            if any(not isinstance(v, (str, bool)) or (isinstance(v, str) and len(v) > 4096) for v in fields.values()):
                raise ToolError("invalid_argument", "Form values must be strings or booleans.")
            normalized = {}
            for name, value in fields.items():
                field = existing[name]
                if field.get("/FT") == "/Sig":
                    raise ToolError("unsupported_form_field", "Signature fields cannot be filled.")
                if field.get("/FT") == "/Btn":
                    states = [str(v) for v in field.get("/_States_", []) if str(v) != "/Off"]
                    if isinstance(value, bool):
                        if value and len(states) != 1:
                            raise ToolError("unsupported_form_field", "Button on-state is ambiguous; provide its exact state name.")
                        value = states[0] if value else "/Off"
                    if value not in ["/Off", *states]:
                        raise ToolError("invalid_argument", "Invalid checkbox or radio state.")
                elif not isinstance(value, str):
                    raise ToolError("invalid_argument", "Text/choice fields require string values.")
                normalized[name] = value
            native_values = {name: value for name, value in normalized.items()
                             if not isinstance(value, str) or all(ord(char) < 128 for char in value)}
            if native_values:
                writer.update_page_form_field_values(None, native_values, auto_regenerate=False)
            writer.set_need_appearances_writer(False)
            _unicode_form_appearances(writer, normalized, options)
        elif kind == "annotate":
            page = _int(operation, "page", 1, 1, total)
            rect = _rect(operation.get("rect"))
            text = operation.get("text", "")
            if not isinstance(text, str) or len(text) > 4096:
                raise ToolError("invalid_argument", "Annotation text must be at most 4096 characters.")
            if operation.get("kind", "text") == "text":
                annotation = Text(rect=rect, text=text, open=False)
            elif operation["kind"] == "highlight":
                x0, y0, x1, y1 = rect
                annotation = Highlight(rect=rect, quad_points=ArrayObject([FloatObject(v) for v in (x0, y1, x1, y1, x0, y0, x1, y0)]))
            else:
                raise ToolError("invalid_argument", "Annotation kind must be text or highlight.")
            writer.add_annotation(page - 1, annotation)
        else:
            raise ToolError("unsupported_operation", f"Unsupported PDF operation: {kind}.")
        data = _serialize(writer)
        if kind == "fill_form":
            verified = _reader(data).get_fields() or {}
            if any(str(verified.get(name, {}).get("/V", "")) != str(value) for name, value in normalized.items()):
                raise ToolError("verification_failed", "PDF form values differ after save/readback.")
        results.append({"op": kind, "verified": True})
    if len(data) > MAX_BYTES:
        raise ToolError("file_too_large", "Generated PDF exceeds 64 MiB.")
    verified_reader = _reader(data)
    target.write_bytes(data)
    return {"path": str(target.resolve()), "page_count": len(verified_reader.pages), "operations": results,
            "verified": True, "sha256": hashlib.sha256(data).hexdigest()}


def _ocr_pages(source, options):
    if source.suffix.lower() in IMAGE_SUFFIXES:
        bitmap = _image(source)
        result = _ocr(bitmap, options)
        result.update(page=1, width=bitmap.width * 72 / 150, height=bitmap.height * 72 / 150,
                      bitmap=bitmap, coordinate_system="pixels_top_left")
        yield result
        return
    import pypdfium2 as pdfium
    data = source.read_bytes()
    _reader(data)
    with pdfium.PdfDocument(data) as document:
        # Conversion is a whole-document operation; never silently drop pages.
        if len(document) > _int(options, "max_conversion_pages", 50, 1, 100):
            raise ToolError("page_limit", "Conversion exceeds max_conversion_pages; split the PDF first.")
        for index in range(len(document)):
            with closing(document[index]) as page:
                native, _ = _native_page(page)
                if native.strip():
                    yield {"page": index + 1, "text": native, "status": "native_text", "blocks": []}
                else:
                    bitmap = _bitmap(page, options)
                    result = _ocr(bitmap, options)
                    for word in result["blocks"]:
                        rect = word["rect"]
                        left_bottom = _to_page_point(page, bitmap, rect[0], rect[3])
                        right_bottom = _to_page_point(page, bitmap, rect[2], rect[3])
                        left_top = _to_page_point(page, bitmap, rect[0], rect[1])
                        dx, dy = right_bottom[0] - left_bottom[0], right_bottom[1] - left_bottom[1]
                        word.update(origin=left_bottom, angle=math.atan2(dy, dx),
                                    width_points=math.hypot(dx, dy),
                                    height_points=math.dist(left_top, left_bottom))
                        word["rect"] = _to_page_rect(page, bitmap, rect)
                    result.update(page=index + 1, width=page.get_width(), height=page.get_height(), coordinate_system="pdf_points_bottom_left")
                    yield result


@_api
def convert(source: Path, destination: Path, options: dict) -> dict:
    source = _source(source)
    suffix = Path(destination).suffix.lower()
    if suffix not in (".pdf", ".docx"):
        raise ToolError("unsupported_format", "Convert scans/images to searchable .pdf or editable .docx.")
    target = _destination(source, destination, suffix)
    if suffix == ".pdf" and source.suffix.lower() == ".pdf":
        _unsigned(_reader(source.read_bytes()))
    if "page_start" in options or "page_count" in options:
        raise ToolError("invalid_argument", "Conversion processes the whole document; use select_pages to split first.")
    pages = list(_ocr_pages(source, options))
    if suffix == ".docx":
        if any(page["status"] == "native_text" for page in pages) and not options.get("allow_text_reflow", False):
            raise ToolError("office_conversion_required", "Native/mixed PDFs require Office PDF-to-DOCX conversion; allow_text_reflow=true explicitly accepts plain-text reflow.")
        from docx import Document
        document = Document()
        for index, page in enumerate(pages):
            if index:
                document.add_page_break()
            for line in page["text"].splitlines():
                document.add_paragraph(line)
        output = io.BytesIO()
        document.save(output)
        data = output.getvalue()
        loaded = Document(io.BytesIO(data))
        actual = "\n".join(p.text for p in loaded.paragraphs)
        if any(page["text"] and page["text"] not in actual for page in pages):
            raise ToolError("verification_failed", "DOCX OCR text did not survive save/readback.")
        limitations = ["OCR text is editable; original columns, tables, images and typography are not reconstructed."]
    else:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas
        from reportlab.lib.utils import ImageReader
        writer = PdfWriter(clone_from=_reader(source.read_bytes())) if source.suffix.lower() == ".pdf" else PdfWriter()
        words = [word for page in pages for word in page["blocks"]]
        font_name = None
        if words:
            font_path = options.get("font_path")
            _font_data(font_path, "".join(word["text"] for word in words))
            font_name = "PDFNativeOCR_" + hashlib.sha256(str(font_path).encode()).hexdigest()[:12]
            if font_name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
        for page in pages:
            if page["status"] == "native_text":
                continue
            buffer = io.BytesIO()
            original_page = None if "bitmap" in page else writer.pages[page["page"] - 1]
            page_size = (page["width"], page["height"]) if original_page is None else (
                float(original_page.mediabox.width), float(original_page.mediabox.height))
            layer = canvas.Canvas(buffer, pagesize=page_size, pageCompression=1)
            if "bitmap" in page:
                layer.drawImage(ImageReader(page["bitmap"]), 0, 0, width=page["width"], height=page["height"])
            for word in page["blocks"]:
                x0, y0, x1, y1 = word["rect"]
                if page["coordinate_system"] == "pixels_top_left":
                    scale = 72 / 150
                    x0, x1, y0, y1 = x0 * scale, x1 * scale, page["height"] - y1 * scale, page["height"] - y0 * scale
                size = max(1.0, word.get("height_points", y1 - y0))
                text = layer.beginText()
                angle = word.get("angle", 0)
                origin = word.get("origin", (x0, y0))
                text.setTextTransform(math.cos(angle), math.sin(angle), -math.sin(angle), math.cos(angle), *origin)
                text.setFont(font_name, size)
                text.setTextRenderMode(3)
                width = pdfmetrics.stringWidth(word["text"], font_name, size)
                if width > 0:
                    text.setHorizScale(word.get("width_points", x1 - x0) * 100 / width)
                text.textOut(word["text"])
                layer.drawText(text)
            layer.showPage()
            layer.save()
            overlay = PdfReader(io.BytesIO(buffer.getvalue())).pages[0]
            if "bitmap" in page:
                writer.add_page(overlay)
            else:
                overlay.mediabox = original_page.mediabox
                overlay.cropbox = original_page.cropbox
                original_page.merge_page(overlay)
        data = _serialize(writer)
        verification = _reader(data)
        for page in pages:
            extracted = verification.pages[page["page"] - 1].extract_text() or ""
            if any(word["text"] not in extracted for word in page["blocks"]):
                raise ToolError("verification_failed", "Searchable PDF OCR text did not survive save/readback.")
        limitations = ["OCR is an invisible searchable text layer; low-confidence words remain marked in the result."]
    if len(data) > MAX_BYTES:
        raise ToolError("file_too_large", "Converted file exceeds 64 MiB.")
    target.write_bytes(data)
    return {"path": str(target.resolve()), "page_count": len(pages), "verified": True,
            "sha256": hashlib.sha256(data).hexdigest(), "limitations": limitations,
            "pages": [{key: page.get(key) for key in ("page", "status", "confidence_mean", "low_confidence_count")} for page in pages]}
