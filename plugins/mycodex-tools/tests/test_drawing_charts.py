"""Focused real-rendering and read-only data-boundary tests for office charts."""
from __future__ import annotations

from copy import deepcopy
import csv
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook
from PIL import Image
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from drawing.charts import render_chart
from drawing.common import DrawingError

FONT = Path("C:/Windows/Fonts/msyh.ttc")
S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


class DrawingChartTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.mpl = tempfile.TemporaryDirectory(prefix="mycodex-chart-font-cache-")
        cls.old_mpl = os.environ.get("MPLCONFIGDIR")
        os.environ["MPLCONFIGDIR"] = cls.mpl.name
        if not FONT.is_file():
            raise RuntimeError("The Chinese rendering fixture requires the installed Microsoft YaHei font.")

    @classmethod
    def tearDownClass(cls):
        if cls.old_mpl is None:
            os.environ.pop("MPLCONFIGDIR", None)
        else:
            os.environ["MPLCONFIGDIR"] = cls.old_mpl
        cls.mpl.cleanup()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="mycodex-chart-test-")
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.root = self.base / "project"
        self.root.mkdir()
        self.context = {"project_root": self.root, "font_path": FONT}
        self.counter = 0

    def workbook(self, rows, name="source.xlsx"):
        book = Workbook()
        book.active.title = "Sheet1"
        for row in rows:
            book.active.append(row)
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        book.save(path)
        return path

    def spec(self, **changes):
        value = {"kind": "chart", "chart_type": "bar", "title": "月度经营统计",
                 "source": {"path": "source.xlsx", "sheet": "Sheet1", "range": "A1:C4", "header": True},
                 "x": "月份", "series": [{"column": "收入"}, {"column": "成本"}], "unit": "万元"}
        value.update(changes)
        return value

    def draw(self, spec):
        self.counter += 1
        output = self.root / f"out-{self.counter}"
        output.mkdir()
        return render_chart(spec, output, self.context)

    def error(self, code, spec):
        self.counter += 1
        output = self.root / f"bad-{self.counter}"
        output.mkdir()
        with self.assertRaises(DrawingError) as raised:
            render_chart(spec, output, self.context)
        self.assertEqual(raised.exception.code, code)
        self.assertEqual(list(output.iterdir()), [])

    def set_cache(self, path, coordinate, value, error=False):
        with zipfile.ZipFile(path) as archive:
            parts = {item.filename: archive.read(item) for item in archive.infolist()}
        root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        cell = root.find(f'.//{S}c[@r="{coordinate}"]')
        cached = cell.find(S + "v")
        if cached is None:
            cached = ET.SubElement(cell, S + "v")
        cached.text = str(value)
        if error:
            cell.set("t", "e")
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in parts.items():
                archive.writestr(name, data)

    def check_outputs(self, result):
        self.assertEqual(set(result["outputs"]), {"png", "svg", "pdf", "data_csv"})
        png = Path(result["outputs"]["png"])
        with Image.open(png) as image:
            style = result["metadata"]["style"]
            self.assertEqual(image.size, (int(style["width"]*style["dpi"]), int(style["height"]*style["dpi"])))
            self.assertGreater(len(image.getcolors(image.width*image.height)), 30)
        self.assertIn(b"<svg", Path(result["outputs"]["svg"]).read_bytes())
        self.assertEqual(len(PdfReader(result["outputs"]["pdf"]).pages), 1)
        self.assertFalse(result["metadata"]["source_modified"])

    def test_real_bar_line_pie_scatter_preserve_source_and_record_references(self):
        path = self.workbook([["月份", "收入", "成本"], ["一月", 12, 8], ["二月", 15, 10], ["三月", 19, 11]])
        before = path.read_bytes()
        for kind in ("bar", "line", "pie", "scatter"):
            with self.subTest(kind=kind):
                spec = self.spec(chart_type=kind)
                if kind == "pie":
                    spec["series"] = [{"column": "收入"}]
                if kind == "scatter":
                    spec["x"] = "收入"
                    spec["series"] = [{"column": "成本"}]
                    spec["x_label"], spec["y_label"] = "收入", "成本"
                result = self.draw(spec)
                self.check_outputs(result)
                self.assertEqual(result["metadata"]["source_sha256"], hashlib.sha256(before).hexdigest())
                self.assertEqual(result["metadata"]["point_rows"], [2, 3, 4])
                self.assertEqual(result["metadata"]["selected_columns"]["series"][0]["column"], "C" if kind == "scatter" else "B")
        self.assertEqual(path.read_bytes(), before)

    def test_actual_four_shipped_templates_render_against_fixture(self):
        self.workbook([["月份", "收入", "成本"], ["一月", 12, 8], ["二月", 15, 10], ["三月", 19, 11],
                       ["四月", 16, 9], ["五月", 21, 12], ["六月", 23, 14]], "drawings/data.xlsx")
        templates = Path(__file__).resolve().parents[1] / "drawing" / "templates"
        for kind in ("bar", "line", "pie", "scatter"):
            with self.subTest(template=kind):
                result = self.draw(json.loads((templates / f"{kind}.json").read_text("utf-8")))
                self.check_outputs(result)
                self.assertEqual(result["metadata"]["data_rows"], 6)

    def test_negative_values_and_pie_negative_zero_fail(self):
        self.workbook([["月份", "收入", "成本"], ["一月", -12, 8], ["二月", 15, -10], ["三月", 19, 11]])
        result = self.draw(self.spec())
        with Path(result["outputs"]["data_csv"]).open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[1][2], "-12")
        self.error("negative_pie_value", self.spec(chart_type="pie", series=[{"column": "收入"}]))
        self.workbook([["月份", "收入", "成本"], ["一月", 0, 8], ["二月", 0, 10], ["三月", 0, 11]])
        self.error("empty_pie", self.spec(chart_type="pie", series=[{"column": "收入"}]))

    def test_missing_skip_is_explicit_but_errors_or_non_numbers_never_skip(self):
        self.workbook([["月份", "收入", "成本"], ["一月", 12, 8], ["二月", None, 10], ["三月", 19, 11]])
        self.error("missing_data", self.spec())
        result = self.draw(self.spec(missing="skip"))
        self.assertEqual(result["metadata"]["point_rows"], [2, 4])
        self.assertEqual(result["metadata"]["skipped_rows"], [3])
        self.assertEqual(result["warnings"][0]["code"], "rows_skipped")
        self.workbook([["月份", "收入", "成本"], ["一月", 12, 8], ["二月", "not-a-number", 10], ["三月", 19, 11]])
        self.error("numeric_text_requires_opt_in", self.spec(missing="skip"))
        self.error("non_numeric_data", self.spec(missing="skip", numeric_strings="decimal"))
        self.workbook([["月份", "收入", "成本"], ["一月", True, 8], ["二月", 15, 10], ["三月", 19, 11]])
        self.error("non_numeric_data", self.spec(missing="skip"))

    def test_missing_skip_does_not_hide_another_bad_selected_value(self):
        self.workbook([["月份", "收入", "成本"], ["一月", 12, 8], ["二月", "bad", None], ["三月", 19, 11]])
        self.error("non_numeric_data", self.spec(missing="skip", numeric_strings="decimal"))
        self.workbook([["月份", "收入", "成本"], ["一月", 12, 8], ["二月", True, None], ["三月", 19, 11]])
        self.error("non_numeric_data", self.spec(missing="skip"))
        self.workbook([["月份", "收入", "成本"], ["一月", 12, 8], [True, 15, None], ["三月", 19, 11]])
        self.error("invalid_category", self.spec(missing="skip"))

    def test_formulas_use_saved_cache_and_missing_or_error_cache_fails(self):
        path = self.workbook([["月份", "收入", "成本"], ["一月", "=2*6", 8], ["二月", 15, 10], ["三月", 19, 11]])
        self.error("formula_cache_missing", self.spec(missing="skip"))
        self.set_cache(path, "B2", 12)
        result = self.draw(self.spec())
        self.assertEqual(result["metadata"]["formula_cache"]["count"], 1)
        self.assertEqual(result["metadata"]["formula_cache"]["sample_cells"], ["B2"])
        self.assertEqual(result["warnings"][0]["code"], "saved_formula_values")
        self.set_cache(path, "B2", "#DIV/0!", error=True)
        self.error("spreadsheet_error", self.spec(missing="skip"))

    def test_csv_decimal_opt_in_is_strict_and_export_preserves_precision(self):
        path = self.root / "values.csv"
        path.write_text("类别,数值\n一月,1.12345678901234567890123456789\n二月,-2.500\n", encoding="utf-8-sig")
        spec = self.spec(source={"path": "values.csv", "range": "A1:B3"}, x="类别", series=[{"column": "数值"}])
        self.error("numeric_text_requires_opt_in", spec)
        result = self.draw({**spec, "numeric_strings": "decimal"})
        csv_text = Path(result["outputs"]["data_csv"]).read_text("utf-8-sig")
        self.assertIn("1.12345678901234567890123456789", csv_text)
        self.assertIn("-2.500", csv_text)
        path.write_text('类别,数值\n一月,"1,000"\n二月,2\n', encoding="utf-8-sig")
        self.error("non_numeric_data", {**spec, "numeric_strings": "decimal"})

    def test_archive_csv_replays_without_source_and_preserves_literal_labels(self):
        path = self.workbook([["月份", "收入", "成本"], ["=literal", 12, 8], ["'already quoted", 15, 10], ["三月", 19, 11]])
        # In the fixture, force the leading '=' category to be literal XLSX text.
        from openpyxl import load_workbook
        book = load_workbook(path)
        book.active["A2"].data_type = "s"
        book.save(path)
        result = self.draw(self.spec())
        snapshot = deepcopy(result["snapshot_spec"])
        snapshot["source"]["path"] = Path(result["outputs"]["data_csv"]).relative_to(self.root).as_posix()
        path.unlink()
        replay = self.draw(snapshot)
        self.assertEqual(replay["metadata"]["data_rows"], 3)
        with Path(replay["outputs"]["data_csv"]).open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[1][1], "'=literal")
        self.assertEqual(rows[2][1], "''already quoted")

    def test_date_line_snapshot_keeps_date_axis(self):
        self.workbook([["月份", "收入", "成本"], [datetime(2026, 1, 1), 12, 8],
                       [datetime(2026, 1, 15), 15, 10], [datetime(2026, 3, 1), 19, 11]])
        result = self.draw(self.spec(chart_type="line"))
        self.assertEqual(result["metadata"]["x_scale"], "date")
        snapshot = deepcopy(result["snapshot_spec"])
        snapshot["source"]["path"] = Path(result["outputs"]["data_csv"]).relative_to(self.root).as_posix()
        replay = self.draw(snapshot)
        self.assertEqual(replay["metadata"]["x_scale"], "date")

    def test_paths_and_limits_fail_without_output(self):
        self.workbook([["月份", "收入", "成本"], ["一月", 12, 8], ["二月", 15, 10], ["三月", 19, 11]])
        for path in ("../source.xlsx", "C:/outside.xlsx", ".git/data.xlsx", "secrets/data.xlsx"):
            with self.subTest(path=path):
                self.error("invalid_path", self.spec(source={"path": path, "range": "A1:C4"}))
        self.error("chart_range_too_large", self.spec(source={"path": "source.xlsx", "range": "A1:C2002"}))
        self.error("invalid_chart_spec", self.spec(series=[{"column": 1, "label": str(i)} for i in range(9)]))
        self.error("column_not_found", self.spec(x="不存在"))

    def test_symlink_source_is_rejected(self):
        path = self.workbook([["月份", "收入", "成本"], ["一月", 12, 8], ["二月", 15, 10], ["三月", 19, 11]])
        link = self.root / "linked.xlsx"
        try:
            link.symlink_to(path)
        except OSError:
            self.skipTest("The host does not allow synthetic symlink fixtures")
        self.error("invalid_path", self.spec(source={"path": "linked.xlsx", "range": "A1:C4"}))


if __name__ == "__main__":
    unittest.main()
