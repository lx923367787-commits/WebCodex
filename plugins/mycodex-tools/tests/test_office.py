"""Real Office round trips and externally meaningful format guarantees."""

import hashlib
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

import plugin
from mycodex_tools.common import Context, ToolError, validate_args
from mycodex_tools.office import HANDLERS, TOOL_SPECS


class OfficeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.ctx = Context(self.root)

    def call(self, name, **args):
        spec = next(item for item in TOOL_SPECS if item["name"] == name)
        validate_args(args, spec["inputSchema"])
        return HANDLERS[name](self.ctx, args)

    def digest(self, path):
        return hashlib.sha256((self.root / path).read_bytes()).hexdigest()

    def test_word_roundtrip_updates_preserve_source(self):
        created = self.call("docx_create", output_path="报告.docx", blocks=[
            {"type": "heading", "text": "阶段汇报", "level": 1},
            {"type": "paragraph", "text": "项目进展顺利"},
            {"type": "table", "rows": [["指标", "结果"], ["任务", "完成"]]},
        ])
        self.assertTrue(created["verified"])
        before = self.call("docx_read", path="报告.docx")
        self.assertEqual(before["paragraph_count"], 2)
        self.assertEqual(before["paragraphs"][0]["text"], "阶段汇报")
        self.assertEqual(before["paragraphs"][0]["style"], "Heading 1")
        self.call("docx_update", path="报告.docx", expected_sha256=before["sha256"],
                  output_path="报告-修订.docx", paragraph_updates=[{"index": 1, "text": "已经验收"}],
                  cell_updates=[{"table": 0, "row": 1, "column": 1, "text": "通过"}],
                  append_blocks=[{"type": "paragraph", "text": "后续计划"}])
        after = self.call("docx_read", path="报告-修订.docx")
        self.assertEqual(after["paragraphs"][1]["text"], "已经验收")
        self.assertEqual(after["paragraphs"][2]["text"], "后续计划")
        self.assertEqual(after["tables"][0]["rows"][1]["cells"][1]["text"], "通过")
        self.assertEqual(self.digest("报告.docx"), before["sha256"])

    def test_word_ranges_and_long_text_have_explicit_continuation(self):
        self.call("docx_create", output_path="long.docx", blocks=[
            {"type": "paragraph", "text": "abcdefghij"}, {"type": "paragraph", "text": "第二段"}])
        result = self.call("docx_read", path="long.docx", paragraph_count=1, text_offset=2, max_chars=3)
        self.assertEqual(result["paragraphs"][0]["text"], "cde")
        self.assertEqual(result["paragraphs"][0]["text_length"], 10)
        self.assertTrue(result["paragraphs"][0]["has_more"])
        self.assertTrue(result["selected_paragraphs"]["has_more"])
        with self.assertRaises(ToolError):
            self.call("docx_read", path="long.docx", paragraph_start=-1)

    def test_updates_reject_stale_hash_and_existing_target(self):
        self.call("docx_create", output_path="original.docx", blocks=[{"type": "paragraph", "text": "a"}])
        digest = self.digest("original.docx")
        with self.assertRaises(ToolError):
            self.call("docx_update", path="original.docx", expected_sha256="0" * 64,
                      output_path="bad.docx", paragraph_updates=[{"index": 0, "text": "b"}])
        self.assertFalse((self.root / "bad.docx").exists())
        with self.assertRaises(ToolError):
            self.call("docx_update", path="original.docx", expected_sha256=digest,
                      output_path="original.docx", paragraph_updates=[{"index": 0, "text": "b"}])
        self.assertEqual(digest, self.digest("original.docx"))

    def test_excel_types_formula_and_copy_update(self):
        self.call("xlsx_create", output_path="数据.xlsx", sheets=[{"name": "数据", "cells": [
            {"coordinate": "A1", "type": "text", "text": "=SUM(B1:B2)"},
            {"coordinate": "B1", "type": "number", "number": 12.5},
            {"coordinate": "C1", "type": "formula", "formula": "=B1*2"},
            {"coordinate": "D1", "type": "boolean", "boolean": True},
            {"coordinate": "E1", "type": "empty"},
        ]}, {"name": "备注", "cells": [{"coordinate": "A1", "type": "text", "text": "原样保留"}]}])
        before = self.call("xlsx_read", path="数据.xlsx", sheet="数据", range="A1:E1")
        self.assertEqual(before["sheet_count"], 2)
        cells = before["rows"][0]["cells"]
        self.assertEqual([c["type"] for c in cells], ["text", "number", "formula", "boolean", "empty"])
        self.assertEqual(cells[0]["text"], "=SUM(B1:B2)")
        self.assertEqual(cells[2]["formula"], "=B1*2")
        self.call("xlsx_update", path="数据.xlsx", expected_sha256=before["sha256"], output_path="数据-修订.xlsx",
                  updates=[{"sheet": "数据", "cells": [{"coordinate": "B1", "type": "number", "number": 30}]}])
        after = self.call("xlsx_read", path="数据-修订.xlsx", sheet="数据", range="A1:D1")
        self.assertEqual(after["rows"][0]["cells"][1]["number"], 30)
        self.assertEqual(after["rows"][0]["cells"][2]["formula"], "=B1*2")
        untouched = self.call("xlsx_read", path="数据-修订.xlsx", sheet="备注", range="A1")
        self.assertEqual(untouched["rows"][0]["cells"][0]["text"], "原样保留")
        self.assertEqual(self.digest("数据.xlsx"), before["sha256"])
        workbook = load_workbook(self.root / "数据-修订.xlsx", data_only=False)
        try:
            self.assertEqual(workbook["数据"]["A1"].data_type, "s")
            self.assertEqual(workbook["数据"]["C1"].data_type, "f")
        finally:
            workbook.close()

    def test_excel_rejects_ambiguous_cells_duplicate_names_and_huge_ranges(self):
        bad_cells = [
            {"coordinate": "A1", "type": "text", "number": 5},
            {"coordinate": "A1", "type": "formula", "formula": "SUM(A1:A2)"},
            {"coordinate": "XFE1", "type": "empty"},
        ]
        for item in bad_cells:
            with self.subTest(item=item), self.assertRaises(ToolError):
                self.call("xlsx_create", output_path="bad.xlsx", sheets=[{"name": "Sheet", "cells": [item]}])
        with self.assertRaises(ToolError):
            self.call("xlsx_create", output_path="bad.xlsx", sheets=[{"name": "Sheet", "cells": []}, {"name": "sheet", "cells": []}])
        self.call("xlsx_create", output_path="ok.xlsx", sheets=[{"name": "Sheet", "cells": []}])
        with self.assertRaises(ToolError):
            self.call("xlsx_read", path="ok.xlsx", sheet="Sheet", range="A1:Z1000")
        self.assertFalse((self.root / "bad.xlsx").exists())

    def test_excel_empty_text_remains_distinct_and_precision_loss_is_rejected(self):
        self.call("xlsx_create", output_path="empty.xlsx", sheets=[{"name": "Sheet", "cells": [
            {"coordinate": "A1", "type": "text", "text": ""},
            {"coordinate": "B1", "type": "empty"},
            {"coordinate": "C1", "type": "number", "number": 0},
            {"coordinate": "D1", "type": "boolean", "boolean": False},
        ]}])
        cells = self.call("xlsx_read", path="empty.xlsx", sheet="Sheet", range="A1:D1")["rows"][0]["cells"]
        self.assertEqual([cell["type"] for cell in cells], ["text", "empty", "number", "boolean"])
        self.assertEqual(cells[0]["text"], "")
        with self.assertRaises(ToolError) as raised:
            self.call("xlsx_create", output_path="inexact.xlsx", sheets=[{"name": "Sheet", "cells": [
                {"coordinate": "A1", "type": "number", "number": 1234567890123456789}]}])
        self.assertEqual(raised.exception.code, "verification_failed")
        self.assertFalse((self.root / "inexact.xlsx").exists())

    def test_excel_update_preserves_unedited_empty_text_and_sparse_cells(self):
        self.call("xlsx_create", output_path="source.xlsx", sheets=[{"name": "Sheet", "cells": [
            {"coordinate": "A1", "type": "text", "text": ""},
            {"coordinate": "XFD1048576", "type": "text", "text": "稀疏末尾"},
        ]}])
        digest = self.digest("source.xlsx")
        self.call("xlsx_update", path="source.xlsx", expected_sha256=digest, output_path="changed.xlsx",
                  updates=[{"sheet": "Sheet", "cells": [{"coordinate": "B1", "type": "number", "number": 1}]}])
        cell = self.call("xlsx_read", path="changed.xlsx", sheet="Sheet", range="A1")["rows"][0]["cells"][0]
        self.assertEqual((cell["type"], cell["text"]), ("text", ""))
        edge = self.call("xlsx_read", path="changed.xlsx", sheet="Sheet", range="XFD1048576")["rows"][0]["cells"][0]
        self.assertEqual(edge["text"], "稀疏末尾")
        self.assertEqual(self.digest("source.xlsx"), digest)

    def test_excel_update_rejects_precision_loss_in_unedited_cells(self):
        workbook = Workbook()
        workbook.active["A1"] = 1
        workbook.save(self.root / "base.xlsx")
        workbook.close()
        with zipfile.ZipFile(self.root / "base.xlsx") as source:
            with zipfile.ZipFile(self.root / "source.xlsx", "w", zipfile.ZIP_DEFLATED) as destination:
                for entry in source.infolist():
                    data = source.read(entry)
                    if entry.filename == "xl/worksheets/sheet1.xml":
                        data = data.replace(b"<v>1</v>", b"<v>1.2345678901234567</v>")
                    destination.writestr(entry, data)
        digest = self.digest("source.xlsx")
        with self.assertRaises(ToolError) as raised:
            self.call("xlsx_update", path="source.xlsx", expected_sha256=digest, output_path="changed.xlsx",
                      updates=[{"sheet": "Sheet", "cells": [{"coordinate": "B1", "type": "text", "text": "edit"}]}])
        self.assertEqual(raised.exception.code, "verification_failed")
        self.assertFalse((self.root / "changed.xlsx").exists())
        self.assertEqual(self.digest("source.xlsx"), digest)

    def test_many_excel_sheets_are_pageable_and_update_result_stays_bounded(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for index in range(1600):
            workbook.create_sheet("页" * 26 + f"{index:05d}")
        first_name = workbook.sheetnames[0]
        workbook.save(self.root / "many.xlsx")
        workbook.close()
        first = plugin.call_tool(self.ctx, "xlsx_read", {"path": "many.xlsx", "sheet": first_name, "range": "A1"})
        self.assertFalse(first["isError"])
        self.assertEqual(len(first["structuredContent"]["sheets"]), 20)
        self.assertEqual(first["structuredContent"]["selected_sheets"]["total"], 1600)
        self.assertTrue(first["structuredContent"]["selected_sheets"]["has_more"])
        last = self.call("xlsx_read", path="many.xlsx", sheet_start=1590, sheet_count=100)
        self.assertEqual(len(last["sheets"]), 10)
        self.assertFalse(last["selected_sheets"]["has_more"])
        updated = plugin.call_tool(self.ctx, "xlsx_update", {
            "path": "many.xlsx", "expected_sha256": first["structuredContent"]["sha256"],
            "output_path": "changed.xlsx", "updates": [{"sheet": first_name,
                "cells": [{"coordinate": "A1", "type": "number", "number": 1}]}]})
        self.assertFalse(updated["isError"], updated)
        self.assertTrue(updated["structuredContent"]["verified"])
        self.assertEqual(updated["structuredContent"]["sheet_count"], 1600)
        self.assertNotIn("sheets", updated["structuredContent"])
        self.assertTrue((self.root / "changed.xlsx").exists())

    def test_powerpoint_two_chinese_slides_and_read_external_table(self):
        self.call("pptx_create", output_path="汇报.pptx", slides=[
            {"title": "项目汇报", "body": "第一阶段已完成", "bullets": ["质量达标", "下阶段计划"]},
            {"title": "阶段结论", "body": "按期完成", "bullets": []},
        ])
        result = self.call("pptx_read", path="汇报.pptx")
        self.assertEqual(result["slide_count"], 2)
        self.assertEqual(result["slides"][0]["title"]["text"], "项目汇报")
        self.assertIn("第一阶段已完成", result["slides"][0]["texts"][1]["text"])
        self.assertEqual(result["slides"][1]["title"]["text"], "阶段结论")
        presentation = Presentation(self.root / "汇报.pptx")
        self.assertAlmostEqual(presentation.slide_width / presentation.slide_height, 16 / 9, places=5)
        table = presentation.slides[0].shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
        table.cell(0, 0).text = "指标"
        table.cell(1, 1).text = "完成"
        presentation.save(self.root / "含表格.pptx")
        table_result = self.call("pptx_read", path="含表格.pptx", slide_count=1)
        self.assertEqual(table_result["slides"][0]["tables"][0]["rows"][1]["cells"][1]["text"], "完成")
        self.assertTrue(table_result["selected_slides"]["has_more"])

    def test_powerpoint_overfull_content_fails_before_publication(self):
        with self.assertRaises(ToolError):
            self.call("pptx_create", output_path="too-big.pptx", slides=[{"title": "测试", "body": "太长" * 400}])
        self.assertFalse((self.root / "too-big.pptx").exists())
        for slide in ({"title": "汇" * 60}, {"title": "列表", "bullets": ["项目"] * 9}):
            with self.subTest(slide=slide), self.assertRaises(ToolError):
                self.call("pptx_create", output_path="overflow.pptx", slides=[slide])
            self.assertFalse((self.root / "overflow.pptx").exists())

    def test_dense_powerpoint_group_shapes_are_pageable(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        for index in range(400):
            group.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(1)).text = str(index)
        presentation.save(self.root / "dense.pptx")
        first = plugin.call_tool(self.ctx, "pptx_read", {"path": "dense.pptx", "slide_count": 1})
        self.assertFalse(first["isError"], first)
        page = first["structuredContent"]["slides"][0]
        self.assertEqual(page["shape_count"], 400)
        self.assertEqual(len(page["texts"]), 20)
        self.assertTrue(page["selected_shapes"]["has_more"])
        last = self.call("pptx_read", path="dense.pptx", slide_count=1, shape_start=390, shape_count=100)["slides"][0]
        self.assertEqual(len(last["texts"]), 10)
        self.assertEqual(last["texts"][0]["shape_index"], "0.390")
        self.assertEqual(last["texts"][-1]["text"], "399")
        self.assertFalse(last["selected_shapes"]["has_more"])

    def test_corrupt_document_is_actionable_and_no_output_created(self):
        (self.root / "bad.docx").write_bytes(b"not a docx")
        with self.assertRaises(ToolError):
            self.call("docx_read", path="bad.docx")
        self.assertEqual(len(list(self.root.iterdir())), 1)


if __name__ == "__main__":
    unittest.main()
