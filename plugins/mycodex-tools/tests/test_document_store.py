import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from mycodex_tools.common import Context, ToolError
from mycodex_tools.document_store import DocumentStore, atomic_json, digest_file
from mycodex_tools.documents import document_restore, dispatch


class DocumentStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="webcodex-document-store-")
        self.root = Path(self.temp.name).resolve()
        self.ctx = Context(self.root)
        self.path = self.root / "report.docx"
        self.make_doc(self.path, "Original")
        self.sha = digest_file(self.path)

    def tearDown(self):
        self.temp.cleanup()

    @staticmethod
    def make_doc(path, text):
        doc = Document()
        doc.add_paragraph(text)
        doc.save(path)

    def staged(self, store, text="Changed"):
        staged = store.work / "edited.docx"
        self.make_doc(staged, text)
        return staged

    def test_commit_preserves_original_backup_and_replays_receipt(self):
        store = DocumentStore(self.ctx)
        _, snapshot, sha = store.snapshot("report.docx", self.sha)
        self.assertEqual(digest_file(snapshot), self.sha)
        result = store.publish("report.docx", sha, self.staged(store), "fingerprint", {"count": 1})
        self.assertEqual(Document(self.path).paragraphs[0].text, "Changed")
        self.assertEqual(digest_file(self.root / result["backup_path"]), self.sha)
        self.assertEqual(store.receipt("fingerprint")["after_sha256"], result["after_sha256"])
        self.assertTrue(store.receipt("fingerprint")["replayed_receipt"])
        with self.assertRaises(ToolError) as caught:
            store.receipt("different")
        self.assertEqual(caught.exception.code, "request_conflict")

    def test_source_changed_during_edit_is_not_overwritten(self):
        store = DocumentStore(self.ctx)
        store.snapshot("report.docx", self.sha)
        staged = self.staged(store)
        self.make_doc(self.path, "User changed it")
        latest = digest_file(self.path)
        with self.assertRaises(ToolError) as caught:
            store.publish("report.docx", self.sha, staged, "fp", {})
        self.assertEqual(caught.exception.code, "stale_source")
        self.assertEqual(digest_file(self.path), latest)
        self.assertFalse(list(store.backups.glob("*.docx")))

    def test_office_owner_file_blocks_writeback_but_allows_copy(self):
        (self.root / "~$port.docx").write_bytes(b"fixture owner")
        store = DocumentStore(self.ctx)
        with self.assertRaises(ToolError) as caught:
            store.publish("report.docx", self.sha, self.staged(store), "fp", {})
        self.assertEqual(caught.exception.code, "file_busy")
        self.assertEqual(digest_file(self.path), self.sha)
        copy_store = DocumentStore(self.ctx)
        result = copy_store.publish("report.docx", self.sha, self.staged(copy_store), "fp2", {}, "copy", "copy.docx")
        self.assertEqual(result["save_mode"], "copy")

    def test_copy_does_not_change_source_or_overwrite_existing(self):
        store = DocumentStore(self.ctx)
        result = store.publish("report.docx", self.sha, self.staged(store), "fp", {}, "copy", "copy.docx")
        self.assertEqual(digest_file(self.path), self.sha)
        self.assertIsNone(result["backup_id"])
        another = DocumentStore(self.ctx)
        with self.assertRaises(ToolError) as caught:
            another.publish("report.docx", self.sha, self.staged(another), "fp2", {}, "copy", "copy.docx")
        self.assertEqual(caught.exception.code, "target_exists")

    def test_restore_is_guarded_and_backs_up_current_document(self):
        store = DocumentStore(self.ctx)
        edited = store.publish("report.docx", self.sha, self.staged(store), "fp", {})
        listing = document_restore(self.ctx, {"path": "report.docx"})
        self.assertEqual(listing["backups"][0]["backup_id"], edited["backup_id"])
        with self.assertRaises(ToolError) as caught:
            document_restore(self.ctx, {"path": "report.docx", "backup_id": edited["backup_id"],
                                       "expected_sha256": self.sha})
        self.assertEqual(caught.exception.code, "stale_source")
        restored = document_restore(self.ctx, {"path": "report.docx", "backup_id": edited["backup_id"],
                                               "expected_sha256": edited["after_sha256"]})
        self.assertEqual(digest_file(self.path), self.sha)
        self.assertEqual(digest_file(self.root / restored["backup_path"]), edited["after_sha256"])

    def test_corrupt_backup_and_cross_document_restore_are_rejected(self):
        store = DocumentStore(self.ctx)
        result = store.publish("report.docx", self.sha, self.staged(store), "fp", {})
        with self.assertRaises(ToolError) as caught:
            store.backup_snapshot("other.docx", result["backup_id"])
        self.assertEqual(caught.exception.code, "backup_mismatch")
        (self.root / result["backup_path"]).write_bytes(b"corrupt")
        with self.assertRaises(ToolError) as caught:
            store.backup_snapshot("report.docx", result["backup_id"])
        self.assertEqual(caught.exception.code, "backup_corrupt")

    def test_failed_atomic_publication_keeps_original_and_backup(self):
        store = DocumentStore(self.ctx)
        staged = self.staged(store)
        with patch("mycodex_tools.document_store._atomic_replace", side_effect=PermissionError("simulated file lock")):
            with self.assertRaises(ToolError) as caught:
                store.publish("report.docx", self.sha, staged, "fp", {})
        self.assertEqual(caught.exception.code, "file_busy")
        self.assertEqual(digest_file(self.path), self.sha)
        self.assertEqual(len(list(store.backups.glob("*.docx"))), 1)
        with self.assertRaises(ToolError) as caught:
            store.receipt("fp")
        self.assertEqual(caught.exception.code, "outcome_unknown")

    def test_post_commit_crash_reconciles_without_edit_replay(self):
        store = DocumentStore(self.ctx)
        result = store.publish("report.docx", self.sha, self.staged(store), "fp", {})
        receipt = store.receipts / (store.request_id + ".json")
        record = json.loads(receipt.read_text())
        record["state"] = "prepared"
        atomic_json(receipt, record)
        recovered = store.receipt("fp")
        self.assertTrue(recovered["reconciled_receipt"])
        self.assertEqual(recovered["sha256"], result["sha256"])

    def test_historical_receipt_missing_review_is_readable_without_rewriting_history(self):
        store = DocumentStore(self.ctx)
        store.publish("report.docx", self.sha, self.staged(store), "fp", {})
        receipt = store.receipts / (store.request_id + ".json")
        record = json.loads(receipt.read_text())
        for key in ("review", "review_reason", "review_status", "review_material"):
            record["result"].pop(key)
        atomic_json(receipt, record)
        original = receipt.read_bytes()
        restored = store.receipt("fp")
        self.assertEqual(restored["review_record"], "not_recorded")
        self.assertEqual((restored["review"], restored["review_status"]), (1, "pending"))
        self.assertEqual(receipt.read_bytes(), original)

    def test_receipt_failure_after_commit_reports_unknown_not_file_busy(self):
        store = DocumentStore(self.ctx)
        staged = self.staged(store)
        real_write = atomic_json
        def fail_completed(path, value):
            if value.get("state") == "completed":
                raise PermissionError("simulated receipt write failure")
            return real_write(path, value)
        with patch("mycodex_tools.document_store.atomic_json", side_effect=fail_completed):
            with self.assertRaises(ToolError) as caught:
                store.publish("report.docx", self.sha, staged, "fp", {})
        self.assertEqual(caught.exception.code, "outcome_unknown")
        self.assertEqual(Document(self.path).paragraphs[0].text, "Changed")
        self.assertTrue(store.receipt("fp")["reconciled_receipt"])

    def test_paths_cannot_escape_or_read_internal_backups(self):
        store = DocumentStore(self.ctx)
        for path in (str(self.path), "../report.docx", ".webcodex-office/backups/private.docx"):
            with self.assertRaises(ToolError):
                store.source_path(path)

    def test_internal_directory_link_is_rejected(self):
        other = self.root / "other"
        other.mkdir()
        try:
            (self.root / ".webcodex-office").symlink_to(other, target_is_directory=True)
        except OSError:
            self.skipTest("Host does not allow creating symlinks")
        with self.assertRaises(ToolError) as caught:
            DocumentStore(self.ctx)
        self.assertEqual(caught.exception.code, "path_rejected")

    def test_model_cannot_supply_operator_options(self):
        with self.assertRaises(ToolError) as caught:
            dispatch(self.ctx, "present_document", {"path": "report.docx", "options": {"tesseract_path": "evil.exe"}})
        self.assertEqual(caught.exception.code, "invalid_options")

    def test_worker_rejects_result_path_outside_request(self):
        worker = Path(__file__).resolve().parents[1] / "document_worker.py"
        request = {"request_id": "a" * 32, "operation": "present_document", "arguments": {"path": "report.docx"}}
        result = subprocess.run([sys.executable, "-B", str(worker), "--request", "-", "--result", "outside.json"],
                                cwd=self.root, input=json.dumps(request), text=True, capture_output=True, timeout=20)
        self.assertEqual(result.returncode, 1)
        self.assertFalse((self.root / "outside.json").exists())
        self.assertEqual(json.loads(result.stdout)["error"]["code"], "invalid_result_path")


if __name__ == "__main__":
    unittest.main()
