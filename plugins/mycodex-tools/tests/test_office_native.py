"""Document semantics, preservation and opt-in native Office integration tests.

Every artifact is synthetic and lives in TemporaryDirectory. Native tests use
the same hidden COM worker and ownership boundary as production, never GUI input.
"""
from io import BytesIO
import hashlib
import os
from pathlib import Path
import sys
import subprocess
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from mycodex_tools.common import ToolError
from mycodex_tools import office_native as native, office_ooxml as ox


class OfficeNativeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="mycodex-office-test-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def word(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Before ").bold = True
        p.add_run("target").italic = True
        p.add_run(" after").underline = True
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table content"
        doc.add_paragraph("After table")
        doc.sections[0].header.paragraphs[0].text = "Header"
        doc.sections[0].footer.paragraphs[0].text = "Footer"
        source = self.root / "source.docx"
        doc.save(source)
        return source

    def ppt(self):
        pres = Presentation()
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
        p = box.text_frame.paragraphs[0]
        p.add_run().text = "Before "
        run = p.add_run()
        run.text = "target"
        run.font.bold = True
        slide.notes_slide.notes_text_frame.text = "Original notes"
        image = self.root / "old.png"
        Image.new("RGB", (20, 20), "red").save(image)
        pic = slide.shapes.add_picture(str(image), Inches(1), Inches(4), Inches(1), Inches(1))
        chart_data = CategoryChartData()
        chart_data.categories = ["A", "B"]
        chart_data.add_series("First", (1, 2))
        chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(1), Inches(4), Inches(3), chart_data)
        source = self.root / "source.pptx"
        pres.save(source)
        return source, box.shape_id, pic.shape_id, chart.shape_id

    def assert_code(self, code, function, *args):
        with self.assertRaises(ToolError) as raised:
            function(*args)
        self.assertEqual(raised.exception.code, code)

    def test_word_order_all_stories_and_precise_format_preservation(self):
        source = self.word()
        original = source.read_bytes()
        data = native.inspect(source, {"block_count": 100})
        body = [b for b in data["blocks"] if b["part"] == "word/document.xml"]
        self.assertEqual([b["type"] for b in body], ["paragraph", "table", "paragraph", "paragraph"])
        self.assertIn("Header", str(data))
        self.assertIn("Footer", str(data))
        target = self.root / "changed.docx"
        result = native.edit(source, target, [{"op": "replace_text", "target": {"id": body[0]["id"]},
                                               "find": "target", "text": "replacement"}], {})
        self.assertEqual(result["changed_parts"], ["word/document.xml"])
        doc = Document(target)
        self.assertEqual(doc.paragraphs[0].text, "Before replacement after")
        self.assertTrue(doc.paragraphs[0].runs[0].bold)
        self.assertTrue(doc.paragraphs[0].runs[1].italic)
        self.assertTrue(doc.paragraphs[0].runs[2].underline)
        a, b = ox.Package(source), ox.Package(target)
        for name in a.parts:
            if name != "word/document.xml":
                self.assertEqual(a.parts[name], b.parts[name], name)
        self.assertEqual(source.read_bytes(), original)

    def test_word_revisions_textbox_and_explicit_decision(self):
        source = self.word()
        package = ox.Package(source)
        root = package.get("word/document.xml")
        body = root.find("w:body", ox.NS)
        p = ox.ET.Element(ox.q("w", "p"))
        inserted = ox.ET.SubElement(p, ox.q("w", "ins"), {ox.q("w", "id"): "8", ox.q("w", "author"): "Tester"})
        ox.ET.SubElement(ox.ET.SubElement(inserted, ox.q("w", "r")), ox.q("w", "t")).text = "inserted"
        deleted = ox.ET.SubElement(p, ox.q("w", "del"), {ox.q("w", "id"): "9"})
        ox.ET.SubElement(ox.ET.SubElement(deleted, ox.q("w", "r")), ox.q("w", "delText")).text = "deleted"
        body.insert(0, p)
        box = ox.ET.SubElement(ox.ET.SubElement(ox.ET.SubElement(p, ox.q("w", "r")), ox.q("w", "pict")), ox.q("w", "txbxContent"))
        inner = ox.ET.SubElement(box, ox.q("w", "p"))
        ox.ET.SubElement(ox.ET.SubElement(inner, ox.q("w", "r")), ox.q("w", "t")).text = "Textbox text"
        package.put("word/document.xml", root)
        revised = self.root / "revised.docx"
        package.save(revised)
        final = native.inspect(revised, {"block_count": 100, "revision_view": "final"})
        self.assertEqual(final["blocks"][0]["text"], "inserted")
        self.assertTrue(any(b.get("story") == "textbox" and b["text"] == "Textbox text" for b in final["blocks"]))
        self.assertEqual(final["revision_count"], 2)
        target = self.root / "decision.docx"
        native.edit(revised, target, [{"op": "reject_revision", "target": {"id": final["revisions"][0]["id"]}},
                                     {"op": "reject_revision", "target": {"id": final["revisions"][1]["id"]}}], {})
        data = native.inspect(target, {"revision_view": "markup"})
        self.assertNotIn("inserted", data["blocks"][0]["text"])
        self.assertIn("deleted", data["blocks"][0]["text"])
        self.assertEqual(data["revision_count"], 0)

    def test_notes_can_be_created_without_reserializing_slide_shapes(self):
        pres = Presentation()
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Untouched"
        source, output = self.root / "no-notes.pptx", self.root / "notes-added.pptx"
        pres.save(source)
        before = ox.Package(source).parts["ppt/slides/slide1.xml"]
        native.edit(source, output, [{"op": "set_notes", "target": {"slide": 0}, "text": "New notes"}], {})
        self.assertEqual(ox.Package(output).parts["ppt/slides/slide1.xml"], before)
        self.assertEqual(Presentation(output).slides[0].notes_slide.notes_text_frame.text, "New notes")

    def test_empty_ppt_paragraph_keeps_end_format_last_and_com_paths_are_normalized(self):
        paragraph = ox.ET.Element(ox.q("a", "p"))
        ox.ET.SubElement(paragraph, ox.q("a", "endParaRPr"))
        ox.replace_paragraph(paragraph, {"op": "set_text", "text": "Inserted"}, "a")
        self.assertEqual(ox.ET.QName(paragraph[-1]).localname, "endParaRPr")
        self.assertEqual(ox._paragraph_text(paragraph, prefix="a"), "Inserted")
        self.assertEqual(native._com_path(r"\\?\C:\reports\book.xlsx", "excel"), r"C:\reports\book.xlsx")
        self.assertEqual(native._com_path(r"\\?\UNC\host\share\slides.pptx", "powerpoint"), r"\\host\share\slides.pptx")
        self.assert_code("office_path_too_long", native._com_path, "C:\\" + "a"*220 + ".xlsx", "excel")

    def test_word_cell_and_explicit_format(self):
        source = self.word()
        destination = self.root / "format.docx"
        native.edit(source, destination, [
            {"op": "set_cell", "target": {"table": 0, "row": 0, "column": 0}, "text": "New cell"},
            {"op": "set_format", "target": {"id": "word/document.xml#p0"}, "font_name": "Arial", "font_size": 18,
             "bold": False, "font_color": "123456", "alignment": "center"}], {})
        doc = Document(destination)
        self.assertEqual(doc.tables[0].cell(0, 0).text, "New cell")
        self.assertEqual(doc.paragraphs[0].runs[0].font.size.pt, 18)
        self.assertFalse(doc.paragraphs[0].runs[0].bold)

    def test_powerpoint_notes_chart_image_and_untouched_parts(self):
        source, shape_id, image_id, chart_id = self.ppt()
        output = self.root / "updated.pptx"
        image = self.root / "new.png"
        Image.new("RGB", (30, 40), "blue").save(image)
        result = native.edit(source, output, [
            {"op": "replace_text", "target": {"slide": 0, "shape_id": shape_id}, "find": "target", "text": "updated"},
            {"op": "set_notes", "target": {"slide": 0}, "text": "New notes"},
            {"op": "set_geometry", "target": {"slide": 0, "shape_id": shape_id}, "left": 80, "rotation": 15},
            {"op": "set_format", "target": {"slide": 0, "shape_id": shape_id}, "font_size": 28, "font_color": "112233"},
            {"op": "replace_image", "target": {"slide": 0, "shape_id": image_id}, "image_path": str(image)},
            {"op": "set_chart_data", "target": {"slide": 0, "shape_id": chart_id}, "categories": ["X", "Y"],
             "series": [{"name": "New series", "values": [10, 20]}]}], {})
        data = native.inspect(output, {"include_notes": True})
        self.assertIn("New notes", str(data))
        self.assertIn("updated", str(data))
        self.assertIn("New series", str(data))
        self.assertIn("20", str(data))
        before, after = ox.Package(source), ox.Package(output)
        for name in before.parts:
            if name not in result["changed_parts"]:
                self.assertEqual(before.parts[name], after.parts[name], name)
        pres = Presentation(output)
        shape = next(s for s in pres.slides[0].shapes if s.shape_id == shape_id)
        self.assertTrue(shape.text_frame.paragraphs[0].runs[1].font.bold)
        self.assertEqual(shape.text_frame.paragraphs[0].runs[1].font.size.pt, 28)

    def test_pagination_and_source_target_guards(self):
        source = self.word()
        data = native.inspect(source, {"block_count": 1})
        self.assertEqual(data["next_options"]["block_start"], 1)
        self.assert_code("invalid_path", native.edit, source, source,
                         [{"op": "append_paragraph", "text": "bad"}], {})
        self.assert_code("stale_target", native.edit, source, self.root / "bad.docx",
                         [{"op": "set_text", "target": {"id": "word/document.xml#p0"}, "expected_text": "wrong", "text": "new"}], {})
        self.assertFalse((self.root / "bad.docx").exists())

    def test_excel_formula_cache_is_not_misrepresented_as_calculated(self):
        workbook = Workbook()
        workbook.active["A1"] = 7
        workbook.active["B1"] = "=A1*2"
        source = self.root / "formula.xlsx"
        workbook.save(source)
        data = native.inspect(source, {"sheet": "Sheet", "range": "A1:B1"})
        self.assertEqual(data["rows"][0][0]["value"], 7)
        self.assertEqual(data["rows"][0][1]["formula"], "=A1*2")
        self.assertEqual(data["rows"][0][1]["value_status"], "missing_cache")

    def test_native_numeric_guard_rejects_loss_in_unedited_cells(self):
        book = Workbook()
        book.active["A1"], book.active["B1"] = 7, 10
        source = self.root / "precision-source.xlsx"
        book.save(source)
        package = ox.Package(source)
        root = package.get("xl/worksheets/sheet1.xml")
        root.find('.//s:c[@r="B1"]/s:v', ox.NS).text = "0.1234567890123456789"
        package.put("xl/worksheets/sheet1.xml", root)
        exact = self.root / "exact.xlsx"
        package.save(exact)
        output = self.root / "rounded.xlsx"
        book.active["A1"], book.active["B1"] = 9, 0.1234567890123457
        book.save(output)
        self.assert_code("verification_failed", ox.verify_excel_numeric_preservation, exact, output,
                         [{"op": "set_cells", "sheet": "Sheet", "cells": [{"coordinate": "A1", "type": "number", "value": 9}]}])

    def test_new_external_formulas_are_rejected_without_guessing_table_references(self):
        blocked = [r"='C:\outside\[secret.xlsx]Sheet1'!A1", "=[Book.xlsx]Sheet1!A1", "=[1]Sheet1!A1",
                   "=[Book.xlsx]NamedValue", "=[Book.xlsx]名称", '=CUBEVALUE("connection","member")',
                   '=CUBESETCOUNT(A1)', '=RTD("server",,"topic")', '=_xlfn.CUBEMEMBER("connection","member")']
        for formula in blocked:
            with self.subTest(formula=formula):
                self.assertIsNotNone(ox.formula_external_risk(formula, True))
        allowed = ["=SUM(Table1[列])", "=SUM(Table1[[#Headers],[Sales]])", "=[@Column]",
                   "=SUM(Table1[Sales!])", "=SUM(Table1[CUBEVALUE(foo)])", "='Local sheet'!A1+SUM(Table1[列])",
                   '=IF(A1="CUBEVALUE(foo)",1,0)', '="[Book.xlsx]Sheet1!A1"', "=A1*2"]
        for formula in allowed:
            with self.subTest(formula=formula):
                self.assertIsNone(ox.formula_external_risk(formula, True))
        self.assertIsNone(ox.formula_external_risk(blocked[0], False))

    def test_source_cube_preflight_and_external_cache_preservation_are_offline(self):
        workbook = Workbook()
        workbook.active["A1"] = '=CUBEVALUE("connection","member")'
        source = self.root / "cube.xlsx"
        workbook.save(source)
        self.assert_code("active_content_rejected", native._preflight, source, "edit")
        formulas = SimpleNamespace(Cells=[SimpleNamespace(Formula="=[Book.xlsx]Sheet1!A1")])
        used = SimpleNamespace(HasFormula=True, SpecialCells=lambda kind: formulas)
        doc = SimpleNamespace(Worksheets=[SimpleNamespace(UsedRange=used)])
        self.assert_code("external_refresh_rejected", native._assert_no_external_recalculation, doc)
        package = ox.Package(source)
        root = package.get("xl/worksheets/sheet1.xml")
        cell = root.find('.//s:c[@r="A1"]', ox.NS)
        cell.find("s:f", ox.NS).text = "'[Book.xlsx]Sheet1'!A1"
        cell.find("s:v", ox.NS).text = "12"
        package.put("xl/worksheets/sheet1.xml", root)
        ox.edit_workbook_cells(package, [{"sheet": "Sheet", "cells": [{"coordinate": "B1", "type": "number", "value": 9}]}])
        result = package.get("xl/worksheets/sheet1.xml")
        self.assertEqual(result.find('.//s:c[@r="A1"]/s:f', ox.NS).text, "'[Book.xlsx]Sheet1'!A1")
        self.assertEqual(result.find('.//s:c[@r="A1"]/s:v', ox.NS).text, "12")

    def test_active_content_preflight_and_pid_reuse_guard(self):
        source = self.word()
        package = ox.Package(source)
        root = package.get("word/document.xml")
        p = next(root.iter(ox.q("w", "p")))
        ox.ET.SubElement(p, ox.q("w", "instrText")).text = " DDEAUTO test "
        package.put("word/document.xml", root)
        danger = self.root / "active.docx"
        package.save(danger)
        self.assert_code("active_content_rejected", native._preflight, danger, "render")
        with patch.object(native, "_process_identity", return_value={"pid": 42, "created": 2}):
            self.assertFalse(native._terminate_owned({"pid": 42, "created": 1}))

    @unittest.skipUnless(os.name == "nt", "Windows Job lifecycle contract")
    def test_kernel_job_closure_kills_only_owned_process(self):
        child = subprocess.Popen([sys.executable, "-B", "-c", "import time; time.sleep(120)"],
                                 creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        try:
            identity = native._process_identity(child.pid)
            job = native._guard_owned_process(identity)
            job.Close()
            child.wait(timeout=5)
            self.assertIsNotNone(child.returncode)
        finally:
            if child.poll() is None:
                child.kill()
                child.wait()


@unittest.skipUnless(os.environ.get("MYCODEX_NATIVE_OFFICE_TESTS") == "1" and os.name == "nt",
                     "Set MYCODEX_NATIVE_OFFICE_TESTS=1 for owned, hidden native Office fixtures")
class NativeOfficeIntegration(unittest.TestCase):
    setUp = OfficeNativeTests.setUp
    word = OfficeNativeTests.word
    ppt = OfficeNativeTests.ppt
    def test_native_excel_recalculate_and_complex_object_preservation(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"], sheet["B1"] = 3, "=A1*2"
        sheet["C1"] = "unmodified"
        from openpyxl.chart import BarChart, Reference
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=1, max_col=2, min_row=1, max_row=1))
        sheet.add_chart(chart, "E4")
        source, output = self.root / "native.xlsx", self.root / "native-out.xlsx"
        workbook.save(source)
        original = hashlib.sha256(source.read_bytes()).hexdigest()
        result = native.edit(source, output, [
            {"op": "set_cells", "sheet": "Sheet", "cells": [{"coordinate": "A1", "type": "number", "value": 9}]},
            {"op": "set_format", "sheet": "Sheet", "range": "A1", "bold": True, "number_format": "0.00"},
            {"op": "recalculate"}], {"timeout_seconds": 55})
        self.assertTrue(result["calculated"])
        saved = load_workbook(output, data_only=True)
        self.assertEqual(saved.active["B1"].value, 18)
        self.assertEqual(saved.active["C1"].value, "unmodified")
        self.assertEqual(len(saved.active._charts), 1)
        self.assertTrue(saved.active["A1"].font.bold)
        saved.close()
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original)
        legacy = self.root / "legacy.xls"
        native.convert(output, legacy, {"timeout_seconds": 55})
        legacy_read = native.inspect(legacy, {"sheet": "Sheet", "range": "A1:C1", "timeout_seconds": 55})
        self.assertEqual(legacy_read["cells"][0]["value"], 9)

    def test_excel_error_values_are_distinct_from_equal_negative_numbers(self):
        workbook = Workbook()
        workbook.active["A1"] = "=1/0"
        workbook.active["B1"] = -2146826281
        source = self.root / "errors-source.xlsx"
        calculated = self.root / "errors-calculated.xlsx"
        legacy = self.root / "errors-legacy.xls"
        workbook.save(source)
        native.edit(source, calculated, [{"op": "recalculate"}], {"timeout_seconds": 55})
        modern = native.inspect(calculated, {"sheet": "Sheet", "range": "A1:B1"})["rows"][0]
        self.assertEqual(modern[0]["type"], "error")
        self.assertEqual(modern[0]["value"], "#DIV/0!")
        self.assertEqual(modern[0]["display_value"], "#DIV/0!")
        self.assertEqual(modern[0]["formula"], "=1/0")
        self.assertEqual(modern[1]["value"], -2146826281)
        self.assertNotEqual(modern[1]["type"], "error")
        native.convert(calculated, legacy, {"timeout_seconds": 55})
        cells = native.inspect(legacy, {"sheet": "Sheet", "range": "A1:B1", "timeout_seconds": 55})["cells"]
        self.assertEqual(cells[0]["type"], "error")
        self.assertEqual(cells[0]["value"], "#DIV/0!")
        self.assertEqual(cells[0]["display_value"], "#DIV/0!")
        self.assertEqual(cells[0]["formula"], "=1/0")
        self.assertEqual(cells[1]["value"], -2146826281)
        self.assertNotEqual(cells[1].get("type"), "error")

    def test_native_word_render_and_legacy_roundtrip(self):
        source = self.word()
        formatted = self.root / "formatted.docx"
        native.edit(source, formatted, [{"op": "set_format", "target": {"id": "word/document.xml#p0"},
                                        "font_name": "Arial", "font_size": 16, "font_color": "123456", "alignment": "center"}], {})
        source = formatted
        artifacts = native.render(source, self.root, {"page_count": 1, "timeout_seconds": 55})
        self.assertTrue(Path(artifacts[0]["path"]).is_file())
        legacy = self.root / "legacy.doc"
        native.convert(source, legacy, {"timeout_seconds": 55})
        content = native.inspect(legacy, {"timeout_seconds": 55})
        self.assertIn("Before target after", str(content))
        modern = self.root / "modern.docx"
        native.convert(legacy, modern, {"timeout_seconds": 55})
        self.assertEqual(Document(modern).paragraphs[0].text, "Before target after")

    def test_native_powerpoint_render_without_user_instance(self):
        source, shape_id, _, chart_id = self.ppt()
        if native._processes("POWERPNT.EXE"):
            self.skipTest("A pre-existing PowerPoint instance belongs to the user")
        formatted = self.root / "formatted.pptx"
        native.edit(source, formatted, [
            {"op": "set_format", "target": {"slide": 0, "shape_id": shape_id}, "font_name": "Arial", "font_color": "123456", "font_size": 24},
            {"op": "set_chart_data", "target": {"slide": 0, "shape_id": chart_id}, "categories": ["X", "Y"],
             "series": [{"name": "New", "values": [8, 12]}]}], {})
        source = formatted
        artifacts = native.render(source, self.root, {"slide_count": 1, "timeout_seconds": 55})
        self.assertEqual(artifacts[0]["slide"], 0)
        self.assertTrue(Path(artifacts[0]["path"]).is_file())
        legacy = self.root / "legacy.ppt"
        native.convert(source, legacy, {"timeout_seconds": 55})
        self.assertIn("Before target", str(native.inspect(legacy, {"timeout_seconds": 55})))

    def test_native_pdf_to_word_returns_editable_content(self):
        from reportlab.pdfgen import canvas
        source = self.root / "source.pdf"
        pdf = canvas.Canvas(str(source))
        pdf.drawString(72, 720, "Editable PDF conversion fixture")
        pdf.save()
        destination = self.root / "converted.docx"
        with patch.object(native, "_run_native", side_effect=AssertionError("PDF conversion must not start Word")):
            native.convert(source, destination, {"timeout_seconds": 55})
        self.assertIn("Editable PDF conversion fixture", "\n".join(p.text for p in Document(destination).paragraphs))


if __name__ == "__main__":
    unittest.main()
