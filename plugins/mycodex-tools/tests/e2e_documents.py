#!/usr/bin/env python3
"""Real HTTP MCP document workflow acceptance in disposable Server/Runner roots."""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
import os
from pathlib import Path
import sys
import time

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

sys.path.insert(0, str(Path(__file__).resolve().parent))
from e2e_mcp import Harness, McpClient, PLUGIN_ID, content, ok, output, require


class StatelessMcpClient(McpClient):
    def __init__(self, port, token):
        super().__init__(port, token)
        self.protocol = "2026-07-28"

    def rpc(self, method, params=None, timeout=75):
        params = {**(params or {}), "_meta": {
            "io.modelcontextprotocol/protocolVersion": self.protocol,
            "io.modelcontextprotocol/clientCapabilities": {},
        }}
        return super().rpc(method, params, timeout)

    def initialize(self):
        response = self.rpc("server/discover")
        require(response.get("_http_status") == 200 and "result" in response, f"Stateless MCP discovery failed: {response}")
        return {"serverInfo": response["result"].get("_meta", {}).get("io.modelcontextprotocol/serverInfo", {})}


class DocumentHarness(Harness):
    def environment(self):
        env = super().environment()
        if getattr(self.args, "native_user_profile", False):
            for key in ("USERPROFILE", "APPDATA", "LOCALAPPDATA"):
                if os.environ.get(key):
                    env[key] = os.environ[key]
        self.report["office_environment"] = {
            "profile_mode": "signed-in Windows user" if getattr(self.args, "native_user_profile", False) else "synthetic profile directories",
            "reason": "Installed Office COM uses the current Windows identity; fake profile directories caused reproducible Excel/PPT initialization failures.",
            "webcodex_state_and_credentials": "temporary and synthetic",
            "native_assertions_weakened": False,
        }
        if getattr(self.args, "office_diagnostics", False):
            env["MYCODEX_OFFICE_DIAGNOSTICS"] = "1"
        # Only our generated fixture configuration is amended here, before its
        # first Server/Runner launch. Existing services are never reconfigured.
        q = lambda value: json.dumps(str(value))
        worker = self.args.plugin.with_name("document_worker.py")
        worker_args = ["--font-path", str(self.args.font), "--tesseract-path", str(self.args.tesseract),
                       "--tessdata-path", str(self.args.tesseract.parent / "tessdata")]
        extra = f'document_worker = {q(worker)}\ndocument_worker_args = {json.dumps(worker_args)}\n'
        config = self.runner_config.read_text(encoding="utf-8")
        marker = f'cwd = {q(self.project)}\ntimeout_secs = 60\n'
        require(config.count(marker) == 1, "Fixture provider insertion is ambiguous")
        config = config.replace(marker, marker + extra, 1)
        self.other = self.root / "other-project"
        self.other.mkdir()
        (self.root / "registry" / "other.toml").write_text(
            f'id="other"\npath={q(self.other)}\nallow_patch=true\n', encoding="utf-8")
        config = config.replace(f'allowed_roots = [{q(self.project)}]',
                                f'allowed_roots = [{q(self.project)}, {q(self.other)}]')
        config += (f'\n[[plugins.providers]]\nid="office-other"\nname="Other project document tools"\ncommand={q(self.args.python)}\n'
                   f'args=[{q(self.args.plugin)},"--font-path",{q(self.args.font)}]\n'
                   f'cwd={q(self.other)}\ntimeout_secs=60\n' + extra)
        self.runner_config.write_text(config, encoding="utf-8")
        return env

    def seed_auth(self):
        clients = super().seed_auth()
        return {name: StatelessMcpClient(client.port, client.token) for name, client in clients.items()}

    def fixtures(self):
        for root, marker in ((self.project, "FIRST PROJECT"), (self.other, "SECOND PROJECT")):
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run(marker).bold = True
            paragraph.add_run(" original")
            document.sections[0].header.paragraphs[0].text = "Private fixture header"
            document.save(root / "same.docx")
        pdf = canvas.Canvas(str(self.project / "visual.pdf"))
        pdf.drawString(72, 720, "Visual proof: the shape below is not described in text.")
        pdf.setFillColorRGB(1, 0, 0)
        pdf.circle(200, 400, 50, fill=1)
        pdf.save()
        workbook = Workbook()
        workbook.active["A1"] = 6
        workbook.active["A2"] = 2
        workbook.active["B1"] = "=A1*A2"
        workbook.save(self.project / "calc.xlsx")
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        shape.text = "Original slide"
        self.shape_id = shape.shape_id
        presentation.save(self.project / "slides.pptx")
        return {"two_same_name_documents": True, "local_only_fixture_files": True}

    def document_call(self, tool, args, gateway=False):
        arguments = {"project": self.project_id, **args}
        deadline = time.monotonic() + 100
        job_id = None
        while True:
            result = self.client.call("call_runtime_tool", {"tool": tool, "arguments": arguments}) if gateway else self.client.call(tool, arguments)
            value = ok(result)
            if value.get("status") != "pending":
                return result
            require(value.get("job_id"), "Pending document call omitted Job identity")
            if job_id:
                require(value["job_id"] == job_id, "Continuation started another Job")
            job_id = value["job_id"]
            require(time.monotonic() < deadline, "Document Job did not complete in its deadline")
            arguments = {"project": arguments["project"], "job_id": job_id}
            time.sleep(0.1)

    def word_roundtrip(self):
        read = ok(self.document_call("present_document", {"path": "same.docx", "view": "text"}))
        require("FIRST PROJECT" in json.dumps(read), "Wrong primary project content")
        require("Private fixture header" in json.dumps(read), "Header was omitted")
        edited = ok(self.document_call("document_edit", {"path": "same.docx", "expected_sha256": read["source_sha256"],
            "operations": [{"op": "replace_text", "target": {"id": "word/document.xml#p0"}, "find": "original", "text": "updated"}]}))
        saved = Document(self.project / "same.docx")
        require(saved.paragraphs[0].text == "FIRST PROJECT updated" and saved.paragraphs[0].runs[0].bold, "Word text/style did not survive edit")
        require((self.project / edited["backup_path"]).is_file(), "No recoverable original backup")
        restored = ok(self.document_call("document_restore", {"path": "same.docx", "backup_id": edited["backup_id"], "expected_sha256": edited["after_sha256"]}))
        require(restored["after_sha256"] == read["source_sha256"], "Restore did not return exact original bytes")
        return {"writeback": True, "preserved_bold": True, "backup_and_restore": True}

    def project_isolation(self):
        second = ok(self.document_call("present_document", {"project": "agent:" + self.runner_id + ":other", "path": "same.docx", "view": "text"}))
        require("SECOND PROJECT" in json.dumps(second) and "FIRST PROJECT" not in json.dumps(second), "Project binding crossed roots")
        return {"same_name_different_roots": True}

    def images_and_export(self):
        result = self.document_call("present_document", {"path": "visual.pdf", "view": "auto", "export_file": True}, gateway=True)
        value = ok(result)
        images = [item for item in result.get("content", []) if item.get("type") == "image"]
        resources = [item for item in result.get("content", []) if item.get("type") == "resource_link"]
        require(len(images) == 1, "Expected a real MCP native page image, not a path")
        require(len(resources) == 1, "Original-file ResourceLink missing")
        payload = base64.b64decode(images[0]["data"], validate=True)
        require(len(payload) <= 1024 * 1024, "Native image exceeds its byte limit")
        image_path = self.args.report.with_suffix(".mcp-preview.jpg")
        image_path.parent.mkdir(parents=True, exist_ok=True)
        image_path.write_bytes(payload)
        require("red circle" not in json.dumps(value).lower(), "Visual-only fact accidentally leaked into text")
        downloaded = self.client.rpc("resources/read", {"uri": resources[0]["uri"]})
        require("result" in downloaded, "Original resource could not be downloaded")
        contents = downloaded["result"]["contents"]
        binary = base64.b64decode(contents[0]["blob"], validate=True)
        require(hashlib.sha256(binary).hexdigest() == hashlib.sha256((self.project / "visual.pdf").read_bytes()).hexdigest(), "Original export bytes changed")
        return {"native_image": True, "original_resource_bytes_verified": True, "visual_evidence_path": str(image_path)}

    def excel_recalculation(self):
        before = hashlib.sha256((self.project / "calc.xlsx").read_bytes()).hexdigest()
        edited = ok(self.document_call("document_edit", {"path": "calc.xlsx", "expected_sha256": before,
            "operations": [{"op": "set_cells", "sheet": "Sheet", "cells": [{"coordinate": "A1", "type": "number", "value": 9}]}, {"op": "recalculate"}]}))
        workbook = load_workbook(self.project / "calc.xlsx", data_only=True)
        try:
            require(workbook.active["B1"].value == 18, "Excel stored calculation result was not refreshed")
        finally:
            workbook.close()
        require(edited["backup_id"], "Excel original was not backed up")
        return {"formula_result": 18, "original_backup": True}

    def ppt_edit_and_preview(self):
        path = self.project / "slides.pptx"
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        edited = ok(self.document_call("document_edit", {"path": "slides.pptx", "expected_sha256": before,
            "operations": [{"op": "set_text", "target": {"slide": 0, "shape_id": self.shape_id}, "text": "Updated slide"}]}))
        require(Presentation(path).slides[0].shapes[0].text == "Updated slide", "PPT change was not saved")
        preview = self.document_call("present_document", {"path": "slides.pptx", "view": "pages"})
        require(any(item.get("type") == "image" for item in preview["content"]), "PPT preview is not native image content")
        return {"ppt_edit": True, "native_slide_preview": True, "backup": bool(edited["backup_id"])}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("server", "runner", "python", "plugin", "font", "tesseract", "report"):
        parser.add_argument("--" + name, type=Path, required=True)
    parser.add_argument("--surface", default="adaptive-runtime-v1", choices=("adaptive-runtime-v1", "full-operator-v1"))
    parser.add_argument("--only", nargs="+", choices=("discovery", "word_writeback_restore",
                        "same_name_project_isolation", "mcp_native_image_and_original_export",
                        "excel_native_recalculation", "ppt_edit_and_native_preview"),
                        help="Run only selected scenarios; synthetic fixtures are always created.")
    parser.add_argument("--office-diagnostics", action="store_true",
                        help="Include sanitized COM descriptions for these synthetic fixture calls.")
    parser.add_argument("--native-user-profile", action="store_true",
                        help="Use the signed-in Windows profile for installed Office COM; WebCodex state and credentials remain temporary.")
    args = parser.parse_args()
    for name in vars(args):
        if isinstance(getattr(args, name), Path):
            setattr(args, name, getattr(args, name).resolve())
    harness = DocumentHarness(args)
    try:
        harness.setup()
        for name, function in [("discovery", harness.discovery), ("fixtures", harness.fixtures),
                               ("word_writeback_restore", harness.word_roundtrip),
                               ("same_name_project_isolation", harness.project_isolation),
                               ("mcp_native_image_and_original_export", harness.images_and_export),
                               ("excel_native_recalculation", harness.excel_recalculation),
                               ("ppt_edit_and_native_preview", harness.ppt_edit_and_preview)]:
            if args.only and name != "fixtures" and name not in args.only:
                continue
            harness.scenario(name, function)
    except Exception as error:
        harness.report["scenarios"].append(harness.sanitized({"name": "setup", "passed": False, "error": str(error)}))
    finally:
        try:
            harness.close()
        except Exception as error:
            harness.report["scenarios"].append(harness.sanitized({"name": "cleanup", "passed": False, "error": str(error)}))
        harness.report["passed"] = bool(harness.report["scenarios"]) and all(row["passed"] for row in harness.report["scenarios"])
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(harness.sanitized(harness.report), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{'PASS' if harness.report['passed'] else 'FAIL'}: document MCP acceptance; {args.report}")
    return 0 if harness.report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
