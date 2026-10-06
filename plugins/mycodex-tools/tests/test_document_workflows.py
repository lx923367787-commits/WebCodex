"""Real document/worker round trips, isolated from user documents and Office UI."""
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import unittest
import uuid
from unittest.mock import patch
import zipfile

from docx import Document
from PIL import Image
from reportlab.pdfgen import canvas

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from mycodex_tools.common import Context, ToolError
from mycodex_tools.document_store import digest_file
from mycodex_tools.documents import dispatch


class DocumentWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="webcodex-document-workflow-")
        self.root = Path(self.temp.name).resolve()
        self.ctx = Context(self.root)
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("Quarterly ").bold = True
        paragraph.add_run("report")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Amount 100"
        doc.sections[0].header.paragraphs[0].text = "Department header"
        doc.save(self.root / "report.docx")
        pdf = canvas.Canvas(str(self.root / "report.pdf"))
        for index in range(3):
            pdf.drawString(72, 720, f"Native page {index + 1}")
            # A visible graphic that is absent from text extraction.
            pdf.setFillColorRGB(1, 0, 0)
            pdf.circle(200, 400, 50, fill=1)
            pdf.showPage()
        pdf.save()

    def tearDown(self):
        self.temp.cleanup()

    def test_word_read_edit_restore_preserves_unrelated_format(self):
        before = digest_file(self.root / "report.docx")
        read = dispatch(self.ctx, "present_document", {"path": "report.docx", "view": "text"})
        serialized = json.dumps(read["content"])
        self.assertIn("Department header", serialized)
        self.assertIn("Amount 100", serialized)
        result = dispatch(self.ctx, "document_edit", {
            "path": "report.docx", "expected_sha256": read["source_sha256"],
            "operations": [{"op": "replace_text", "target": {"id": "word/document.xml#p0"},
                            "find": "report", "text": "summary"}],
        })
        self.assertEqual(result["save_mode"], "in_place")
        updated = Document(self.root / "report.docx")
        self.assertEqual(updated.paragraphs[0].text, "Quarterly summary")
        self.assertTrue(updated.paragraphs[0].runs[0].bold)
        self.assertEqual(updated.tables[0].cell(0, 0).text, "Amount 100")
        self.assertEqual(digest_file(self.root / result["backup_path"]), before)
        restored = dispatch(self.ctx, "document_restore", {
            "path": "report.docx", "expected_sha256": result["after_sha256"], "backup_id": result["backup_id"],
        })
        self.assertEqual(restored["after_sha256"], before)

    def test_pdf_presentation_returns_real_bounded_images_and_cursor(self):
        first = dispatch(self.ctx, "present_document", {"path": "report.pdf"})
        self.assertEqual(first["visual_delivery"], "ready")
        self.assertEqual(len(first["previews"]), 2)
        self.assertIn("Native page 1", json.dumps(first["content"]))
        for preview in first["previews"]:
            path = self.root / preview["path"]
            self.assertLessEqual(path.stat().st_size, 1024 * 1024)
            self.assertEqual(digest_file(path), preview["sha256"])
            self.assertEqual(preview["source_sha256"], first["source_sha256"])
            with Image.open(path) as image:
                self.assertGreater(image.width, 300)
                self.assertGreater(image.height, 300)
        self.assertIsNotNone(first["next_cursor"])
        following = dispatch(self.ctx, "present_document", {"path": "report.pdf", "cursor": first["next_cursor"]})
        self.assertIn("Native page 3", json.dumps(following["content"]))

    def test_cursor_cannot_be_used_for_another_document_or_new_version(self):
        first = dispatch(self.ctx, "present_document", {"path": "report.pdf", "view": "text"})
        with self.assertRaises(ToolError) as caught:
            dispatch(self.ctx, "present_document", {"path": "report.docx", "cursor": first["next_cursor"]})
        self.assertEqual(caught.exception.code, "cursor_mismatch")
        with (self.root / "report.pdf").open("ab") as stream:
            stream.write(b"\n% local saved edit\n")
        with self.assertRaises(ToolError) as caught:
            dispatch(self.ctx, "present_document", {"path": "report.pdf", "cursor": first["next_cursor"]})
        self.assertEqual(caught.exception.code, "stale_source")

    def test_auto_view_handles_single_page_without_forcing_two_pages(self):
        pdf = canvas.Canvas(str(self.root / "one.pdf"))
        pdf.drawString(72, 720, "Only page")
        pdf.save()
        result = dispatch(self.ctx, "present_document", {"path": "one.pdf"})
        self.assertEqual(result["visual_delivery"], "ready")
        self.assertEqual(len(result["previews"]), 1)

    def test_pdf_edit_is_real_replacement_and_copy_mode_keeps_source(self):
        before = digest_file(self.root / "report.pdf")
        result = dispatch(self.ctx, "document_edit", {
            "path": "report.pdf", "expected_sha256": before, "output_path": "edited.pdf",
            "operations": [{"op": "replace_text", "page": 1, "old_text": "Native page 1", "new_text": "New page 1"}],
        })
        self.assertEqual(result["save_mode"], "copy")
        self.assertEqual(digest_file(self.root / "report.pdf"), before)
        content = dispatch(self.ctx, "present_document", {"path": "edited.pdf", "view": "text"})
        self.assertIn("New page 1", json.dumps(content["content"]))
        self.assertNotIn("Native page 1", json.dumps(content["content"]))

    def test_auxiliary_file_inputs_must_be_project_relative(self):
        with self.assertRaises(ToolError):
            dispatch(self.ctx, "document_edit", {
                "path": "report.pdf", "expected_sha256": digest_file(self.root / "report.pdf"),
                "operations": [{"op": "merge", "paths": [str(self.root / "report.pdf")]}],
            })

    def test_operation_receipt_prevents_repeating_successful_mutation(self):
        operation_id = uuid.uuid4().hex
        args = {"path": "report.pdf", "expected_sha256": digest_file(self.root / "report.pdf"),
                "operations": [{"op": "rotate_pages", "pages": [1], "angle": 90}]}
        first = dispatch(self.ctx, "document_edit", args, request_id=operation_id)
        again = dispatch(self.ctx, "document_edit", args, request_id=operation_id)
        self.assertTrue(again["replayed_receipt"])
        self.assertEqual(first["after_sha256"], digest_file(self.root / "report.pdf"))

    def test_request_id_is_bound_before_concurrent_backend_work(self):
        from mycodex_tools import pdf_native
        started, release = threading.Event(), threading.Event()
        request_id = uuid.uuid4().hex
        args = {"path": "report.pdf", "expected_sha256": digest_file(self.root / "report.pdf"),
                "operations": [{"op": "rotate_pages", "pages": [1], "angle": 90}]}
        real_edit = pdf_native.edit
        def delayed(*positional):
            started.set()
            if not release.wait(10):
                raise RuntimeError("test release timeout")
            return real_edit(*positional)
        completed = []
        def first():
            try:
                completed.append(dispatch(self.ctx, "document_edit", args, request_id=request_id))
            except BaseException as error:
                completed.append(error)
        with patch("mycodex_tools.pdf_native.edit", side_effect=delayed) as mocked:
            thread = threading.Thread(target=first)
            thread.start()
            self.assertTrue(started.wait(10))
            try:
                with self.assertRaises(ToolError) as caught:
                    dispatch(self.ctx, "document_edit", args, request_id=request_id)
                self.assertEqual(caught.exception.code, "file_busy")
            finally:
                release.set()
                thread.join(10)
            self.assertEqual(mocked.call_count, 1)
        self.assertIsInstance(completed[0], dict)
        with self.assertRaises(ToolError) as caught:
            dispatch(self.ctx, "document_edit", {**args, "output_path": "elsewhere.pdf"}, request_id=request_id)
        self.assertEqual(caught.exception.code, "request_conflict")

    def test_worker_stdout_is_small_and_manifest_contains_source_bound_preview(self):
        request_id = uuid.uuid4().hex
        relative = f".webcodex-office/cache/{request_id}/result.json"
        worker = Path(__file__).resolve().parents[1] / "document_worker.py"
        request = {"request_id": request_id, "operation": "present_document", "arguments": {"path": "report.pdf"}}
        process = subprocess.run([sys.executable, "-B", str(worker), "--request", "-", "--result", relative],
                                 cwd=self.root, input=json.dumps(request).encode(), capture_output=True, timeout=30)
        self.assertEqual(process.returncode, 0, process.stderr.decode(errors="replace"))
        self.assertLess(len(process.stdout), 1024)
        manifest = json.loads((self.root / relative).read_text(encoding="utf-8"))
        self.assertTrue(manifest["success"])
        self.assertEqual(manifest["request_id"], request_id)
        self.assertEqual(manifest["output"]["source_sha256"], digest_file(self.root / "report.pdf"))
        self.assertEqual(len(manifest["output"]["previews"]), 2)

    def test_ordinary_read_failure_remains_structured_in_worker_manifest(self):
        request_id = uuid.uuid4().hex
        relative = f".webcodex-office/cache/{request_id}/result.json"
        worker = Path(__file__).resolve().parents[1] / "document_worker.py"
        request = {"request_id": request_id, "operation": "document_edit", "arguments": {"path": "report.docx", "operations": []}}
        process = subprocess.run([sys.executable, "-B", str(worker), "--request", "-", "--result", relative],
                                 cwd=self.root, input=json.dumps(request).encode(), capture_output=True, timeout=20)
        self.assertEqual(process.returncode, 1)
        manifest = json.loads((self.root / relative).read_text(encoding="utf-8"))
        self.assertFalse(manifest["success"])
        self.assertEqual(manifest["error"]["code"], "expected_sha256_required")

    def test_review_is_strict_integer_and_read_only_defaults_stay_lightweight(self):
        for value in (True, False, None, "1", 1.0, -1, 2):
            with self.subTest(review=value), self.assertRaises(ToolError) as caught:
                dispatch(self.ctx, "present_document", {"path": "report.pdf", "review": value})
            self.assertEqual(caught.exception.code, "invalid_review")
        ordinary = dispatch(self.ctx, "present_document", {"path": "report.pdf", "view": "text"})
        self.assertEqual((ordinary["review"], ordinary["review_status"]), (0, "not_required"))
        requested = dispatch(self.ctx, "present_document", {"path": "report.pdf", "review": 1})
        self.assertEqual((requested["review"], requested["review_status"]), (1, "pending"))
        self.assertEqual(requested["review_material"]["delivery_status"], "awaiting_native_image_delivery")

    def test_layout_override_is_in_durable_receipt_not_only_response(self):
        request_id = uuid.uuid4().hex
        arguments = {"path": "report.pdf", "review": 0,
                     "expected_sha256": digest_file(self.root / "report.pdf"),
                     "operations": [{"op": "rotate_pages", "pages": [1], "angle": 90}]}
        first = dispatch(self.ctx, "document_edit", arguments, request_id=request_id)
        self.assertEqual((first["review"], first["review_status"]), (1, "pending"))
        self.assertIn("corrected", first["review_reason"])
        material = first["review_material"]
        self.assertEqual(material["source_sha256"], digest_file(self.root / "report.pdf"))
        self.assertEqual(material["coverage"]["prepared"], [])
        receipt = json.loads((self.root / ".webcodex-office" / "receipts" / (request_id + ".json")).read_text())
        self.assertEqual(receipt["result"]["review_material"], material)
        replay = dispatch(self.ctx, "document_edit", arguments, request_id=request_id)
        self.assertEqual(replay["review_material"], material)

    def test_all_pages_and_crop_have_distinct_source_bound_coverage(self):
        first = dispatch(self.ctx, "present_document", {"path": "report.pdf", "view": "pages", "review": 1,
                                                       "options": {"page_count": 2}})
        coverage = first["review_material"]["coverage"]
        self.assertEqual(coverage["required"], ["page:1", "page:2", "page:3"])
        self.assertEqual(coverage["missing"], ["page:3"])
        self.assertFalse(coverage["complete"])
        following = dispatch(self.ctx, "present_document", {"path": "report.pdf", "cursor": first["next_cursor"], "review": 1})
        self.assertEqual(following["review_material"]["coverage"]["prepared"], ["page:3"])
        self.assertIsNone(following["next_cursor"])
        self.assertEqual(first["review_material"]["render_dependencies_sha256"], following["review_material"]["render_dependencies_sha256"])
        cropped = dispatch(self.ctx, "present_document", {"path": "report.pdf", "review": 1,
                            "options": {"page_count": 1, "crop": [0, 0, 400, 400]}})
        self.assertIn("page:1", cropped["review_material"]["coverage"]["missing"])
        self.assertTrue(cropped["previews"][0]["coverage_id"].startswith("page:1:crop:"))
        for preview in first["previews"]:
            self.assertEqual(preview["render_source_sha256"], first["source_sha256"])
            self.assertEqual(preview["render_source_path"], "report.pdf")
            self.assertEqual(digest_file(self.root / preview["path"]), preview["sha256"])

    def test_renderer_unavailable_is_explicit_pending_not_approval(self):
        with patch("mycodex_tools.office_native.render", side_effect=ToolError("unsupported_render", "No renderer")):
            result = dispatch(self.ctx, "present_document", {"path": "report.docx", "view": "pages", "review": 1})
        self.assertEqual(result["review_status"], "pending")
        self.assertEqual(result["visual_delivery"], "unavailable")
        self.assertEqual(result["previews"], [])
        self.assertEqual(result["review_material"]["coverage"]["missing"], ["all_pages:unresolved"])
        self.assertEqual(result["warnings"][0]["code"], "unsupported_render")

    def test_independent_pdf_preview_cannot_prove_real_pptx(self):
        from pptx import Presentation
        from mycodex_tools import pdf_native
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[0])
        deck.save(self.root / "report.pptx")
        def wrong_renderer(source, output, options):
            return pdf_native.render(self.root / "report.pdf", output, {"page_count": 1})
        with patch("mycodex_tools.office_native.render", side_effect=wrong_renderer):
            with self.assertRaises(ToolError) as caught:
                dispatch(self.ctx, "present_document", {"path": "report.pptx", "review": 1})
        self.assertEqual(caught.exception.code, "render_source_mismatch")

    def test_corrupt_workbook_and_manifest_only_zip_cannot_be_published(self):
        from openpyxl import Workbook
        from mycodex_tools.document_store import DocumentStore
        Workbook().save(self.root / "report.xlsx")
        (self.root / "corrupt.xlsx").write_bytes(b"not a real workbook")
        with self.assertRaises(ToolError):
            dispatch(self.ctx, "present_document", {"path": "corrupt.xlsx", "review": 1})
        store = DocumentStore(self.ctx)
        staged = store.work / "broken.xlsx"
        with zipfile.ZipFile(staged, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("manifest.json", json.dumps({"verified": True, "sheet_count": 8}))
        before = digest_file(self.root / "report.xlsx")
        with self.assertRaises((ToolError, KeyError, ValueError)):
            store.publish("report.xlsx", before, staged, "fixture", {"verified": True})
        self.assertEqual(digest_file(self.root / "report.xlsx"), before)
        self.assertFalse((store.receipts / (store.request_id + ".json")).exists())

    def test_excel_print_pages_cannot_hide_uncovered_sheet_regions(self):
        from openpyxl import Workbook
        from mycodex_tools.document_review import material
        workbook = Workbook()
        workbook.active["A1"] = "Data outside declared print area"
        workbook.create_sheet("Second")
        workbook.save(self.root / "report.xlsx")
        result = dispatch(self.ctx, "present_document", {"path": "report.xlsx", "view": "text", "review": 1})
        sha = result["source_sha256"]
        rendered = material("report.xlsx", sha, result["content"],
            [{"page": 1, "page_count_total": 1, "render_source_sha256": sha}], {}, "a" * 64)
        self.assertEqual(rendered["coverage"]["prepared"], ["page:1"])
        self.assertEqual(len(rendered["coverage"]["missing"]), 2)
        self.assertFalse(rendered["coverage"]["complete"])

    def test_review_probe_rechecks_current_dependencies_without_rendering(self):
        dependency = self.root / "configured-font.fixture"
        dependency.write_bytes(b"renderer dependency version one")
        def environment(_source):
            return {"actual_file_sha256": digest_file(dependency)}
        with patch("mycodex_tools.document_review._environment_fingerprints", side_effect=environment):
            shown = dispatch(self.ctx, "present_document", {"path": "report.pdf", "view": "pages", "review": 1})
            arguments = {"path": "report.pdf", "view": "text", "review": 1,
                         "expected_sha256": shown["source_sha256"],
                         "options": {**shown["review_material"]["render_options"], "review_probe": True}}
            with patch("mycodex_tools.pdf_native.render", side_effect=AssertionError("probe must not render")):
                fresh = dispatch(self.ctx, "present_document", arguments)
                self.assertEqual(fresh["review_probe"]["status"], "verified")
                self.assertEqual(fresh["previews"], [])
                self.assertEqual(fresh["review_probe"]["render_dependencies_sha256"], shown["review_material"]["render_dependencies_sha256"])
                dependency.write_bytes(b"renderer dependency version two")
                changed = dispatch(self.ctx, "present_document", arguments)
                self.assertNotEqual(changed["review_probe"]["render_dependencies_sha256"], fresh["review_probe"]["render_dependencies_sha256"])
                self.assertEqual(changed["source_sha256"], shown["source_sha256"])

    def test_unavailable_dependencies_leave_no_usable_review_receipt(self):
        with patch("mycodex_tools.document_review._environment_fingerprints", side_effect=ToolError("dependency_probe_unavailable", "fixture unavailable")):
            result = dispatch(self.ctx, "present_document", {"path": "report.pdf", "review": 1})
        self.assertEqual(result["review_status"], "pending")
        self.assertEqual(result["visual_delivery"], "unavailable")
        self.assertEqual(result["previews"], [])
        self.assertIsNone(result["review_material"]["render_dependencies_sha256"])
        self.assertEqual(result["review_material"]["coverage"]["missing"], ["page:1", "page:2", "page:3"])
        with patch("mycodex_tools.document_review._environment_fingerprints", side_effect=ToolError("dependency_probe_unavailable", "fixture unavailable")):
            ordinary = dispatch(self.ctx, "present_document", {"path": "report.pdf", "view": "pages"})
        self.assertEqual(ordinary["review"], 0)
        self.assertEqual(ordinary["visual_delivery"], "ready")
        self.assertEqual(len(ordinary["previews"]), 2)
        self.assertFalse(ordinary["review_material"]["dependencies_verified"])
        self.assertIsNone(ordinary["previews"][0]["render_dependencies_sha256"])

    def test_replayed_presentation_preserves_original_images_and_raw_material(self):
        request_id = uuid.uuid4().hex
        arguments = {"path": "report.pdf", "review": 1, "view": "pages"}
        original = dispatch(self.ctx, "present_document", arguments, request_id=request_id)
        image_bytes = [(self.root / preview["path"]).read_bytes() for preview in original["previews"]]
        with patch("mycodex_tools.pdf_native.render", side_effect=AssertionError("an archived read must not rerender")):
            repeated = dispatch(self.ctx, "present_document", arguments, request_id=request_id)
        self.assertEqual(original, repeated)
        self.assertEqual(image_bytes, [(self.root / preview["path"]).read_bytes() for preview in repeated["previews"]])


if __name__ == "__main__":
    unittest.main()
